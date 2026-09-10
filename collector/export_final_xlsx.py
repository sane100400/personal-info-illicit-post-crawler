#!/usr/bin/env python3
"""Build a minimal public-release XLSX from the masked positive corpus."""

from __future__ import annotations

import argparse
import csv
from collections import Counter
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


PUBLIC_COLUMNS = [
    ("번호", 9),
    ("수집일", 13),
    ("게시 위치 유형", 16),
    ("거래 대상", 24),
    ("마스킹 제목", 42),
    ("마스킹 본문", 82),
    ("비슷한 게시글 여부", 22),
    ("같거나 비슷한 게시글 수", 24),
]


VALUE_LABELS = {
    "source_unit_kind": {
        "board": "게시판",
        "site": "독립 사이트",
        "social_account": "공개 SNS",
    },
    "target_category": {
        "personal_data_db": "개인정보 DB",
        "bulk_account_or_authentication": "대량 계정·인증정보",
        "financial_account": "통장·계좌",
        "identity_document": "신원 문서",
    },
}


HEADER_FILL = "D9E2F3"
HEADER_BORDER = Side(style="thin", color="A6A6A6")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--input",
        type=Path,
        default=Path("output/final/public/positive_posts_masked.csv"),
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("output/final/논문_최종데이터.xlsx"),
    )
    return parser.parse_args()


def load_rows(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def public_rows(rows: list[dict[str, str]]) -> list[list[object]]:
    exact_counts = Counter(row["review_unit_id"] for row in rows)
    similar_counts = Counter(row["near_duplicate_fingerprint"] for row in rows)
    exported: list[list[object]] = []

    for number, row in enumerate(rows, start=1):
        exact_count = exact_counts[row["review_unit_id"]]
        similar_count = similar_counts[row["near_duplicate_fingerprint"]]
        if exact_count > 1:
            repetition = "제목·본문이 같음"
            repetition_count = exact_count
        elif similar_count > 1:
            repetition = "제목·본문이 비슷함"
            repetition_count = similar_count
        else:
            repetition = "없음"
            repetition_count = 1

        collected_date = date.fromisoformat(row["collected_at"][:10])
        exported.append(
            [
                number,
                collected_date,
                VALUE_LABELS["source_unit_kind"].get(
                    row["source_unit_kind"], row["source_unit_kind"]
                ),
                VALUE_LABELS["target_category"].get(
                    row["target_category"], row["target_category"]
                ),
                row["masked_title"],
                row["masked_text"],
                repetition,
                repetition_count,
            ]
        )
    return exported


def build_workbook(rows: list[list[object]]) -> Workbook:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "공개데이터"
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = True
    sheet.sheet_view.zoomScale = 90

    for column_index, (title, width) in enumerate(PUBLIC_COLUMNS, start=1):
        cell = sheet.cell(1, column_index, title)
        cell.font = Font(name="맑은 고딕", size=10, bold=True)
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=HEADER_BORDER)
        sheet.column_dimensions[get_column_letter(column_index)].width = width
    sheet.row_dimensions[1].height = 24

    for row_index, values in enumerate(rows, start=2):
        for column_index, value in enumerate(values, start=1):
            cell = sheet.cell(row_index, column_index, value)
            cell.font = Font(name="맑은 고딕", size=9)
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=column_index in {5, 6},
            )
            if column_index == 2:
                cell.number_format = "yyyy-mm-dd"
        sheet.row_dimensions[row_index].height = 36

    last_column = get_column_letter(len(PUBLIC_COLUMNS))
    sheet.auto_filter.ref = f"A1:{last_column}{len(rows) + 1}"
    return workbook


def validate(path: Path, expected_rows: int) -> None:
    workbook = load_workbook(path, read_only=False, data_only=False)
    if workbook.sheetnames != ["공개데이터"]:
        raise ValueError("unexpected worksheet structure")
    sheet = workbook["공개데이터"]
    if sheet.max_row != expected_rows + 1 or sheet.max_column != len(PUBLIC_COLUMNS):
        raise ValueError("XLSX dimensions do not match the public dataset")
    headers = [sheet.cell(1, column).value for column in range(1, len(PUBLIC_COLUMNS) + 1)]
    if headers != [title for title, _ in PUBLIC_COLUMNS]:
        raise ValueError("public headers do not match")
    if sheet._charts or sheet.tables or sheet.merged_cells.ranges:
        raise ValueError("the public workbook contains unexpected decorations")
    if any(
        sheet.column_dimensions[get_column_letter(index)].hidden
        for index in range(1, len(PUBLIC_COLUMNS) + 1)
    ):
        raise ValueError("the public workbook contains hidden columns")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if cell.data_type == "f":
                raise ValueError("formula cell found in public workbook")
    workbook.close()


def main() -> int:
    args = parse_args()
    source_rows = load_rows(args.input)
    rows = public_rows(source_rows)
    workbook = build_workbook(rows)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output)
    validate(args.output, len(rows))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
