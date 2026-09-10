from __future__ import annotations

import base64
import csv
import datetime as dt
import json
import tempfile
import unittest
from collections import Counter
from pathlib import Path
from unittest.mock import Mock, patch
from urllib.parse import quote_plus

from openpyxl import load_workbook
from selenium.common.exceptions import WebDriverException

from collector.build_labeling_pilot import (
    SourceRow,
    assign_near_duplicate_clusters,
    intent_bucket,
    prioritize_rows,
    select_rows,
    strict_bucket,
    write_labeling_sheet,
)
from collector.labeling_workbook import write_labeling_workbook

from collector.collect_candidates import (
    Candidate,
    backup_candidate_queue,
    CollectionLog,
    DetectionEntry,
    QuerySpec,
    RESTRICTED_REVIEW_SCHEMA,
    SCHEMA,
    append_detection_entries,
    attributed_destination_url,
    campaign_fingerprint_text,
    canonicalize_url,
    classify_page_type,
    constrain_query_specs,
    contact_campaign_id,
    data_manifest,
    effective_domain_record_limit,
    effective_minimum_domains,
    discover_candidates,
    discover_google_api_candidates,
    discover_serpapi_candidates,
    discover_related_internal_links,
    discovery_candidate_relevant,
    discovery_candidate_passes,
    discovery_relevance_score,
    expand_query_specs,
    extract_title_text,
    extraction_failure_record,
    exclude_known_source_unit_candidates,
    existing_fingerprints,
    infer_collection_type,
    interleave_candidates_by_domain,
    load_candidate_queue,
    load_excluded_fingerprints,
    load_excluded_urls,
    load_query_specs,
    load_seed_candidates,
    limit_query_specs_by_group,
    mask_text,
    masking_validation,
    make_record,
    merge_candidates,
    mine_keyword_expansions,
    near_duplicate_id,
    nate_search_url,
    ordered_provider_names,
    post_identity_descriptor,
    prefilter_seed_candidates,
    prepare_detection_workbook,
    prioritize_candidates_by_domain_deficit,
    public_content_fallback_url,
    relevance_gate_reason,
    read_html,
    render_public_text,
    registrable_domain,
    revalidate_existing_records,
    safe_spreadsheet_text,
    save_candidate_queue,
    save_restricted_workbook,
    should_reserve_for_domain_diversity,
    source_unit_descriptor,
    source_unit_token,
    text_quality_reason,
    terminal_attempt_hashes,
    unwrap_search_result_url,
    upgrade_collection_log_schema,
    upgrade_existing_csv_schema,
    url_in_excluded_domain,
    yahoo_japan_search_url,
    write_collector_labeling_workbook,
)

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "(양식) 탐지내역.xlsx"


class CollectorTests(unittest.TestCase):
    def test_metadata_result_url_resolves_attributed_public_site(self) -> None:
        self.assertEqual(
            attributed_destination_url("https://host.io/ggidmc.com"),
            "https://ggidmc.com/",
        )
        self.assertEqual(
            attributed_destination_url("https://www.host.io/store.example.co.kr"),
            "https://store.example.co.kr/",
        )

    def test_metadata_result_url_rejects_non_domain_paths(self) -> None:
        invalid = [
            "https://host.io/127.0.0.1",
            "https://host.io/example",
            "https://host.io/example.com/details",
            "https://host.io/example.com%2Fadmin",
        ]
        for url in invalid:
            self.assertEqual(attributed_destination_url(url), url)
        ordinary = "https://example.com/host.io/ggidmc.com"
        self.assertEqual(attributed_destination_url(ordinary), ordinary)

    def test_html_response_cap_allows_template_heavy_storefronts(self) -> None:
        class FakeResponse:
            headers = {"Content-Type": "text/html; charset=UTF-8"}
            encoding = "utf-8"
            apparent_encoding = "utf-8"

            def iter_content(self, _chunk_size: int):
                yield b"<html>"
                yield b"x" * 1_100_000
                yield b"</html>"

        html, reason = read_html(FakeResponse())
        self.assertEqual(reason, "")
        self.assertIsNotNone(html)
        self.assertGreater(len(html or ""), 1_000_000)

    def test_html_response_cap_still_rejects_oversized_pages(self) -> None:
        class FakeResponse:
            headers = {
                "Content-Type": "text/html; charset=UTF-8",
                "Content-Length": "2000001",
            }

        html, reason = read_html(FakeResponse())
        self.assertIsNone(html)
        self.assertEqual(reason, "content_too_large")

    def test_html_meta_charset_overrides_requests_latin1_default(self) -> None:
        korean = "민증위조 제작 전문업체"
        payload = (
            '<html><head><meta http-equiv="Content-Type" '
            'content="text/html; charset=euc-kr"></head><body>'
            + korean
            + "</body></html>"
        ).encode("euc-kr")

        class FakeResponse:
            headers = {"Content-Type": "text/html"}
            encoding = "ISO-8859-1"
            apparent_encoding = "EUC-KR"

            def iter_content(self, _chunk_size: int):
                yield payload

        html, reason = read_html(FakeResponse())
        self.assertEqual(reason, "")
        self.assertIn(korean, html or "")

    def test_missing_content_type_accepts_clear_legacy_html(self) -> None:
        payload = (
            "2005\r\n<html><head><meta http-equiv=Content-Type "
            "content='text/html; charset=euc-kr'></head>"
            "<body>개인 법인 통장 판매합니다</body></html>"
        ).encode("euc-kr")

        class FakeResponse:
            headers = {}
            encoding = None
            apparent_encoding = "EUC-KR"

            def iter_content(self, _chunk_size: int):
                yield payload

        html, reason = read_html(FakeResponse())
        self.assertEqual(reason, "")
        self.assertIn("통장 판매합니다", html or "")

    def test_missing_content_type_does_not_accept_arbitrary_bytes(self) -> None:
        class FakeResponse:
            headers = {}
            encoding = None
            apparent_encoding = "utf-8"

            def iter_content(self, _chunk_size: int):
                yield b"plain response without html structure"

        html, reason = read_html(FakeResponse())
        self.assertIsNone(html)
        self.assertEqual(reason, "non_html_content")

    def test_mislabeled_xml_feed_is_not_parsed_as_html(self) -> None:
        class FakeResponse:
            headers = {"Content-Type": "text/html; charset=UTF-8"}
            encoding = "utf-8"
            apparent_encoding = "utf-8"

            def iter_content(self, _chunk_size: int):
                yield (
                    b'<?xml version="1.0" encoding="UTF-8"?>'
                    b'<rss version="2.0"><channel><title>DB sale feed</title>'
                    b'</channel></rss>'
                )

        html, reason = read_html(FakeResponse())
        self.assertIsNone(html)
        self.assertEqual(reason, "non_html_content")

    def test_extraction_removes_css_before_account_relevance_checks(self) -> None:
        html = """
        <html><head><title>가방 대량 주문 문의</title>
        <style>@media screen { .account-sale { content: '계정 판매'; } }</style>
        </head><body><main>
        프리미엄 여행 가방과 백팩을 기업 고객에게 대량 주문으로 제공합니다.
        제품 색상, 수량, 각인과 배송 일정을 문의 양식에 작성해 주세요.
        정상적인 여행용품 주문 서비스입니다.
        </main></body></html>
        """
        title, text, _ = extract_title_text(html, "https://brand.example/bulk")
        self.assertNotIn("계정 판매", text)
        self.assertNotIn("@media", text)
        self.assertNotEqual(
            relevance_gate_reason(
                title,
                text,
                "https://brand.example/bulk",
                "unknown",
                "intent",
            ),
            "",
        )

    def test_registrable_domain_uses_public_suffix_rules(self) -> None:
        self.assertEqual(registrable_domain("forum.audio.com.pl"), "audio.com.pl")
        self.assertEqual(registrable_domain("m.example.co.kr"), "example.co.kr")
        self.assertEqual(registrable_domain("sub.example.com"), "example.com")
        self.assertEqual(registrable_domain("127.0.0.1"), "127.0.0.1")

    def test_excluded_domain_applies_to_subdomains_only(self) -> None:
        excluded = {"dcinside.com"}
        self.assertTrue(
            url_in_excluded_domain(
                "https://gall.dcinside.com/board/view/?id=test&no=1",
                excluded,
            )
        )
        self.assertFalse(
            url_in_excluded_domain("https://example.com/dcinside.com", excluded)
        )

    def test_five_percent_domain_share_requires_twenty_domains_for_500(self) -> None:
        limit = effective_domain_record_limit(500, 0, 0.05)
        self.assertEqual(limit, 25)
        self.assertEqual(effective_minimum_domains(500, limit, 0), 20)

    def test_absolute_domain_limit_uses_the_stricter_setting(self) -> None:
        self.assertEqual(effective_domain_record_limit(500, 40, 0.05), 25)
        self.assertEqual(effective_domain_record_limit(500, 10, 0.05), 10)

    def test_candidates_are_interleaved_and_capped_by_domain(self) -> None:
        candidates = [
            Candidate(f"https://one.example/{index}", "g", "기타")
            for index in range(4)
        ] + [
            Candidate(f"https://two.example/{index}", "g", "기타")
            for index in range(2)
        ]
        ordered = interleave_candidates_by_domain(candidates, max_per_domain=2)
        self.assertEqual(
            [row.url for row in ordered],
            [
                "https://one.example/0",
                "https://two.example/0",
                "https://one.example/1",
                "https://two.example/1",
            ],
        )

    def test_unseen_domains_are_prioritized_until_domain_floor_is_met(self) -> None:
        candidates = [
            Candidate("https://known.example/new-board/1", "g", "기타"),
            Candidate("https://fresh-one.example/post/1", "g", "기타"),
            Candidate("https://known.example/new-board/2", "g", "기타"),
            Candidate("https://fresh-two.example/post/1", "g", "기타"),
        ]
        ordered = prioritize_candidates_by_domain_deficit(
            candidates,
            {"known.example"},
            minimum_domains=3,
        )
        self.assertEqual(
            [item.url for item in ordered],
            [
                "https://fresh-one.example/post/1",
                "https://fresh-two.example/post/1",
                "https://known.example/new-board/1",
                "https://known.example/new-board/2",
            ],
        )
        unchanged = prioritize_candidates_by_domain_deficit(
            candidates,
            {"known.example", "other.example", "third.example"},
            minimum_domains=3,
        )
        self.assertEqual(unchanged, candidates)

    def test_last_slots_are_reserved_for_required_new_domains(self) -> None:
        counts = Counter({"one.example": 3, "two.example": 2})
        self.assertTrue(
            should_reserve_for_domain_diversity(
                "one.example",
                counts,
                retained_source_unit_count=8,
                target=10,
                minimum_domains=4,
            )
        )
        self.assertFalse(
            should_reserve_for_domain_diversity(
                "fresh.example",
                counts,
                retained_source_unit_count=8,
                target=10,
                minimum_domains=4,
            )
        )
        self.assertFalse(
            should_reserve_for_domain_diversity(
                "one.example",
                counts,
                retained_source_unit_count=7,
                target=10,
                minimum_domains=4,
            )
        )

    def test_search_switches_provider_after_repeated_navigation_errors(self) -> None:
        class BrokenDriver:
            def __init__(self) -> None:
                self.calls = 0

            def get(self, _url: str) -> None:
                self.calls += 1
                raise WebDriverException("connection closed")

        driver = BrokenDriver()
        candidates = discover_candidates(
            driver,  # type: ignore[arg-type]
            [
                QuerySpec("group", "기타", f"검색어 {index}")
                for index in range(10)
            ],
            desired=20,
            pages=1,
            delay=0,
            providers_enabled=["bing"],
        )
        self.assertEqual(candidates, [])
        self.assertEqual(driver.calls, 3)

    def test_provider_staleness_does_not_skip_later_queries(self) -> None:
        class EmptyDriver:
            current_url = ""

            def __init__(self) -> None:
                self.urls: list[str] = []

            def get(self, url: str) -> None:
                self.current_url = url
                self.urls.append(url)

            def execute_script(self, script: str, *_args):
                if "document.body.innerText" in script:
                    return ""
                return []

        driver = EmptyDriver()
        with patch("collector.collect_candidates.time.sleep"):
            candidates = discover_candidates(
                driver,  # type: ignore[arg-type]
                [
                    QuerySpec("bank", "여권 및 통장", '"empty phrase one"'),
                    QuerySpec("bank", "여권 및 통장", '"empty phrase two"'),
                ],
                desired=20,
                pages=2,
                delay=0,
                providers_enabled=["yahoo_japan"],
                provider_stale_pages_limit=2,
            )

        self.assertEqual(candidates, [])
        self.assertEqual(len(driver.urls), 4)
        self.assertTrue(any("empty+phrase+one" in url for url in driver.urls))
        self.assertTrue(any("empty+phrase+two" in url for url in driver.urls))

    def test_nate_discovery_uses_rendered_organic_web_results(self) -> None:
        class NateDriver:
            current_url = ""

            def __init__(self) -> None:
                self.urls: list[str] = []

            def get(self, url: str) -> None:
                self.current_url = url
                self.urls.append(url)

            def execute_script(self, script: str, *_args):
                if "Boolean(document.querySelector" in script:
                    return True
                if "document.body.innerText" in script:
                    return ""
                return [
                    {
                        "href": "https://fresh.example/article/6/123/",
                        "text": "대출DB 판매",
                        "context": "대출DB 판매 텔레그램 문의",
                        "ignored": False,
                    }
                ]

        driver = NateDriver()
        with patch("collector.collect_candidates.time.sleep"):
            candidates = discover_candidates(
                driver,  # type: ignore[arg-type]
                [QuerySpec("db", "개인정보DB", "대출DB 판매")],
                desired=20,
                pages=1,
                delay=0,
                providers_enabled=["nate"],
                prefilter_mode="labeling",
            )

        self.assertEqual(driver.urls, [nate_search_url("대출DB 판매", 0)])
        self.assertEqual(len(candidates), 1)
        self.assertEqual(candidates[0].search_provider, "nate")
        self.assertEqual(candidates[0].url, "https://fresh.example/article/6/123/")

    def test_naver_blog_discovery_uses_each_result_card_context(self) -> None:
        class NaverBlogDriver:
            current_url = ""

            def __init__(self) -> None:
                self.context_selector = ""

            def get(self, url: str) -> None:
                self.current_url = url

            def execute_script(self, script: str, *args):
                if "document.body.innerText" in script:
                    return ""
                self.context_selector = str(args[1])
                return [
                    {
                        "href": "https://m.blog.naver.com/directseller/123",
                        "text": "네이버 계정 대량 판매",
                        "context": (
                            "네이버 비실명 계정 대량 판매합니다. "
                            "텔레그램 raw_handle 문의"
                        ),
                        "ignored": False,
                    },
                    {
                        "href": "https://m.blog.naver.com/phoneguide/456",
                        "text": "중고폰 초기화 안내",
                        "context": (
                            "중고 폰 팔 때 계정 로그아웃 후 공장 초기화하는 "
                            "순서를 안내합니다."
                        ),
                        "ignored": False,
                    },
                ]

        driver = NaverBlogDriver()
        with patch("collector.collect_candidates.time.sleep"):
            candidates = discover_candidates(
                driver,  # type: ignore[arg-type]
                [QuerySpec("account", "포털ID", "네이버 계정 대량 판매")],
                desired=20,
                pages=1,
                delay=0,
                providers_enabled=["naver_blog"],
                prefilter_mode="intent",
            )

        self.assertIn("[data-template-id='ugcItem']", driver.context_selector)
        self.assertEqual(len(candidates), 1)
        self.assertEqual(
            candidates[0].url,
            "https://m.blog.naver.com/directseller/123",
        )

    def test_yahoo_japan_discovery_uses_organic_card_context(self) -> None:
        class YahooDriver:
            current_url = ""

            def __init__(self) -> None:
                self.urls: list[str] = []

            def get(self, url: str) -> None:
                self.current_url = url
                self.urls.append(url)

            def execute_script(self, script: str, *_args):
                if "document.body.innerText" in script:
                    return ""
                return [
                    {
                        "href": "https://fresh-yahoo.example/post/7",
                        "text": "신분증 위조 제작",
                        "context": "신분증 위조 제작 텔레그램 seller 문의",
                        "ignored": False,
                    }
                ]

        driver = YahooDriver()
        with patch("collector.collect_candidates.time.sleep"):
            candidates = discover_candidates(
                driver,  # type: ignore[arg-type]
                [QuerySpec("identity", "기타", "신분증위조 텔레그램")],
                desired=20,
                pages=1,
                delay=0,
                providers_enabled=["yahoo_japan"],
                prefilter_mode="labeling",
            )

        self.assertEqual(
            driver.urls,
            [yahoo_japan_search_url("신분증위조 텔레그램", 0)],
        )
        self.assertEqual(len(candidates), 1)
        self.assertEqual(candidates[0].search_provider, "yahoo_japan")

    def test_social_accounts_are_distinct_but_posts_share_the_account_unit(self) -> None:
        first = source_unit_descriptor("https://t.me/s/channel_a/10")
        second = source_unit_descriptor("https://t.me/s/channel_a?before=20")
        other = source_unit_descriptor("https://t.me/s/channel_b/1")
        self.assertEqual(first, second)
        self.assertNotEqual(first, other)
        self.assertEqual(first[0], "social_account")

    def test_instagram_and_twitter_count_each_account_as_a_source(self) -> None:
        instagram_a = source_unit_descriptor("https://www.instagram.com/seller_a/")
        instagram_b = source_unit_descriptor("https://instagram.com/seller_b/")
        twitter_first = source_unit_descriptor(
            "https://x.com/seller_c/status/100"
        )
        twitter_second = source_unit_descriptor(
            "https://twitter.com/seller_c/status/999"
        )
        self.assertNotEqual(instagram_a, instagram_b)
        self.assertEqual(twitter_first, twitter_second)
        self.assertEqual(twitter_first[1], "x-twitter:seller_c")
        self.assertTrue(
            all(
                item[0] == "social_account"
                for item in (
                    instagram_a,
                    instagram_b,
                    twitter_first,
                    twitter_second,
                )
            )
        )

    def test_unresolved_instagram_posts_are_not_counted_as_distinct_accounts(self) -> None:
        first = source_unit_descriptor("https://instagram.com/p/POST_A/")
        second = source_unit_descriptor("https://instagram.com/reel/POST_B/")
        self.assertEqual(first, second)

    def test_board_posts_share_one_unit_but_separate_boards_do_not(self) -> None:
        first = source_unit_descriptor("https://creativebox.kr/igtrade/100")
        second = source_unit_descriptor("https://creativebox.kr/igtrade/999")
        other = source_unit_descriptor("https://creativebox.kr/ttmarket/100")
        self.assertEqual(first, second)
        self.assertNotEqual(first, other)
        generic_first = source_unit_descriptor(
            "https://forum.example/bbs/board.php?bo_table=free&wr_id=1"
        )
        generic_second = source_unit_descriptor(
            "https://forum.example/bbs/board.php?bo_table=free&wr_id=900"
        )
        self.assertEqual(generic_first, generic_second)
        cafe24_first = source_unit_descriptor(
            "https://shop.example/article/product-qa/6/225083/"
        )
        cafe24_second = source_unit_descriptor(
            "https://shop.example/article/product-qa/6/224969/"
        )
        self.assertEqual(cafe24_first, cafe24_second)

    def test_naver_and_daum_cafes_count_each_cafe_as_one_board(self) -> None:
        naver_first = source_unit_descriptor(
            "https://cafe.naver.com/ca-fe/cafes/12345/articles/10"
        )
        naver_second = source_unit_descriptor(
            "https://cafe.naver.com/ca-fe/cafes/12345/articles/999"
        )
        naver_other = source_unit_descriptor(
            "https://cafe.naver.com/ca-fe/cafes/67890/articles/10"
        )
        legacy = source_unit_descriptor(
            "https://cafe.naver.com/ArticleRead.nhn?clubid=12345&articleid=77"
        )
        daum_first = source_unit_descriptor(
            "https://cafe.daum.net/cafe_a/AbCd/1"
        )
        daum_second = source_unit_descriptor(
            "https://cafe.daum.net/cafe_a/AbCd/200"
        )
        daum_other = source_unit_descriptor(
            "https://cafe.daum.net/cafe_b/AbCd/1"
        )
        self.assertEqual(naver_first, naver_second)
        self.assertEqual(naver_first, legacy)
        self.assertNotEqual(naver_first, naver_other)
        self.assertEqual(daum_first, daum_second)
        self.assertNotEqual(daum_first, daum_other)

    def test_standalone_site_pages_count_as_one_source_unit(self) -> None:
        self.assertEqual(
            source_unit_descriptor("https://seller.example/service/a"),
            source_unit_descriptor("https://seller.example/service/b"),
        )

    def test_source_unit_token_is_hmac_and_does_not_reveal_account(self) -> None:
        descriptor = source_unit_descriptor("https://t.me/s/private_channel/1")
        token = source_unit_token(b"test-key", descriptor)
        self.assertTrue(token.startswith("social_account-hmac:"))
        self.assertNotIn("private_channel", token)

    def test_candidate_interleave_caps_repeated_board_posts(self) -> None:
        candidates = [
            Candidate(f"https://creativebox.kr/igtrade/{index}", "g", "기타")
            for index in range(5)
        ] + [
            Candidate("https://creativebox.kr/ttmarket/1", "g", "기타")
        ]
        ordered = interleave_candidates_by_domain(
            candidates,
            max_per_source_unit=2,
        )
        self.assertEqual(len(ordered), 3)
        self.assertEqual(
            [source_unit_descriptor(item.url) for item in ordered],
            [
                ("board", "creativebox:igtrade"),
                ("board", "creativebox:ttmarket"),
                ("board", "creativebox:igtrade"),
            ],
        )

    def test_known_board_and_social_account_candidates_are_excluded(self) -> None:
        known = {
            source_unit_descriptor("https://forum.example/bbs/board.php?bo_table=free&wr_id=1"),
            source_unit_descriptor("https://instagram.com/seller_a/p/POST1"),
        }
        candidates = [
            Candidate(
                "https://forum.example/bbs/board.php?bo_table=free&wr_id=999",
                "g",
                "기타",
            ),
            Candidate(
                "https://forum.example/bbs/board.php?bo_table=trade&wr_id=2",
                "g",
                "기타",
            ),
            Candidate("https://instagram.com/seller_a/p/POST2", "g", "기타"),
            Candidate("https://instagram.com/seller_b/p/POST3", "g", "기타"),
        ]
        kept = exclude_known_source_unit_candidates(candidates, known)
        self.assertEqual(
            [source_unit_descriptor(candidate.url) for candidate in kept],
            [
                ("board", "forum.example:bo_table=trade"),
                ("social_account", "instagram.com:seller_b"),
            ],
        )

    def test_collection_type_strata_are_mutually_exclusive(self) -> None:
        self.assertEqual(
            infer_collection_type("기타", "위조여권 제작", "판매 문의"),
            "신분증·여권 위조/제작",
        )
        self.assertEqual(
            infer_collection_type(
                "기타",
                "간호사면허증위조 위조제작전문",
                "24시 상담, 모든 작업 당일 진행",
            ),
            "신분증·여권 위조/제작",
        )
        self.assertEqual(
            infer_collection_type("기타", "법인통장 매입", "텔레그램 문의"),
            "통장·계좌",
        )
        self.assertEqual(
            infer_collection_type("기타", "대출DB 판매", "실시간 자료"),
            "개인정보DB",
        )
        self.assertEqual(
            infer_collection_type(
                "기타", "개인정보 판매", "연락처 명단을 제공합니다"
            ),
            "개인정보DB",
        )
        self.assertEqual(
            infer_collection_type("기타", "네이버 아이디 매입", "대량 문의"),
            "계정·아이디·가입인증",
        )
        self.assertEqual(
            infer_collection_type(
                "기타",
                "구글 아이디 판매",
                "Google 계정 구매 후 OTP 2FA 설정 방법과 판매 문의",
            ),
            "계정·아이디·가입인증",
        )

    def test_collection_type_ignores_settlement_and_form_documents(self) -> None:
        self.assertEqual(
            infer_collection_type(
                "기타",
                "2330명 유튜브계정 7만원에 팝니다",
                "토스계좌로 돈을 받은 후 브랜드계정을 제공할 예정입니다.",
            ),
            "계정·아이디·가입인증",
        )
        self.assertEqual(
            infer_collection_type(
                "기타",
                "네이버 계정 대여",
                "계정을 매입합니다. 양식은 신분증 앞면과 계좌번호입니다.",
            ),
            "계정·아이디·가입인증",
        )
        self.assertEqual(
            infer_collection_type(
                "기타",
                "강남 VVIP 디비 판매",
                "주민등록증 인증을 거친 고객 DB 5000개를 판매합니다.",
            ),
            "개인정보DB",
        )
        self.assertEqual(
            infer_collection_type(
                "기타",
                "게임 매입방",
                "KYC 인증 삽니다. 진행 양식에는 신분증 하나가 필요합니다.",
            ),
            "계정·아이디·가입인증",
        )

    def test_intent_gate_rejects_platform_policy_and_bank_product_pages(self) -> None:
        self.assertEqual(
            relevance_gate_reason(
                "법적 고지 - Apple 미디어 서비스 이용 약관",
                "계정과 콘텐츠를 구매하거나 판매하는 거래에 관한 서비스 약관",
                "https://www.apple.com/kr/legal/terms.html",
                "unknown",
                "intent",
            ),
            "excluded_domain",
        )
        self.assertEqual(
            relevance_gate_reason(
                "계정 탈취, 거래, 양도, 교환 등",
                "운영정책상 계정 판매와 구매 행위는 허용되지 않습니다.",
                "https://talksafety.kakao.com/policy/account",
                "unknown",
                "intent",
            ),
            "excluded_domain",
        )
        self.assertEqual(
            relevance_gate_reason(
                "자유입출금 예금 상품",
                "사업자 간 금 거래를 위한 정상 은행 계좌를 제공합니다.",
                "https://www.kebhana.com/product/1479724",
                "unknown",
                "intent",
            ),
            "excluded_domain",
        )

    def test_intent_gate_rejects_warning_post_titled_as_related_tips(self) -> None:
        self.assertEqual(
            relevance_gate_reason(
                "청소년보호법 위반 관련 팁",
                "텔레그램에서 위조신분증을 제작해 준다니 조심하세요.",
                "https://forum.example/post/1",
                "unknown",
                "intent",
            ),
            "excluded_document_type",
        )

    def test_labeling_workbook_has_links_dropdown_and_notes(self) -> None:
        row = {name: "" for name in SCHEMA}
        row.update(
            {
                "sample_id": "LP-000001",
                "registrable_domain": "example.com",
                "masked_title": "고객 DB 판매",
                "masked_text": "판매 의사 확인용 본문",
                "page_type": "unknown",
                "live_status": "accessible",
            }
        )
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "label.xlsx"
            write_labeling_workbook(
                path,
                [row],
                {"LP-000001": "https://example.com/public/post/1"},
            )
            workbook = load_workbook(path)
            sheet = workbook["라벨링"]
            self.assertEqual(sheet["C2"].hyperlink.target, "https://example.com/public/post/1")
            self.assertEqual(sheet["F1"].value, "판정")
            self.assertEqual(sheet["G1"].value, "메모")
            self.assertEqual(len(sheet.data_validations.dataValidation), 1)
            self.assertIn("안내", workbook.sheetnames)
            self.assertNotIn("본문 전체", workbook.sheetnames)

    def test_labeling_workbook_supports_an_empty_result(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "label.xlsx"
            write_labeling_workbook(path, [], {})
            workbook = load_workbook(path)
            sheet = workbook["라벨링"]
            self.assertEqual(sheet.max_row, 1)
            self.assertEqual(len(sheet.data_validations.dataValidation), 0)
            self.assertIn("안내", workbook.sheetnames)

    def test_collector_writes_shareable_labeling_workbook(self) -> None:
        row = {
            "sample_id": "EG-000001",
            "collected_at": "2026-08-28T12:00:00+09:00",
            "source_url": "https://example.com/public/post/2",
            "registrable_domain": "example.com",
            "title": "고객 DB 판매",
            "text": "텔레그램 raw_handle 문의",
        }
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            csv_path = root / "data.csv"
            with csv_path.open("w", encoding="utf-8-sig", newline="") as handle:
                writer = csv.DictWriter(handle, fieldnames=RESTRICTED_REVIEW_SCHEMA)
                writer.writeheader()
                writer.writerow(row)
            workbook_path = root / "label.xlsx"
            count = write_collector_labeling_workbook(
                csv_path, workbook_path
            )
            workbook = load_workbook(workbook_path)
            sheet = workbook["라벨링"]
            self.assertEqual(count, 1)
            self.assertEqual(
                sheet["C2"].hyperlink.target,
                "https://example.com/public/post/2",
            )
            self.assertEqual(workbook_path.stat().st_mode & 0o777, 0o644)

    def test_restricted_labeling_sheet_includes_source_url(self) -> None:
        row = {name: "" for name in SCHEMA}
        row.update(
            {
                "sample_id": "LP-000001",
                "masked_title": "고객 DB 판매",
                "masked_text": "판매 의사 확인용 본문",
                "page_type": "unknown",
                "live_status": "accessible",
            }
        )
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "labeling_with_urls.csv"
            write_labeling_sheet(
                path,
                [row],
                {"LP-000001": "https://example.com/public/post/1"},
            )
            with path.open(encoding="utf-8-sig", newline="") as handle:
                written = list(csv.DictReader(handle))
        self.assertEqual(
            written[0]["source_url"],
            "https://example.com/public/post/1",
        )
        self.assertEqual(written[0]["final_label"], "")

    def test_intent_priority_does_not_require_contact_details(self) -> None:
        self.assertEqual(intent_bucket(""), "intent_priority")
        self.assertEqual(
            intent_bucket("missing_concrete_contact"), "intent_priority"
        )
        self.assertEqual(
            intent_bucket("missing_body_offer"), "boundary_review"
        )
        self.assertEqual(
            intent_bucket("excluded_reporting_context"), "hard_negative"
        )

    def test_labeling_pilot_prioritizes_strict_and_balances_reasons(self) -> None:
        def item(index: int, reason: str) -> SourceRow:
            return SourceRow(
                row={"registrable_domain": f"d{index}.example"},
                candidate=Candidate(f"https://d{index}.example/post", "g", "기타"),
                source=Path("source"),
                source_tier="primary",
                gate_reason=reason,
                selection_bucket=strict_bucket(reason),
            )

        rows = [
            item(1, "excluded_reporting_context"),
            item(2, ""),
            item(3, "excluded_reporting_context"),
            item(4, "excluded_page_type"),
            item(5, "missing_concrete_contact"),
        ]
        ordered = prioritize_rows(rows, priority_enabled=True)
        self.assertEqual(
            [row.gate_reason for row in ordered],
            [
                "",
                "missing_concrete_contact",
                "excluded_reporting_context",
                "excluded_page_type",
                "excluded_reporting_context",
            ],
        )

    def test_labeling_pilot_honors_domain_cap(self) -> None:
        rows = [
            SourceRow(
                row={"registrable_domain": "same.example"},
                candidate=Candidate(f"https://same.example/{index}", "g", "기타"),
                source=Path("source"),
                source_tier="primary",
            )
            for index in range(3)
        ]
        rows.append(
            SourceRow(
                row={"registrable_domain": "other.example"},
                candidate=Candidate("https://other.example/1", "g", "기타"),
                source=Path("source"),
                source_tier="primary",
            )
        )
        selected = select_rows(rows, target=2, max_per_domain=1)
        self.assertEqual(
            [row.row["registrable_domain"] for row in selected],
            ["same.example", "other.example"],
        )

    def test_labeling_pilot_assigns_duplicate_clusters(self) -> None:
        rows = [
            {"masked_title": "계정 판매", "masked_text": "대량 계정 판매 문의"},
            {"masked_title": "계정 판매", "masked_text": "대량 계정 판매 문의"},
            {"masked_title": "다른 글", "masked_text": "전혀 다른 정상 문맥"},
        ]
        assign_near_duplicate_clusters(rows)
        self.assertEqual(
            rows[0]["near_duplicate_cluster"],
            rows[1]["near_duplicate_cluster"],
        )
        self.assertNotEqual(
            rows[0]["near_duplicate_cluster"],
            rows[2]["near_duplicate_cluster"],
        )

    def test_bing_redirect_is_unwrapped_without_request(self) -> None:
        target = "https://example.com/public/post/1"
        encoded = base64.urlsafe_b64encode(target.encode()).decode().rstrip("=")
        wrapped = f"https://www.bing.com/ck/a?u=a1{encoded}&ntb=1"
        self.assertEqual(unwrap_search_result_url(wrapped), target)

    def test_duckduckgo_redirect_is_unwrapped_without_request(self) -> None:
        target = "https://example.com/public/post/2"
        wrapped = f"https://duckduckgo.com/l/?uddg={target}"
        self.assertEqual(unwrap_search_result_url(wrapped), target)

    def test_contact_data_is_masked(self) -> None:
        text = (
            "문의 test@example.com, 010-1234-5678, 텔레그램 sample_id "
            "홈페이지https://example.com/path"
        )
        masked = mask_text(text)
        self.assertIn("[EMAIL]", masked)
        self.assertIn("[PHONE]", masked)
        self.assertIn("[MESSENGER_ID]", masked)
        self.assertIn("홈페이지[CONTACT_URL]", masked)
        self.assertNotIn("sample_id", masked)
        self.assertNotIn("https://", masked)

    def test_restricted_review_keeps_messenger_id_only(self) -> None:
        raw = (
            "문의 텔레그램 raw_handle, https://t.me/raw_channel, "
            "카카오 아이디: kakao_raw, ㅌㄹ short_raw, "
            "test@example.com, 010-1234-5678"
        )
        restricted = mask_text(raw, preserve_messenger_ids=True)
        self.assertIn("raw_handle", restricted)
        self.assertIn("https://t.me/raw_channel", restricted)
        self.assertIn("kakao_raw", restricted)
        self.assertIn("short_raw", restricted)
        self.assertIn("[EMAIL]", restricted)
        self.assertIn("[PHONE]", restricted)
        self.assertNotIn("test@example.com", restricted)
        self.assertNotIn("010-1234-5678", restricted)

    def test_spaced_fullwidth_phone_is_masked(self) -> None:
        raw = "유선 문의 0１0 - 9 ６４9 -７１４７ 상담 가능합니다"
        masked = mask_text(raw, preserve_messenger_ids=True)
        self.assertIn("[PHONE]", masked)
        self.assertNotIn("７１４７", masked)
        self.assertNotIn("9649", masked)

    def test_obfuscated_mobile_phone_is_masked_and_groups_campaign(self) -> None:
        standard = contact_campaign_id(
            b"test-key",
            campaign_fingerprint_text("문의", "010-3435-5177"),
        )
        for raw in (
            "ⓞ①ⓞ - ③④③⑤ - ⑤①⑦⑦",
            "0.1.0-3.4.3.5-5.1.7.7",
            "0+1+0-3+4+3+5-5+1+7+7",
            "o.1.o - 3.4.3.5 - 5.I.7.7",
        ):
            with self.subTest(raw=raw):
                masked = mask_text(f"문의 {raw}", preserve_messenger_ids=True)
                self.assertIn("[PHONE]", masked)
                self.assertNotIn("3435", masked)
                self.assertEqual(
                    standard,
                    contact_campaign_id(
                        b"test-key",
                        campaign_fingerprint_text("문의", raw),
                    ),
                )
        self.assertIn("[PHONE]", mask_text("문의 ⓞ①ⓞ-③④③⑤-⑤①⑦"))
        long_campaign = contact_campaign_id(
            b"test-key",
            campaign_fingerprint_text(
                "통장 판매",
                ("반복 판매 문구 " * 1_000) + " 문의 010-3435-5177",
            ),
        )
        self.assertEqual(standard, long_campaign)

    def test_spaced_fullwidth_email_is_masked(self) -> None:
        raw = "메일 문의 ｆｏｒｍ88@gmail. Ｃom 또는 normal@example.com"
        masked = mask_text(raw, preserve_messenger_ids=True)
        self.assertEqual(masked.count("[EMAIL]"), 2)
        self.assertNotIn("gmail", masked)
        self.assertNotIn("example.com", masked)

    def test_telegram_shorthand_contact_is_masked(self) -> None:
        masked = mask_text("디비 텔그 sample_id 문의")
        self.assertIn("텔그 [MESSENGER_ID]", masked)
        self.assertNotIn("sample_id", masked)

    def test_messenger_contact_url_keeps_channel_but_masks_handle(self) -> None:
        masked = mask_text(
            "상담 https://t.me/private_handle 또는 https://open.kakao.com/o/secret"
        )
        self.assertIn("텔레그램 [MESSENGER_ID]", masked)
        self.assertIn("카카오톡 [MESSENGER_ID]", masked)
        self.assertNotIn("private_handle", masked)
        self.assertNotIn("/secret", masked)

    def test_spreadsheet_formula_prefix_is_neutralized(self) -> None:
        self.assertEqual(mask_text('=HYPERLINK("bad")'), '\'=HYPERLINK("bad")')
        self.assertEqual(safe_spreadsheet_text("+cmd"), "'+cmd")

    def test_tracking_parameters_and_fragment_are_removed(self) -> None:
        url = canonicalize_url("https://Example.com/post?id=7&utm_source=test#part")
        self.assertEqual(url, "https://example.com/post?id=7")

    def test_public_naver_blog_frame_has_mobile_fallback(self) -> None:
        self.assertEqual(
            public_content_fallback_url(
                "https://blog.naver.com/public_writer/223456789012"
            ),
            "https://m.blog.naver.com/public_writer/223456789012",
        )
        self.assertIsNone(
            public_content_fallback_url("https://example.com/public/post/1")
        )

    def test_contact_campaign_uses_hmac_without_plaintext(self) -> None:
        campaign = contact_campaign_id(
            b"test-key", "문의 test@example.com 010-1234-5678"
        )
        self.assertTrue(campaign.startswith("contact-hmac:"))
        self.assertNotIn("example.com", campaign)
        self.assertEqual(
            campaign,
            contact_campaign_id(b"test-key", "010-1234-5678 / test@example.com"),
        )

    def test_direct_contact_campaign_ignores_varying_page_urls(self) -> None:
        first = contact_campaign_id(
            b"test-key",
            "텔레그램 @same_seller https://example.com/post/1",
        )
        second = contact_campaign_id(
            b"test-key",
            "텔레그램 @same_seller https://other.example/post/9",
        )
        self.assertEqual(first, second)

    def test_english_kakao_id_creates_contact_campaign(self) -> None:
        campaign = contact_campaign_id(
            b"test-key",
            campaign_fingerprint_text(
                "최적화 블로그 임대 및 매매 문의",
                "KAKAO ID: wlgks3787",
            ),
        )
        self.assertTrue(campaign.startswith("contact-hmac:"))

    def test_campaign_uses_shared_primary_handle_despite_secondary_contacts(self) -> None:
        first = contact_campaign_id(
            b"test-key",
            "카톡 cdc000011 텔레 https://t.me/ydc2019 010-1234-5678",
        )
        second = contact_campaign_id(
            b"test-key",
            "텔레그램 cdc000011 인스타 계정 4천개 판매",
        )
        self.assertEqual(first, second)

    def test_repeated_operator_brand_groups_storefronts_with_missing_contact(self) -> None:
        without_contact = contact_campaign_id(
            b"test-key",
            "아이디피플에서는 계정을 판매합니다. 아이디피플은 전문 업체이며 "
            "아이디피플 서비스를 메신저로 문의하세요.",
        )
        with_contact = contact_campaign_id(
            b"test-key",
            "아이디피플에서는 계정을 판매합니다. 아이디피플은 전문 업체이며 "
            "아이디피플 텔레그램 @seller_handle 문의",
        )
        self.assertTrue(without_contact)
        self.assertEqual(without_contact, with_contact)

    def test_operator_brand_in_professional_supplier_copy_groups_campaign(self) -> None:
        storefront = contact_campaign_id(
            b"test-key",
            "아이디피플에서는 계정을 판매합니다. 아이디피플은 전문 업체이며 "
            "아이디피플 서비스를 제공합니다.",
        )
        copied_post = contact_campaign_id(
            b"test-key",
            "국내 010 인증 N사 계정 전문 공급처 아이디피플입니다. "
            "아이디피플은 수작업 계정만 공급하며 아이디피플 재구매율이 높습니다. "
            "텔레그램 @different_handle",
        )
        self.assertEqual(storefront, copied_post)

    def test_google_id_mc_template_groups_renamed_storefront(self) -> None:
        old_site = contact_campaign_id(
            b"test-key",
            "정상·실사용 Google 계정 대량 보유. "
            "결제부터 계정 제공까지 최소 대기. "
            "개인정보 안전 보장 및 익명 거래 지원.",
        )
        renamed_site = contact_campaign_id(
            b"test-key",
            "Google ID MC 신규 사이트입니다. 개인정보 안전 보장 및 익명 거래 지원. "
            "정상·실사용 Google 계정 대량 보유. 결제부터 계정 제공까지 최소 대기.",
        )
        self.assertTrue(old_site)
        self.assertEqual(old_site, renamed_site)
        product_page = contact_campaign_id(
            b"test-key",
            "업무용 구글 계정 - Business Google Account | Google ID MC",
        )
        self.assertEqual(old_site, product_page)

    def test_obfuscated_messenger_label_keeps_same_campaign(self) -> None:
        first = contact_campaign_id(
            b"test-key",
            "문의 텔레•Run55 계정 판매",
        )
        second = contact_campaign_id(
            b"test-key",
            "문의텔ﾩ RUN55 아이디 임대",
        )
        self.assertEqual(first, second)

    def test_halfwidth_messenger_label_is_masked(self) -> None:
        masked = mask_text("문의텔ﾩ RUN55 네이버 아이디 판매")
        self.assertIn("텔ﾩ [MESSENGER_ID]", masked)
        self.assertNotIn("RUN55", masked)

    def test_misspelled_messenger_label_masks_and_groups_campaign(self) -> None:
        standard = contact_campaign_id(
            b"test-key",
            "개인통장 매입 문의 텔레그램 @mbzang",
        )
        shortened = contact_campaign_id(
            b"test-key",
            "개인장 판매 문의 탤@mbzang",
        )
        misspelled = contact_campaign_id(
            b"test-key",
            "법인통장 매입 문의 탤레@mbzang",
        )
        self.assertEqual(standard, shortened)
        self.assertEqual(standard, misspelled)
        for label in ("탤@mbzang", "탤레@mbzang", "탤레그램@mbzang"):
            masked = mask_text(label)
            self.assertNotIn("mbzang", masked)
            self.assertIn("[MESSENGER_ID]", masked)

    def test_hangul_adjacent_handle_is_masked_and_grouped(self) -> None:
        standard = contact_campaign_id(
            b"test-key",
            "통장 매입 문의 텔레그램 @same_handle",
        )
        for contact in ("문의@same_handle", "텔@same_handle", "텔/레@same_handle"):
            masked = mask_text(contact)
            self.assertNotIn("same_handle", masked)
            self.assertRegex(masked, r"\[(?:ACCOUNT|MESSENGER_ID)\]")
            self.assertEqual(
                standard,
                contact_campaign_id(b"test-key", contact),
            )

    def test_contact_campaign_is_stable_after_review_masking(self) -> None:
        title = "문의 010-1234-5678 텔레그램 @same_seller"
        text = "고객 DB 판매 test@example.com"
        first = campaign_fingerprint_text(title, text)
        second = campaign_fingerprint_text(
            mask_text(title, preserve_messenger_ids=True),
            mask_text(text, preserve_messenger_ids=True),
        )
        self.assertEqual(
            contact_campaign_id(b"test-key", first),
            contact_campaign_id(b"test-key", second),
        )
        self.assertIn("same_seller", first)

    def test_tree_db_brand_is_not_a_personal_database_target(self) -> None:
        reason = relevance_gate_reason(
            "왕대추나무 100주 삽니다 : 트리디비",
            "조경수는 수종과 규격으로 검색합니다. 왕대추나무를 삽니다.",
            "https://treedb.example/tree/100",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "missing_relevant_target")

    def test_db_engineering_machinery_listing_is_not_personal_data(self) -> None:
        reason = relevance_gate_reason(
            "판매 중인 중고 DB Engineering 폐기물 재활용 굴착기 광고",
            (
                "중고 건설장비와 이동식 선별기, 스크리너, 트롬멜을 판매합니다. "
                "DB Engineering DB-40 Siebanlage 제작년도와 장비 가격을 "
                "확인하고 판매자에게 장비 구매를 문의할 수 있습니다."
            ),
            "https://machinery.example/db_engineering",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_db_brand_or_stock")

        illicit = relevance_gate_reason(
            "DB 엔지니어링 고객정보 판매",
            (
                "관리자 서버에서 추출한 이름과 전화번호 고객DB를 판매합니다. "
                "최신 해킹 DB 단가와 샘플은 텔레그램 seller123 문의 바랍니다."
            ),
            "https://seller.example/hacked-db",
            "unknown",
            "intent",
        )
        self.assertEqual(illicit, "")

    def test_niche_personal_database_offer_is_kept_and_stratified(self) -> None:
        title = "병원DB 판매 텔레그램 문의"
        text = "성형 상담 고객 병원디비를 판매합니다. 대량 공급 가능합니다."
        self.assertEqual(
            relevance_gate_reason(
                title,
                text,
                "https://seller.example/db/1",
                "unknown",
                "intent",
            ),
            "",
        )
        self.assertEqual(
            infer_collection_type("기타", title, text),
            "개인정보DB",
        )

    def test_art_team_identity_card_request_is_not_illicit_forgery(self) -> None:
        reason = relevance_gate_reason(
            "미술팀 한그림 - 신분증 제작의뢰합니다",
            "PROJECT WORKS 제작의뢰 Professional Art work and Design service",
            "https://art.example/project/1",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_legal_prop_or_security_guide")

    def test_account_purchase_warning_is_not_a_buying_inquiry(self) -> None:
        reason = relevance_gate_reason(
            "네이버카페 계정 구매하시려는 사람들은 꼭 보세요",
            "신규계정은 절대 구매하지 마세요. 판매자가 AS도 해주지 않습니다.",
            "https://forum.example/post/1",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_trade_warning")

    def test_blog_sale_do_not_warning_is_not_an_offer(self) -> None:
        reason = relevance_gate_reason(
            "네이버 블로그 판매/양도/대여/매매 하면 안되는 이유",
            (
                "요약: 네이버 계정 판매는 약관 위반이며 위험하다. "
                "블로그 판매 안합니다. 댓글, 메일, 문자, 카톡 하지마세요. "
                "계정을 다른 사람에게 판매 또는 양도할 수 없습니다. "
                "결국 골치아파집니다. 하지마세요."
            ),
            "https://m.blog.naver.com/example/1",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_trade_warning")

    def test_do_not_hesitate_copy_does_not_negate_account_offer(self) -> None:
        reason = relevance_gate_reason(
            "네이버 블로그 계정 판매",
            (
                "네이버맘카페계정판매 이용을 망설이고 계신다면 너무 "
                "고민하지 마세요. 네이버 계정 판매합니다. 아이디 다량 "
                "보유, 즉시 거래 가능. 카톡 seller123 문의 바랍니다."
            ),
            "https://seller.example/account-offer",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "")

    def test_account_sale_howto_question_is_not_a_buying_inquiry(self) -> None:
        reason = relevance_gate_reason(
            "구글 계정 판매시 문의사항입니다",
            (
                "판매하고자 하는 구글 계정에 개인정보가 많이 연결되어 있는데 "
                "구매자가 복구 이메일과 전화번호를 수정하면 제 아이디는 그대로 "
                "보존되는 건가요? 거래 방식이 명확하지 않아 문의드립니다."
            ),
            "https://community.example/account-question",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_question_or_guide")

    def test_past_account_sale_legal_dispute_is_not_a_current_offer(self) -> None:
        reason = relevance_gate_reason(
            "네이버 비실명 계정 판매했는데 실명인증 해달라고 연락와요",
            (
                "지인이 2015년경 비실명 계정을 팔았습니다. 판매했던 업체가 "
                "보호조치 해제를 위해 실명인증을 해달라고 요구합니다. "
                "개인정보를 사고판 것이 법적으로 문제가 되는 것을 알고 "
                "있었는데, 협박죄로 고소하고 싶습니다. 제가 당할 불이익이 "
                "있을까요?"
            ),
            "https://kin.example/question/past-account-sale",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_question_or_guide")

    def test_casual_username_scam_story_is_a_trade_warning(self) -> None:
        reason = relevance_gate_reason(
            "인스타 username 판매, 구매 사기?",
            (
                "친구한테 사용자 이름을 사겠다는 디엠이 왔대요. "
                "돈을 보냈다고 주장했지만 입금 내역은 없었고, "
                "이런 연락을 받으면 조심하라고 알려드려요. "
                "인스타 계정 판매는 불법입니다. 여러분 하시면 안돼요."
            ),
            "https://blog.example/username-scam",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_trade_warning")

    def test_informal_blog_rental_victim_finding_post_is_excluded(self) -> None:
        reason = relevance_gate_reason(
            "최블 임대 피해보신 분 계실지요?",
            (
                "최블 임대 1년 600 계약했는데 입금받고 한 시간 뒤 환불한다고 "
                "했습니다. 나눠서 준다더니 계속 시간을 끌고 있습니다. "
                "비슷한 사례가 있으신 분은 연락주십시오."
            ),
            "https://community.example/victim-post",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_trade_warning")

    def test_account_rental_offer_can_address_prior_victims(self) -> None:
        reason = relevance_gate_reason(
            "최적화 블로그 임대합니다",
            (
                "기존 업체에서 피해보신 분도 연락주세요. 최블을 대량 보유하고 "
                "있으며 6개월 또는 1년 단위로 임대합니다. 계약 즉시 계정을 "
                "공급하고 A/S를 보장합니다. 카톡 seller123 문의 바랍니다."
            ),
            "https://seller.example/blog-rental",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "")

    def test_legal_decision_about_borrowed_account_is_not_trade_intent(self) -> None:
        reason = relevance_gate_reason(
            "차명계좌 대여에 따른 사례금으로 보아 과세한 처분의 당부",
            "조세심판원 주문 심판청구를 기각한다. 처분개요 청구인은 "
            "쟁점금액에 대한 청구주장을 제기했고 처분청은 경정·고지했다.",
            "https://law.example/case/100",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_legal_decision")

    def test_repeated_ssr_keyword_block_is_not_a_direct_post(self) -> None:
        block = "법인장파는곳 법인통장구매 법인통장판매 법인통장매입"
        reason = relevance_gate_reason(
            "상품 Q&A",
            f'"{block}" "{block}" "{block}" 기록을 정리한 일반 안내입니다.',
            "https://shop.example/article/qa/6/100",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_search_spam")

    def test_illicit_keyword_advertising_service_is_not_the_trade_post(self) -> None:
        text = " ".join(
            [
                "법인통장삽니다홍보 상단 노출 광고 도배 문의 폭주 유도 광고주",
                "법인장구매광고 상단 유지 키워드 광고 홍보",
            ]
            * 8
        )
        reason = relevance_gate_reason(
            "마케팅 홍보 상품 Q&A",
            text,
            "https://agency.example/article/qa/6/200",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_search_spam")

    def test_db_promotion_vendor_copy_is_not_a_db_sales_post(self) -> None:
        reason = relevance_gate_reason(
            "농산물 쇼핑몰",
            "DB 디비 관련 홍보, 광고는 트리플마케팅 이용해보세요. "
            "대출DB판매 광고 캠페인과 상단 노출을 대행합니다.",
            "https://farm.example/article/data/7/1",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_search_spam")

    def test_ssr_db_master_guide_is_not_a_direct_sales_post(self) -> None:
        text = (
            "Google 기준 상위 노출 SEO 기준 DB판매 마스터 가이드 "
            "검색 키워드들을 구조적으로 포함 키워드 통합 정리 "
            + "대출DB 판매 코인DB 구매 병원DB 매입 디비 판매 " * 8
        )
        reason = relevance_gate_reason(
            "사업자DB 완전 가이드",
            text,
            "https://shop.example/article/free/5/300",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_search_spam")

    def test_attributable_forged_document_catalog_is_not_search_spam(self) -> None:
        reason = relevance_gate_reason(
            "신분증위조판매 텔레 sos3579 신분증위조가격",
            (
                "여권위조판매 텔레✍sos3579 주문 문의 가능합니다.\n"
                "주민등록증위조제작 텔레✍sos3579 가격 상담합니다.\n"
                "운전면허증위조판매 텔레✍sos3579 제작 가능합니다.\n"
                "졸업증명서위조제작 텔레✍sos3579 판매 문의. "
                + "여권위조판매 주민등록증위조판매 운전면허증위조판매 " * 20
            ),
            "https://forum.example/topic/forged-id-sale",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "")

    def test_in_page_heading_can_attribute_forgery_catalog(self) -> None:
        reason = relevance_gate_reason(
            "의류 쇼핑몰 상품 리뷰",
            (
                "토익성적표위조 카톡 job365 광고 삭제 후에도 문의 가능. "
                "각종서류위조 모두 가능합니다. 카톡 job365 문의.\n"
                "여권위조판매 카톡 job365 제작 문의 가능합니다.\n"
                "주민등록증위조제작 카톡 job365 가격 상담합니다.\n"
                "운전면허증위조판매 카톡 job365 주문 가능합니다.\n"
                + "졸업증명서위조 판매 성적증명서위조 제작 문의 " * 30
            ),
            "https://shop.example/mobile/board/review/200",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "")

    def test_forgery_storefront_survives_generic_service_filler(self) -> None:
        reason = relevance_gate_reason(
            "간호사면허증위조 텔레 [ACCOUNT] 위조제작전문 신용보장업체",
            (
                "간호사면허증위조 텔레 [ACCOUNT] 간호사면허증위조제작. "
                "24시 보안상담, 신용보장업체입니다. 텔레 [ACCOUNT] 문의. "
                "모든 작업 당일 진행하며 작업 전 디테일한 상담을 합니다. "
                "보안을 최우선으로 진행하고 안전은 믿고 맡겨주세요. "
                "텔레 [ACCOUNT] 가격 문의. "
                "예약 방법 원하시는 코스 심야 시간대 평일 주말 모두 영업 "
                "재방문 의사 이용 후기 편하게 즐길 분위기 날짜와 인원."
            ),
            "https://clinic.example/board/forgery-seller",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "")

    def test_contact_attributed_forgery_catalog_list_is_not_keyword_dump(self) -> None:
        document_lines = [
            "운전면허증위조",
            "졸업증명서제작",
            "성적증명서위조",
            "가족관계증명서제작",
            "주민등록등본위조",
            "재직증명서위조",
            "통장거래내역서제작",
            "병원진단서위조",
            "국가기술자격증제작",
            "신분증위조",
        ] * 6
        reason = relevance_gate_reason(
            "대학교졸업증명서위조 카톡 [ACCOUNT] 텔레 [ACCOUNT]",
            "\n".join(document_lines),
            "https://board.example/free/forged-catalog",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "")

    def test_forgery_keywords_in_seo_agency_copy_remain_search_spam(self) -> None:
        reason = relevance_gate_reason(
            "신분증위조판매 키워드 상위 노출",
            (
                "여권위조판매 주민등록증위조제작 운전면허증위조판매 "
                "검색 키워드 광고 대행사 실행사 광고주 광고비 상위 노출 "
                "최상단 고정 구글 검색 1페이지 영역 독점 무한 도배 "
                "문의 폭주 유도 서비스를 제공합니다."
            ),
            "https://agency.example/seo/forgery-keywords",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_search_spam")

    def test_forgery_phrase_in_unrelated_peptide_forum_ad_is_mismatch(self) -> None:
        reason = relevance_gate_reason(
            "CCA-010: WhatsApp contact | wholesale peptides usa",
            (
                "WhatsApp contact | Telegram vendor | wholesale peptides usa\n"
                "WhatsApp contact | Telegram vendor | buy hcg peptide usa\n"
                "신분증위조업체 신분증위조제작 회계사자격증위조 "
                "텔레 [ACCOUNT]\nDisplay mode\n"
                "Display replies flat, with oldest first\n"
                "Display replies in threaded form\nPermalink\nReply\n"
                "WhatsApp contact | wholesale peptides usa\n"
                "WhatsApp contact | buy hcg peptide usa\n"
                "신분증위조업체 신분증위조제작 회계사자격증위조 "
                "텔레 [ACCOUNT]"
            ),
            "https://lms.example/mod/forum/discuss.php?d=10",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_title_body_mismatch")

    def test_unsolicited_db_sales_complaint_is_not_purchase_intent(self) -> None:
        reason = relevance_gate_reason(
            "사업자 DB 판매한다는 연락 다들 받으시나요?",
            "광고 전화가 폭주합니다. 개인정보보호법 위반 아닌가요? "
            "이 전화 어떻게 차단하나요?",
            "https://community.example/post/1",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_reporting_context")

        report_call = relevance_gate_reason(
            "소프트웨어 가품 판매 및 당근 다계정 판매업자",
            (
                "당근마켓 다중계정을 동원한 정책위반 유저입니다. 불법 "
                "복제 소프트웨어를 판매하는 전문판매업자이니 보시면 같이 "
                "신고 해주세요. 당근 신고 바랍니다."
            ),
            "https://report.example/post/account-seller",
            "unknown",
            "intent",
        )
        self.assertEqual(report_call, "excluded_reporting_context")

    def test_fanclub_ticket_page_is_not_identity_document_trade(self) -> None:
        reason = relevance_gate_reason(
            "팬클럽 추첨제 예매 안내",
            "공연 정보 관객 입장 17시 공연 시작 18시 티켓 금액 지정석 "
            "예매 시 본인 확인을 위해 여권을 지참해 주세요.",
            "https://fan.example/event/1",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_event_ticket_context")

    def test_personal_forged_id_ticket_check_question_is_not_trade(self) -> None:
        reason = relevance_gate_reason(
            "미성년자 신분증 위조",
            (
                "미성년자가 신분증 위조하면 이거 걸리냐? 티켓에 있는 "
                "이름이랑 신분증 이름을 대조해서 검사할까, 아니면 성인인 "
                "것만 보고 넘어가냐? 진지하게 궁금하다."
            ),
            "https://community.example/post/id-check-question",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_question_or_guide")

    def test_direct_forged_id_offer_near_event_language_is_still_kept(self) -> None:
        reason = relevance_gate_reason(
            "신분증 위조 제작합니다",
            (
                "성인 확인과 입장 검사에 사용할 신분증을 위조 제작합니다. "
                "여권과 주민등록증도 제작 가능하며 텔레그램 seller123 문의."
            ),
            "https://seller.example/forged-id",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "")

    def test_betting_referral_feed_is_not_identity_or_account_trade(self) -> None:
        reason = relevance_gate_reason(
            "종토방 꽁머니 공지",
            (
                "토사 가입시키고 입금미션 후 환 내줍니다. "
                "거래소 매입중이니 시드 마련하세요. "
                "채널에 입장하면 상담원이 가입 퀴스트와 꽁머니 받는 법을 "
                "자세히 알려주고 무료 환전까지 도와줍니다. "
                "어디서 봤냐고 물어보면 인스타그램에서 봤다고 대답하세요. "
                "40대만 진행하고 신분증은 봅니다."
            ),
            "https://t.me/s/referral_feed",
            "public_messenger_page",
            "intent",
        )
        self.assertEqual(reason, "excluded_gambling_referral")

    def test_repeated_no_db_purchase_copy_is_not_a_buying_post(self) -> None:
        reason = relevance_gate_reason(
            "보험DB 구매 없이 고객을 만나는 방법",
            "보험DB를 구매하지 않습니다. 광고나 보험DB 구매를 하지 않고 "
            "고객이 먼저 찾아오도록 온라인 영업을 합니다. 상담은 연락주세요.",
            "https://blog.example/post/1",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_db_purchase_alternative")

    def test_db_numbered_scale_model_is_not_personal_database_trade(self) -> None:
        reason = relevance_gate_reason(
            "카스 벤치형저울 DB-1 구매문의",
            "DB-1 저울은 목욕탕용 제품입니다. 모델 사양과 견적을 안내합니다.",
            "https://shop.example/product/db-1",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_normal_product_context")

    def test_it_database_procurement_is_not_personal_data_trade(self) -> None:
        reason = relevance_gate_reason(
            "채널계 전용 DB 분리 구축을 위한 eXperDB 구매 입찰공고",
            "시스템 안정성 강화를 위한 DBMS 환경 구축 프로젝트입니다. "
            "eXperDB 소프트웨어 구매 및 설치, 기술지원과 교육을 "
            "제공할 입찰 참가업체를 모집합니다.",
            "https://association.example/notice/1",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_it_database_system")

    def test_privacy_compliance_keyword_article_is_not_a_db_offer(self) -> None:
        reason = relevance_gate_reason(
            "대출DB판매 대신 알아보는 상담 정보",
            "대출DB판매라는 검색어를 볼 때 확인할 내용입니다. "
            "개인정보 수집과 이용 목적, 이용자 동의 여부를 확인해야 "
            "합니다. 개인정보 제3자 제공 절차와 정보의 출처도 "
            "신중하게 확인해야 합니다. 자주 묻는 질문을 정리합니다.",
            "https://cleaning.example/gallery/1",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_db_compliance_guide")

    def test_overseas_futures_rental_account_is_not_bank_account_trade(self) -> None:
        reason = relevance_gate_reason(
            "해외선물 대여계좌 대여업체",
            "나스닥 실시간 미국 선물지수 거래를 지원합니다. "
            "전문가 교육과 실시간 담보금 예치, 실체결 거래 서비스를 "
            "상담하세요.",
            "https://futures.example/rental-account",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_investment_trading_service")

    def test_crypto_wallet_p2p_guide_is_not_bank_account_trade(self) -> None:
        reason = relevance_gate_reason(
            "텔레그램 지갑 P2P 마켓 이용 방법 및 안전 거래 가이드",
            (
                "판매자의 코인은 공식 에스크로 지갑에 잠금 처리됩니다. "
                "P2P 이용 전 본인 확인 KYC를 승인하고 본인 명의 은행 계좌를 "
                "등록합니다. 판매자가 입금을 확인하면 USDT 코인 릴리즈가 "
                "완료되며 분쟁 시스템으로 안전하게 거래합니다."
            ),
            "https://guide.example/wallet/p2p",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_normal_product_context")

    def test_blog_false_positive_contexts_are_excluded(self) -> None:
        cases = [
            (
                "여권사진 과도한 보정, 공항에서 발목 잡히는 이유",
                "여권사진 촬영 규정과 보정 기준, 발급 반려 사유를 "
                "알아봅니다. 여권 위조로 오해받지 않도록 주의하세요.",
                "excluded_identity_photo_guide",
            ),
            (
                "외국인 불법체류와 여권 위조 심층 보고서",
                "목차, 서론, 현황, 사례 비교, 정책 제언, 제도 개선, "
                "결론을 통해 여권 위조 문제를 종합적으로 고찰합니다.",
                "excluded_informational_report",
            ),
            (
                "[채용] 보험설계사 이직 추천 DB 무한생성",
                "GA보험사 본부 채용 안내입니다. 소속 설계사에게 "
                "고객 DB를 무료로 제공하고 영업 시스템을 지원합니다.",
                "excluded_insurance_recruitment",
            ),
            (
                "DB 구매 없이 고객이 찾아오는 방법",
                "유료 DB를 사지 않고 콘텐츠로 고객 유입 구조를 "
                "직접 만드는 인바운드 영업을 소개합니다.",
                "excluded_db_purchase_alternative",
            ),
            (
                "DB하이텍, 지금 매수할까?",
                "DB하이텍은 8인치 파운드리 반도체 업체입니다. "
                "실적과 주가, 매수 시점을 분석합니다.",
                "excluded_db_brand_or_stock",
            ),
            (
                "보험 판매수수료 개편 이후 GA 변화",
                "1,200% 룰과 4년 분급제 규제가 보험 설계사 조직에 "
                "미치는 영향과 전략을 분석합니다.",
                "excluded_insurance_industry_analysis",
            ),
            (
                "네이버 아이디삽니다, 팔지 마세요",
                "본인 명의를 넘기는 것 자체가 잘못된 생각입니다. "
                "계정을 판매하게 되면 사기 범죄에 악용될 수 있으니 무시하세요.",
                "excluded_trade_warning",
            ),
        ]
        for title, text, expected in cases:
            with self.subTest(expected=expected):
                self.assertEqual(
                    relevance_gate_reason(
                        title,
                        text,
                        "https://blog.example/post/1",
                        "unknown",
                        "intent",
                    ),
                    expected,
                )

    def test_simhash_is_stable_for_equivalent_token_order(self) -> None:
        first = near_duplicate_id("제목", "반복 문구 반복 문구")
        second = near_duplicate_id("제목", "반복 문구 반복 문구")
        self.assertEqual(first, second)
        self.assertRegex(first, r"^simhash64:[0-9a-f]{16}$")

    def test_existing_fingerprints_supports_current_and_legacy_columns(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "candidates_masked.csv"
            with path.open("w", encoding="utf-8-sig", newline="") as handle:
                writer = csv.DictWriter(
                    handle,
                    fieldnames=[
                        "near_duplicate_fingerprint",
                        "near_duplicate_cluster",
                    ],
                )
                writer.writeheader()
                writer.writerow(
                    {
                        "near_duplicate_fingerprint": "simhash64:current",
                        "near_duplicate_cluster": "simhash64:legacy-alias",
                    }
                )
                writer.writerow(
                    {
                        "near_duplicate_fingerprint": "",
                        "near_duplicate_cluster": "simhash64:legacy",
                    }
                )
            self.assertEqual(
                existing_fingerprints(path),
                {"simhash64:current", "simhash64:legacy"},
            )

    def test_private_query_yaml_is_validated(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "queries.local.yaml"
            path.write_text(
                "groups:\n"
                "  - name: internal_example\n"
                "    detection_type: 개인정보DB\n"
                "    queries:\n"
                "      - private query\n",
                encoding="utf-8",
            )
            specs = load_query_specs(path)
            self.assertEqual(len(specs), 1)
            self.assertEqual(specs[0].group, "internal_example")
            self.assertEqual(specs[0].detection_type, "개인정보DB")

    def test_query_variants_are_created_in_stable_order(self) -> None:
        specs = [QuerySpec("group", "기타", "테스트 문의")]
        expanded = expand_query_specs(specs, 4)
        self.assertEqual(len(expanded), 4)
        self.assertEqual(expanded[0].query, "테스트 문의")
        self.assertEqual(expanded[-1].query, "테스트 문의 문의")

    def test_requested_search_provider_order_is_preserved(self) -> None:
        available = [
            "naver",
            "daum",
            "nate",
            "yahoo_japan",
            "bing",
            "google",
        ]
        self.assertEqual(
            ordered_provider_names(
                available,
                ["nate", "google", "naver", "nate"],
            ),
            ["nate", "google", "naver"],
        )
        self.assertEqual(ordered_provider_names(available, None), available)

    def test_nate_search_url_uses_live_pagination_contract(self) -> None:
        self.assertEqual(
            nate_search_url("대출DB 판매", 0),
            "https://search.daum.net/nate?w=tot&q=%EB%8C%80%EC%B6%9CDB+%ED%8C%90%EB%A7%A4",
        )
        self.assertEqual(
            nate_search_url("대출DB 판매", 2),
            "https://search.daum.net/nate?w=fusion&q=%EB%8C%80%EC%B6%9CDB+%ED%8C%90%EB%A7%A4&p=3&DA=PGD",
        )
        self.assertEqual(
            nate_search_url("대출DB 판매 -site:t.me -뉴스", 0),
            "https://search.daum.net/nate?w=tot&q=%EB%8C%80%EC%B6%9CDB+%ED%8C%90%EB%A7%A4",
        )

    def test_yahoo_japan_search_url_uses_live_pagination_contract(self) -> None:
        self.assertEqual(
            yahoo_japan_search_url("신분증위조 텔레그램", 0),
            "https://search.yahoo.co.jp/search?p=%EC%8B%A0%EB%B6%84%EC%A6%9D%EC%9C%84%EC%A1%B0+%ED%85%94%EB%A0%88%EA%B7%B8%EB%9E%A8&ei=UTF-8&b=1",
        )
        self.assertEqual(
            yahoo_japan_search_url("신분증위조 텔레그램", 2),
            "https://search.yahoo.co.jp/search?p=%EC%8B%A0%EB%B6%84%EC%A6%9D%EC%9C%84%EC%A1%B0+%ED%85%94%EB%A0%88%EA%B7%B8%EB%9E%A8&ei=UTF-8&b=21",
        )

    def test_google_api_discovery_uses_snippet_relevance_without_key_output(self) -> None:
        class FakeResponse:
            status_code = 200

            def json(self):
                return {
                    "items": [
                        {
                            "link": "https://example.com/public/post/1",
                            "title": "디비 텔그",
                            "snippet": "판매 관련 연락 안내",
                        }
                    ]
                }

            def close(self) -> None:
                pass

        class FakeSession:
            def __init__(self) -> None:
                self.params = {}

            def get(self, _url, **kwargs):
                self.params = kwargs["params"]
                return FakeResponse()

        session = FakeSession()
        candidates = discover_google_api_candidates(
            session,
            [QuerySpec("group", "개인정보DB", "디비 텔그")],
            desired=1,
            pages=1,
            delay=0,
            prefilter_mode="review",
            api_key="secret-key",
            cse_id="engine-id",
        )
        self.assertEqual(len(candidates), 1)
        self.assertEqual(candidates[0].url, "https://example.com/public/post/1")
        self.assertEqual(session.params["key"], "secret-key")

    def test_serpapi_discovery_combines_pages_into_one_paid_request(self) -> None:
        class FakeResponse:
            status_code = 200

            def json(self):
                return {
                    "organic_results": [
                        {
                            "link": "https://example.com/public/post/1",
                            "title": "디비 텔그",
                            "snippet": "판매 관련 연락 안내",
                        }
                    ]
                }

            def close(self) -> None:
                pass

        class FakeSession:
            def __init__(self) -> None:
                self.calls = 0
                self.params = {}

            def get(self, _url, **kwargs):
                self.calls += 1
                self.params = kwargs["params"]
                return FakeResponse()

        session = FakeSession()
        candidates = discover_serpapi_candidates(
            session,
            [QuerySpec("group", "개인정보DB", "디비 텔그")],
            desired=1,
            pages=3,
            delay=0,
            prefilter_mode="review",
            api_key="secret-key",
        )
        self.assertEqual(len(candidates), 1)
        self.assertEqual(candidates[0].search_provider, "serpapi")
        self.assertEqual(session.calls, 1)
        self.assertEqual(session.params["num"], 30)
        self.assertEqual(session.params["api_key"], "secret-key")

    def test_query_limit_round_robins_across_groups(self) -> None:
        specs = [
            QuerySpec("one", "기타", "one-a"),
            QuerySpec("one", "기타", "one-b"),
            QuerySpec("two", "기타", "two-a"),
            QuerySpec("two", "기타", "two-b"),
            QuerySpec("three", "기타", "three-a"),
        ]
        limited = limit_query_specs_by_group(specs, 4)
        self.assertEqual(
            [(item.group, item.query) for item in limited],
            [
                ("one", "one-a"),
                ("two", "two-a"),
                ("three", "three-a"),
                ("one", "one-b"),
            ],
        )

    def test_strict_queries_use_phrases_and_negative_filters(self) -> None:
        specs = [QuerySpec("group", "기타", "고객 DB 판매")]
        constrained = constrain_query_specs(specs)
        self.assertEqual(len(constrained), 4)
        self.assertTrue(constrained[0].query.startswith("고객 DB 판매"))
        self.assertTrue(constrained[1].query.startswith('"고객 DB 판매"'))
        self.assertTrue(all("-개인정보처리방침" in item.query for item in constrained))

    def test_search_snippet_prefilter_requires_target_and_trade_context(self) -> None:
        relevant = Candidate(
            "https://example.com/post",
            "group",
            "기타",
            discovery_text="고객 DB를 대량 판매한다는 게시물과 텔레그램 문의 안내",
        )
        generic = Candidate(
            "https://example.com/help",
            "group",
            "기타",
            discovery_text="고객센터에서 개인정보 처리방침을 확인하세요",
        )
        self.assertTrue(discovery_candidate_relevant(relevant))
        self.assertFalse(discovery_candidate_relevant(generic))
        target_only = Candidate(
            "https://example.com/news",
            "group",
            "기타",
            discovery_text="개인정보 유출 사고를 다룬 국내 기사",
        )
        self.assertTrue(discovery_candidate_passes(target_only, "labeling"))
        self.assertFalse(discovery_candidate_passes(target_only, "review"))
        hard_negative = Candidate(
            "https://example.com/database-guide",
            "group",
            "기타",
            discovery_text="DB 계정 관리 방법을 설명하는 기술 문서",
        )
        self.assertTrue(discovery_candidate_passes(hard_negative, "labeling"))
        self.assertFalse(discovery_candidate_passes(hard_negative, "review"))
        trade_only = Candidate(
            "https://example.com/classified/7",
            "group",
            "기타",
            discovery_text="대량 판매합니다. 자세한 품목은 본문을 확인하세요.",
        )
        self.assertTrue(discovery_candidate_passes(trade_only, "labeling"))
        self.assertFalse(discovery_candidate_passes(trade_only, "review"))
        self.assertGreater(
            discovery_relevance_score(relevant),
            discovery_relevance_score(generic),
        )

    def test_strict_search_prefilter_requires_local_direct_offer(self) -> None:
        direct_offer = Candidate(
            "https://example.com/post/positive",
            "group",
            "개인정보DB",
            discovery_text=(
                "고객 DB 판매합니다. 건당 단가 문의 "
                "https://t.me/private_handle"
            ),
        )
        reporting = Candidate(
            "https://example.com/news/negative",
            "group",
            "개인정보DB",
            discovery_text="고객 DB 판매 사건을 경찰이 적발한 텔레그램 관련 기사",
        )
        destination_contact_deferred = Candidate(
            "https://example.com/post/weak",
            "group",
            "개인정보DB",
            discovery_text="고객 DB 판매합니다. 텔레그램에서 안내합니다.",
        )
        fused_card = Candidate(
            "https://example.com/post/fused",
            "group",
            "개인정보DB",
            discovery_text=(
                "고객 DB 관련 보안 안내 "
                + "일반 설명 " * 80
                + "중고 자동차 부품 판매합니다. 연락 [PHONE]"
            ),
        )
        self.assertTrue(discovery_candidate_passes(direct_offer, "strict"))
        self.assertFalse(discovery_candidate_passes(reporting, "strict"))
        self.assertTrue(
            discovery_candidate_passes(destination_contact_deferred, "strict")
        )
        self.assertFalse(discovery_candidate_passes(fused_card, "strict"))

    def test_intent_search_prefilter_rejects_press_and_known_normal_contexts(self) -> None:
        quoted_news = Candidate(
            "https://regional.example/news/articleView.html?idxno=10",
            "group",
            "기타",
            discovery_text=(
                "SNS 신분증 위조 판매 기승. 취재진이 판매자에게 제작을 "
                "문의했으며 경찰은 범죄 악용이 우려된다고 밝혔다."
            ),
        )
        game_trade = Candidate(
            "https://market.example/post/20",
            "group",
            "포털ID",
            discovery_text="쿠키런 카카오 계정 구매합니다. 희망 가격 문의",
        )
        real_listing = Candidate(
            "https://community.example/post/21",
            "group",
            "개인정보DB",
            discovery_text=(
                "대출DB 판매합니다. 건당 단가 상담은 텔레그램 "
                "raw_handle 로 문의주세요."
            ),
        )
        normal_id_shop = Candidate(
            "https://shop.example/category/id",
            "group",
            "신분증",
            discovery_text=(
                "사원증 학생증 방문증 신분증 제작 상품목록 60 items "
                "장바구니 배송조회"
            ),
        )
        trade_guide = Candidate(
            "https://blog.example/account-risk",
            "group",
            "포털ID",
            discovery_text=(
                "텔레그램 계정 거래 후기 위험 분석. 계정 회수 피해와 "
                "정책 위반 위험을 자세히 알아봅니다."
            ),
        )
        used_phone_reset = Candidate(
            "https://blog.example/used-phone-reset",
            "group",
            "포털ID",
            discovery_text=(
                "중고 폰 팔 때 필수! 아이폰·갤럭시 초기화 및 구글 계정 "
                "해제법. 계정을 로그아웃한 뒤 공장 초기화하는 순서를 "
                "안내합니다."
            ),
        )
        db_promy_trip = Candidate(
            "https://blog.example/basketball-trip",
            "group",
            "개인정보DB",
            discovery_text=(
                "원주 DB프로미 홈 개막전 직관 후기. 랜덤 포토카드도 "
                "판매 중이라 계좌이체로 구매했습니다."
            ),
        )
        db_linked_saas = Candidate(
            "https://blog.example/saas-guide",
            "group",
            "개인정보DB",
            discovery_text=(
                "결제와 인증, DB 연동이 된 실제 판매 가능한 SaaS 출시. "
                "AI 코딩 플랫폼 개발 방법을 소개합니다."
            ),
        )
        frozen_account_remedy = Candidate(
            "https://blog.example/account-freeze-remedy",
            "group",
            "통장",
            discovery_text=(
                "계좌지급정지 해제와 통장대여 전기통신금융사기 구제법. "
                "은행 안내와 거래 내역, 객관적 자료를 갖춰 이의신청하는 "
                "절차입니다."
            ),
        )
        forged_id_reporting = Candidate(
            "https://blog.example/social-issue",
            "group",
            "신분증",
            discovery_text=(
                "[사회] SNS 점령한 위조 신분증 광고, 사기 범죄의 온상 "
                "되나. 텔레그램에서 제작 광고가 기승을 부려 시민들의 "
                "주의가 요구됩니다."
            ),
        )
        public_smartstore_db = Candidate(
            "https://market.example/smartstore-db",
            "group",
            "개인정보DB",
            discovery_text=(
                "스마트스토어 판매자 DB 웹 프로그램. 업체명, 연락처, "
                "스토어 URL과 전 카테고리 업종을 실시간 수집합니다."
            ),
        )
        mobile_id_verification = Candidate(
            "https://government.example/mobile-id-check",
            "group",
            "신분증",
            discovery_text=(
                "모바일 주민등록증 위조 여부는 검증앱에서 QR코드를 "
                "스캔하면 빠르게 확인할 수 있습니다."
            ),
        )
        permission_cpa_db = Candidate(
            "https://leads.example/permission-cpa",
            "group",
            "개인정보DB",
            discovery_text=(
                "퍼미션 DB·노리워드 CPA 디비 전문. 고객이 먼저 찾아오는 "
                "고품질 리드를 공급해 광고비와 상담 전환율을 개선합니다."
            ),
        )
        self.assertFalse(discovery_candidate_passes(quoted_news, "intent"))
        self.assertFalse(discovery_candidate_passes(game_trade, "intent"))
        self.assertFalse(discovery_candidate_passes(normal_id_shop, "intent"))
        self.assertFalse(discovery_candidate_passes(trade_guide, "intent"))
        self.assertFalse(discovery_candidate_passes(used_phone_reset, "intent"))
        self.assertFalse(discovery_candidate_passes(db_promy_trip, "intent"))
        self.assertFalse(discovery_candidate_passes(db_linked_saas, "intent"))
        self.assertFalse(
            discovery_candidate_passes(frozen_account_remedy, "intent")
        )
        self.assertFalse(discovery_candidate_passes(forged_id_reporting, "intent"))
        self.assertFalse(discovery_candidate_passes(public_smartstore_db, "intent"))
        self.assertFalse(
            discovery_candidate_passes(mobile_id_verification, "intent")
        )
        self.assertFalse(discovery_candidate_passes(permission_cpa_db, "intent"))
        self.assertTrue(discovery_candidate_passes(real_listing, "intent"))

    def test_shorthand_target_and_contact_pass_review_prefilter(self) -> None:
        shorthand = Candidate(
            "https://example.com/post/8",
            "group",
            "개인정보DB",
            discovery_text="디비 텔그 문의",
        )
        brand_noise = Candidate(
            "https://example.com/news/8",
            "group",
            "개인정보DB",
            discovery_text="DB손해보험 농구단 소식",
        )
        self.assertTrue(discovery_candidate_relevant(shorthand))
        self.assertFalse(discovery_candidate_relevant(brand_noise))
        self.assertTrue(discovery_candidate_passes(shorthand, "review"))
        self.assertEqual(
            relevance_gate_reason(
                "디비 텔그",
                "보유 자료 관련 연락 안내입니다.",
                "https://example.com/post/8",
                "unknown",
                "review",
            ),
            "",
        )

    def test_review_prefilter_does_not_treat_display_url_as_contact(self) -> None:
        normal_result = Candidate(
            "https://support.example/contacts",
            "group",
            "기타",
            discovery_text=(
                "갤럭시 연락처 가져오기 방법 "
                "https://support.example/contacts 연락처 앱 사용 안내"
            ),
        )
        self.assertFalse(discovery_candidate_passes(normal_result, "review"))

    def test_keyword_expansion_requires_repetition_across_domains(self) -> None:
        candidates = [
            Candidate(
                "https://one.example/post/1",
                "personal_info_db",
                "개인정보DB",
                discovery_text="고객디비 텔그 문의",
            ),
            Candidate(
                "https://two.example/post/2",
                "personal_info_db",
                "개인정보DB",
                discovery_text="고객디비 텔그 연락",
            ),
            Candidate(
                "https://one.example/post/3",
                "portal_accounts",
                "포털ID",
                discovery_text="희귀아이디 텔그 문의",
            ),
        ]
        expansions = mine_keyword_expansions(
            candidates,
            [QuerySpec("seed", "개인정보DB", "디비 텔그")],
            round_number=1,
            limit=10,
            minimum_domains=2,
        )
        self.assertEqual([item.query for item in expansions], ["고객디비 텔그"])
        self.assertEqual(expansions[0].domain_frequency, 2)

    def test_candidate_merge_combines_snippet_evidence(self) -> None:
        current = [
            Candidate(
                "https://example.com/post",
                "group",
                "기타",
                discovery_text="디비 판매",
                search_provider="bing",
            )
        ]
        additions = [
            Candidate(
                "https://example.com/post",
                "group",
                "기타",
                discovery_text="텔그 문의",
                search_provider="google_api",
            )
        ]
        merged = merge_candidates(current, additions)
        self.assertEqual(len(merged), 1)
        self.assertIn("디비 판매", merged[0].discovery_text)
        self.assertIn("텔그 문의", merged[0].discovery_text)
        self.assertEqual(merged[0].search_provider, "bing,google_api")

    def test_contact_required_expansion_rejects_trade_only_pair(self) -> None:
        candidates = [
            Candidate(
                "https://one.example/post/1",
                "portal_accounts",
                "포털ID",
                discovery_text="계정 거래 안내",
            ),
            Candidate(
                "https://two.example/post/2",
                "portal_accounts",
                "포털ID",
                discovery_text="계정 거래 게시물",
            ),
        ]
        expansions = mine_keyword_expansions(
            candidates,
            [],
            round_number=1,
            limit=10,
            minimum_domains=2,
            require_contact=True,
        )
        self.assertEqual(expansions, [])

    def test_identity_document_forgery_passes_review_gate(self) -> None:
        self.assertEqual(
            relevance_gate_reason(
                "신분증 위조 텔그",
                "제작 관련 연락 안내입니다.",
                "https://example.com/post/9",
                "unknown",
                "review",
            ),
            "",
        )

    def test_relevance_gate_keeps_local_trade_and_contact_signals(self) -> None:
        reason = relevance_gate_reason(
            "고객 DB 판매합니다",
            "최신 자료 대량 보유, 건당 단가 문의 텔레그램 [MESSENGER_ID]",
            "https://board.example/post/7",
            "unknown",
            "strict",
        )
        self.assertEqual(reason, "")

    def test_intent_gate_prioritizes_direct_offer_without_contact(self) -> None:
        candidate = Candidate(
            "https://board.example/post/intent",
            "group",
            "개인정보DB",
            discovery_text="고객 DB를 대량 보유하고 판매합니다. 건당 단가 안내",
        )
        self.assertTrue(discovery_candidate_passes(candidate, "intent"))
        self.assertEqual(
            relevance_gate_reason(
                "고객 DB 판매합니다",
                "최신 고객 명단을 대량 보유하고 건당 판매합니다.",
                candidate.url,
                "unknown",
                "intent",
            ),
            "",
        )
        self.assertEqual(
            relevance_gate_reason(
                "고객 DB 판매합니다",
                "최신 고객 명단을 대량 보유하고 건당 판매합니다.",
                candidate.url,
                "unknown",
                "strict",
            ),
            "missing_concrete_contact",
        )

    def test_intent_gate_rejects_reporting_and_missing_offer(self) -> None:
        reporting = relevance_gate_reason(
            "고객 DB 판매 게시물 적발",
            "경찰이 개인정보 명단 거래 사건을 수사하고 있습니다.",
            "https://news.example/article/intent",
            "news_or_education",
            "intent",
        )
        no_offer = relevance_gate_reason(
            "고객 DB 안내",
            "고객정보 데이터베이스의 보관 방식을 설명합니다.",
            "https://board.example/post/no-offer",
            "unknown",
            "intent",
        )
        self.assertEqual(reporting, "excluded_page_type")
        self.assertEqual(no_offer, "missing_body_offer")

    def test_intent_gate_rejects_press_copy_without_news_word(self) -> None:
        reason = relevance_gate_reason(
            '양홍원 "인스타그램 아이디 1천만원에 팝니다"…무슨 일?',
            (
                "[마이데일리 = 이승길 기자] SNS 계정을 판매하겠다는 "
                "글을 올려 이목이 쏠리고 있다. 무단전재 및 재배포 금지."
            ),
            "https://media.example/page/view/123",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_reporting_context")

    def test_v6_press_paths_and_reposted_reporting_are_excluded(self) -> None:
        press_text = (
            "편집자 주 쿠팡 구매자는 판매업체에 계좌번호 등 개인정보를 "
            "제공해야 환불받을 수 있는 것으로 확인됐다."
        )
        self.assertEqual(
            classify_page_type(
                "https://regional.example/news/articleView.html?idxno=10",
                "계좌 등 개인정보 제공해야 환불",
                press_text,
            ),
            "news_or_education",
        )
        repost = relevance_gate_reason(
            "홀로그램 가짜 주민등록증 5만원에 판매",
            (
                "중앙일보 취재를 종합하면 SNS에서 신분증을 만들어드립니다라는 "
                "글을 쉽게 찾을 수 있다. 판매자에게 제작을 문의하자 답했다."
            ),
            "https://community.example/view/10",
            "unknown",
            "intent",
        )
        self.assertEqual(repost, "excluded_reporting_context")

    def test_v6_telegram_shell_requires_offer_copy(self) -> None:
        empty_view = relevance_gate_reason(
            "Telegram: View coffee",
            (
                "커피 바이럴 마케팅 계정 매입\n"
                "If you have Telegram, you can view post and join right away."
            ),
            "https://t.me/channel/1",
            "public_messenger_page",
            "intent",
        )
        ambiguous_contact = relevance_gate_reason(
            "Telegram: Contact account",
            (
                "Download\n계좌매입\n제보센터입니다\nSend Message\n"
                "If you have Telegram, you can contact 계좌매입 right away."
            ),
            "https://t.me/account",
            "public_messenger_page",
            "intent",
        )
        explicit_contact = relevance_gate_reason(
            "Telegram: Contact seller",
            (
                "카톡계정 및 텔레그램 계정 최저가 및 각종 DB전문판매업체 "
                "24시간 상담가능\nIf you have Telegram, you can contact seller right away."
            ),
            "https://t.me/seller",
            "public_messenger_page",
            "intent",
        )
        self.assertEqual(empty_view, "excluded_empty_container")
        self.assertEqual(ambiguous_contact, "excluded_empty_container")
        self.assertEqual(explicit_contact, "")

    def test_empty_marketplace_form_does_not_inherit_intent_from_title(self) -> None:
        reason = relevance_gate_reason(
            "페이스북 실계정 다 삽니다",
            (
                "판매글 등록 유의사항\n"
                "1. 페이지명 :\n2. 팔로워 수 :\n3. 매매가 :\n"
                "4. 안전거래 가능 여부 :\n"
                "5. 거래 문의 연락처 (이메일 주소/오픈카톡방 링크 등) :\n"
                "추가로 작성해주시고 싶은 내용이 있으면 적어주세요."
            ),
            "https://market.example/post/empty",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_empty_listing_template")

    def test_marketplace_category_intro_is_not_a_trade_post(self) -> None:
        reason = relevance_gate_reason(
            "인스타 계정 - 사이트양도",
            (
                "팝니다: 다양한 채널 판매 게시판\n"
                "게시글을 등록하기 위해서는 회원가입이 필요합니다. "
                "판매 게시글은 아래의 카테고리에 맞게 등록해주세요.\n"
                "전체 워드프레스 네이버 카페 네이버 밴드 유튜브 채널 "
                "인스타 계정 틱톡 계정 페이스북 쇼핑몰 그외 채널\n"
                "29만원으로 애드센스 승인 대행 서비스"
            ),
            "https://siteyangdo.com/sell/category/126",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_empty_listing_template")

    def test_v6_normal_id_products_props_and_photo_guide_are_excluded(self) -> None:
        normal_shop = relevance_gate_reason(
            "신분증 제작",
            (
                "사원증 협회신분증 종교신분증 학생증 방문증 Total 60 items "
                "상품명: 학생증01 상품명: 방문증1 장바구니 배송조회"
            ),
            "https://shop.example/category/id",
            "unknown",
            "intent",
        )
        prop = relevance_gate_reason(
            "의사면허증 제작 촬영소품",
            "민감정보를 비식별화한 촬영용 의사면허증 소품을 제작합니다.",
            "https://props.example/medical-license",
            "unknown",
            "intent",
        )
        photo = relevance_gate_reason(
            "자동차운전전문학원",
            "운전면허증 제작용 사진의 표준 규격을 참고해 사진을 제출하세요.",
            "https://academy.example/license-photo",
            "unknown",
            "intent",
        )
        self.assertEqual(normal_shop, "excluded_normal_product_context")
        self.assertEqual(prop, "excluded_legal_prop_or_security_guide")
        self.assertEqual(photo, "excluded_question_or_guide")

    def test_v6_keyword_spam_guides_and_unrelated_commodities_are_excluded(self) -> None:
        seo_template = relevance_gate_reason(
            "010인증 네이버아이디판매 계정대여",
            (
                "네이버계정판매 네이버아이디구매 네이버계정대여. "
                "예약 방법을 안내합니다. 평일 주말 모두 영업합니다. "
                "심야 시간대에도 가능합니다. 날짜와 인원을 알려주세요. "
                "첫 방문 고객 혜택과 이용 후기가 있습니다. 분위기를 찾는 분께 추천합니다."
            ),
            "https://unrelated.example/post/1",
            "unknown",
            "intent",
        )
        guide = relevance_gate_reason(
            "텔레그램 아이디 거래 후기 위험 완벽 분석",
            (
                "계정 회수 피해와 예상치 못한 위험을 자세히 알아봅니다. "
                "비공식 거래는 정책 위반이며 안전한 방법을 확인해야 합니다."
            ),
            "https://blog.example/guide",
            "unknown",
            "intent",
        )
        gift_card = relevance_gate_reason(
            "각종 상품권 판매 및 구매합니다",
            (
                "상품권 도소매 매입 판매 문의. 상품권 현금화와 상품권 구매를 "
                "도와드립니다. 카톡아이디 seller 텔레 아이디 seller2"
            ),
            "https://blog.example/gift-card",
            "unknown",
            "intent",
        )
        self.assertEqual(seo_template, "excluded_keyword_stuffing")
        self.assertEqual(guide, "excluded_informational_article")
        self.assertEqual(gift_card, "excluded_normal_product_context")

        gift_card_shop = relevance_gate_reason(
            "논현상품권",
            (
                "지류상품권 매입과 판매를 합니다. 상품권을 구매/판매하는 "
                "양에 따라 할인율이 달라지며 계좌이체와 현금가 가격은 "
                "차이가 있습니다. 상품권영수증은 매입/매출계산서가 아닙니다."
            ),
            "https://giftcard.example/",
            "unknown",
            "intent",
        )
        self.assertEqual(gift_card_shop, "excluded_normal_product_context")

        gift_card_watch_shop = relevance_gate_reason(
            "상품권·중고시계 매입·판매 정식 업체",
            (
                "사업자 등록을 마친 정식 업체로 상품권을 매입·판매합니다. "
                "지류·모바일 상품권과 중고 명품시계를 취급하며 세금계산서 "
                "발행이 가능합니다. 상품권 진위와 도용 여부를 확인합니다."
            ),
            "https://giftcard.example/watch-shop",
            "unknown",
            "intent",
        )
        self.assertEqual(
            gift_card_watch_shop,
            "excluded_normal_product_context",
        )

    def test_v6_clear_account_sale_on_unrelated_board_is_still_kept(self) -> None:
        reason = relevance_gate_reason(
            "대외활동 게시판 인스타 계정 판매",
            (
                "마케팅용으로 사용하던 인스타그램 계정 판매합니다. "
                "12년 계정부터 다양하게 있으며 입금 가능합니다. "
                "오픈카톡과 텔레그램 seller 로 문의주세요."
            ),
            "https://campus.example/board/1",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "")

    def test_bulk_discord_pva_account_storefront_is_in_scope(self) -> None:
        reason = relevance_gate_reason(
            "디스코드 PVA 계정 대량 판매",
            (
                "휴대폰 인증을 마친 디스코드 계정을 대량 보유하고 판매합니다. "
                "소량과 도매 모두 즉시 공급하며 교환 A/S를 제공합니다. "
                "텔레그램 seller123 문의 바랍니다."
            ),
            "https://accounts.example/discord-pva",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "")

    def test_discord_account_setup_guide_has_no_trade_intent(self) -> None:
        reason = relevance_gate_reason(
            "디스코드 계정 설정 가이드",
            "공식 앱에서 이메일 인증과 2단계 인증을 설정하는 방법입니다.",
            "https://guide.example/discord-account",
            "unknown",
            "intent",
        )
        self.assertNotEqual(reason, "")

    def test_account_storefront_lead_survives_generic_seo_filler(self) -> None:
        reason = relevance_gate_reason(
            "010실명인증 텔레 Run55 발번등록업자 국내본인인증문자판매",
            (
                "010 인증 네이버 아이디 판매 서비스를 전문으로 제공합니다. "
                "비실명 계정과 준최적화 블로그를 수작업 생성하여 판매합니다. "
                "365일 24시간 문의 가능하고 구매 계정은 최초 로그인 때 "
                "교환 및 A/S 사후관리를 제공합니다. 텔레 Run55 문의. "
                "단골 혜택 친절한 응대 투명한 가격 위치와 접근성 "
                "재방문 의사 편하게 즐길 수 있는 분위기 운영 시간 "
                "원하시는 날짜와 인원을 알려주세요."
            ),
            "https://board.example/post/account-storefront",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "")

    def test_marketplace_result_page_is_not_retained_as_a_post(self) -> None:
        repeated_controls = (
            "쪽지보내기 메일보내기 자기소개 아이디로 검색 전체게시물 " * 3
        )
        page_type = classify_page_type(
            "https://market.example/account-trade",
            "틱톡 계정 거래",
            (
                "전체 3,057건 / 1 페이지 팝니다 삽니다 거래완료 "
                "틱톡 계정 판매합니다. "
                + repeated_controls
                + "게시물 검색 검색대상 제목 내용 글쓴이"
            ),
        )
        self.assertEqual(page_type, "board_listing")
        self.assertEqual(
            relevance_gate_reason(
                "틱톡 계정 거래",
                "전체 5건 / 1 페이지 계정 판매합니다 " + repeated_controls,
                "https://market.example/account-trade",
                page_type,
                "intent",
            ),
            "excluded_page_type",
        )
        card_layout = classify_page_type(
            "https://market.example/igtrade",
            "인스타 계정 거래 - 마켓",
            (
                "전체 3,057건 / 1 페이지 판매중 팔로워 250 계정 판매 "
                "판매중 인스타 계정 매입 판매중 계정 팝니다"
            ),
        )
        self.assertEqual(card_layout, "board_listing")

    def test_intent_gate_rejects_reposted_warning_channel(self) -> None:
        reason = relevance_gate_reason(
            "저승사자 박제채널",
            (
                "사건내용: 계정 매입 업자에게 피해를 입어 제보합니다. "
                "이전 홍보글에는 계정 매입합니다라는 문구가 있습니다."
            ),
            "https://t.me/s/report_channel",
            "public_messenger_page",
            "intent",
        )
        self.assertEqual(reason, "excluded_aggregation_or_commentary")

    def test_intent_gate_rejects_legal_question_and_news_repost(self) -> None:
        legal_question = relevance_gate_reason(
            "계정 판매 후 명의 정지 고소 됨?",
            "계정을 넘긴 뒤 고소될 수 있는지 묻는 글입니다.",
            "https://community.example/question/3",
            "unknown",
            "intent",
        )
        repost = relevance_gate_reason(
            "이슈/유머 - 개인정보 DB 팝니다",
            "기사본문을 옮긴 게시물입니다.",
            "https://community.example/issue/4",
            "unknown",
            "intent",
        )
        self.assertEqual(legal_question, "excluded_question_or_guide")
        self.assertEqual(repost, "excluded_aggregation_or_commentary")

    def test_intent_gate_keeps_illicit_passport_issue_offer(self) -> None:
        reason = relevance_gate_reason(
            "여권발급이나 새 신분이 필요한 분",
            "위조 여권 제작 가능합니다. 다크웹 판매업체로 연락하세요.",
            "https://board.example/post/passport",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "")

    def test_intent_gate_rejects_normal_products_and_guides(self) -> None:
        passport_case = relevance_gate_reason(
            "여권 케이스 제작 방법",
            "맞춤형 디자인과 소재를 선택해 주문 제작합니다.",
            "https://shop.example/passport-case",
            "unknown",
            "intent",
        )
        bank_product = relevance_gate_reason(
            "농협 어린이 캐릭터 통장 판매",
            "은행에서 어린이 입출금 통장 상품을 출시해 판매합니다.",
            "https://bank.example/product/1",
            "unknown",
            "intent",
        )
        account_guide = relevance_gate_reason(
            "위탁판매 계정 정지 조건과 스마트스토어 차이",
            "정상 판매자가 계정 정지를 피하는 방법을 설명합니다.",
            "https://guide.example/post/1",
            "unknown",
            "intent",
        )
        self.assertEqual(passport_case, "excluded_question_or_guide")
        self.assertEqual(bank_product, "excluded_normal_product_context")
        self.assertEqual(account_guide, "excluded_question_or_guide")

        bankbook_case = relevance_gate_reason(
            "2026년 통장·카드케이스 구매 단가",
            "은행용 통장 케이스와 카드 비닐의 물품 입찰 공고입니다.",
            "https://bid.example/product/2",
            "unknown",
            "intent",
        )
        self.assertEqual(bankbook_case, "excluded_normal_product_context")

    def test_intent_gate_rejects_extraction_boilerplate(self) -> None:
        privacy_policy = relevance_gate_reason(
            "인스타 계정 판매",
            (
                "개인정보 처리방침 회사는 이용자의 개인정보를 보호합니다. "
                "수집 목적과 보유 기간을 안내합니다."
            ),
            "https://board.example/post/12",
            "unknown",
            "intent",
        )
        marketplace_footer = relevance_gate_reason(
            "인스타 계정 판매합니다",
            (
                "사업자등록번호 123-45-67890 통신판매업신고번호 안내. "
                "회사는 통신판매중개자로서 거래 당사자가 아닙니다."
            ),
            "https://market.example/item/12",
            "unknown",
            "intent",
        )
        self.assertEqual(privacy_policy, "excluded_extraction_boilerplate")
        self.assertEqual(marketplace_footer, "excluded_extraction_boilerplate")

        empty_request = relevance_gate_reason(
            "보험 퍼미션 인바운드 DB 마케팅",
            (
                "통신판매업신고 2018-서울-1234 문의하기. 회사는 "
                "통신판매중개자이며 상품 정보와 거래 책임은 판매회원에게 "
                "있습니다. 저작권법에 따라 무단복제를 금지합니다."
            ),
            "https://marketplace.example/empty-request",
            "unknown",
            "intent",
        )
        self.assertEqual(empty_request, "excluded_extraction_boilerplate")

    def test_insurance_recruitment_with_permission_db_benefit_is_excluded(self) -> None:
        job = relevance_gate_reason(
            "보험설계사 모집 3차 퍼미션DB 영업 특화",
            (
                "대형 GA 보험설계사 채용 공고입니다. 자체 콜센터를 운영해 "
                "3차 동의콜 녹취가 완료된 고객 DB를 무료로 공급합니다."
            ),
            "https://jobs.example/insurance-sales",
            "unknown",
            "intent",
        )
        branch = relevance_gate_reason(
            "신입 보험설계사 정착 지원 제도 안내",
            (
                "신입 설계사에게 정착지원금과 퍼미션 DB 30건을 무료 제공하고 "
                "상담 고객 리드를 배분하는 지사 모집 안내입니다."
            ),
            "https://branch.example/recruitment",
            "unknown",
            "intent",
        )
        self.assertEqual(job, "excluded_insurance_recruitment")
        self.assertEqual(branch, "excluded_insurance_recruitment")

        internal_program = relevance_gate_reason(
            "무료보험디비영업으로 영업 스트레스 받지 마세요",
            (
                "DB사업단에서 무료보험디비영업을 할 수 있게 도와드립니다. "
                "POM디비를 가공해 평균 이상 매출을 하는 설계사분들께 무료로 "
                "테스트를 진행합니다. 현재 사업단에 소속되어 일하는 설계사의 "
                "소득이 중요하며 고객 계약은 모두 본인에게 이관됩니다."
            ),
            "https://blog.example/insurance-agent-support",
            "unknown",
            "intent",
        )
        self.assertEqual(internal_program, "excluded_insurance_recruitment")

    def test_intent_gate_rejects_account_market_guide(self) -> None:
        reason = relevance_gate_reason(
            "구글 계정 판매: 안전하고 신뢰할 수 있는 옵션",
            (
                "구글 계정 판매 시장 동향과 규모를 설명합니다. "
                "공급업체 선정 시 구매자 리뷰와 비교표를 확인하세요."
            ),
            "https://guide.example/account-market",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_market_guide")

    def test_intent_gate_keeps_account_storefront_with_guide_section(self) -> None:
        reason = relevance_gate_reason(
            "구글 아이디 판매 해외-국내 깡통 계정 전문",
            (
                "구글아이디구매사이트로 해외/국내 깡통 계정 판매를 전문으로 "
                "하는 업체입니다. 24시간 문의 시 신속한 구입을 도와드립니다. "
                "텔레그램 문의하기 카카오톡 문의하기. 정상·실사용 Google 계정 "
                "대량 보유, 결제부터 계정 제공까지 즉시 발송합니다. "
                "구글 계정 주요 기능 안내: Gmail, YouTube, Google Ads를 "
                "하나의 계정으로 연동해 사용할 수 있습니다."
            ),
            "https://seller.example/google-accounts",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "")

    def test_intent_gate_rejects_generic_account_purchase_explainer(self) -> None:
        reason = relevance_gate_reason(
            "구글 아이디 판매 - 전문적인 서비스",
            (
                "1. 구글 아이디 판매란 무엇인가요? "
                "2. 왜 구글 아이디 판매가 필요한가요? "
                "3. 어떻게 구글 아이디를 구매할 수 있나요? "
                "4. 구글 아이디 판매 서비스의 중요성. "
                "신뢰할 만한 판매자를 찾고 가격을 협상하는 절차를 설명합니다."
            ),
            "https://blog.example/google-account-guide",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_market_guide")

    def test_intent_gate_rejects_structured_account_risk_guide(self) -> None:
        reason = relevance_gate_reason(
            "쿠팡 아이디 판매 및 네이버 아이디 구매 패키지",
            (
                "서론\n정의\n주요 특징\n장점\n문제점 및 주의사항\n"
                "법적 문제로 이용 약관 위반과 법적 처벌 위험이 있습니다. "
                "보안 문제로 개인정보 유출과 계정 도용 위험이 있습니다. "
                "안전한 대안과 이용 방법\n본인 휴대폰 번호로 직접 인증하고 "
                "2단계 인증과 정기적인 비밀번호 변경을 권합니다.\n결론"
            ),
            "https://guide.example/account-package",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_market_guide")

        direct_seller = relevance_gate_reason(
            "네이버 아이디 판매",
            (
                "법적 문제와 계정 도용 위험을 확인하세요. 인증 계정 10개를 "
                "판매합니다. 단가 2만원, 주문은 텔레그램 [ACCOUNT] 문의하세요."
            ),
            "https://seller.example/account-package",
            "unknown",
            "intent",
        )
        self.assertEqual(direct_seller, "")

    def test_formal_b2b_db_consultation_is_not_an_illicit_listing(self) -> None:
        reason = relevance_gate_reason(
            "보험DB 판매 | 영업DB 전문",
            (
                "B2B 대량 구매 비대면 상담 신청하기. 마스케어DB 대량 구매와 "
                "교육 서비스 안내, 협업을 희망하시면 신청서 작성 후 온라인 "
                "ZOOM 상담을 진행합니다. 신청 시간 30분 전에 입장 링크를 "
                "전달하며 소요시간은 최대 50분입니다. 상담사님의 시간과 "
                "비용 효율을 높이고 고객님과의 상담에 집중하세요."
            ),
            "https://business.example/b2b-consultation",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_formal_b2b_db_service")

    def test_consented_traceable_db_marketplace_is_not_illicit_distribution(self) -> None:
        reason = relevance_gate_reason(
            "디비매입 | 업종별 DB 가치 평가",
            (
                "DB매입 시 고객 동의 여부와 개인정보 활용 동의를 확인합니다. "
                "고객 동의가 확인되는 DB만 검수하며 동의 여부를 다시 봅니다. "
                "수집 경로와 유입 경로, 동의 근거, 개인정보 보호 기준, "
                "적법성 및 유입 시점과 문의 목적을 확인합니다."
            ),
            "https://lead.example/",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_consent_based_db_marketplace")

        illicit = relevance_gate_reason(
            "대출DB 판매",
            (
                "고객 동의 자료라고 홍보하는 최신 대출DB를 판매합니다. "
                "샘플과 단가는 텔레그램 [ACCOUNT] 문의하세요."
            ),
            "https://seller.example/db",
            "unknown",
            "intent",
        )
        self.assertEqual(illicit, "")

        permission_service = relevance_gate_reason(
            "보험 퍼미션 DB 판매업체",
            (
                "소비자의 동의를 받고 자발적으로 캠페인 참여를 유도합니다. "
                "고객 사전승낙과 보험 1차 동의콜로 생성한 DB만 구매할 수 있습니다. "
                "개인정보동의 녹취본을 전량 보관하고 문제 발생 시 소명자료를 제공합니다. "
                "1업체당 1DB 제공 원칙으로 중복 납품되지 않고, "
                "한 번 제공된 DB는 폐기하여 다시 판매하지 않습니다."
            ),
            "https://permission.example/insurance-db",
            "unknown",
            "intent",
        )
        self.assertEqual(
            permission_service,
            "excluded_consent_based_db_marketplace",
        )

        formal_insurance_product = relevance_gate_reason(
            "모딩 | 보험DB 전문 | 검증된 고객 DB 영업 솔루션",
            (
                "프리미엄 방문확정 상품 구매 시 장기부재 및 단박거절 A/S를 "
                "제공합니다. DB를 배분받은 다음 날까지 TA를 진행하고 총 3회 "
                "연락을 시도해야 합니다. DB 구매 시 TA 멘트와 반론 멘트 "
                "스크립트를 함께 제공하며 상담 관련 통화를 마친 고객입니다."
            ),
            "https://formal-leads.example/",
            "unknown",
            "intent",
        )
        self.assertEqual(
            formal_insurance_product,
            "excluded_consent_based_db_marketplace",
        )

        signed_application_product = relevance_gate_reason(
            "대면 100% 만남보장DB",
            (
                "보장분석 신청고객DB 구매하기. 대면영업으로 보장분석을 "
                "안내하고 신청서에 본인이 자필작성 및 서명한 고객정보를 "
                "제공합니다. 고객의 거부나 부재로 만남이 불가능하면 "
                "증빙자료 확인 후 새 고객DB로 교환하며, 100건 이상 주문 시 "
                "가격 협의가 가능합니다."
            ),
            "https://formal-leads.example/product/meeting-db",
            "unknown",
            "intent",
        )
        self.assertEqual(
            signed_application_product,
            "excluded_consent_based_db_marketplace",
        )

        permission_storefront = relevance_gate_reason(
            "퍼미션 메이커 보험 DB 라인업",
            (
                "기본 2차 퍼미션 DB 정가 140,000원 할인 적용가 130,000원. "
                "Perfect A/S 60% 2차 퍼미션 DB, 실버 2차 퍼미션 DB, "
                "가성비 퍼미션 DB를 판매하며 디비 신청 및 문의를 받습니다."
            ),
            "https://formal-leads.example/store",
            "unknown",
            "intent",
        )
        self.assertEqual(
            permission_storefront,
            "excluded_consent_based_db_marketplace",
        )

    def test_inbound_ad_lead_generation_is_not_existing_db_distribution(self) -> None:
        formal = relevance_gate_reason(
            "주식 DB 마케팅",
            (
                "광고대행사로서 광고 매체와 이벤트 페이지를 운영합니다. "
                "목표 타겟이 광고주가 원하는 행동을 할 때만 광고비를 지급하며, "
                "회원가입, 설문지 작성, 앱설치, 개인정보 입력으로 "
                "잠재고객의 개인정보를 수집합니다."
            ),
            "https://agency.example/stock-leads",
            "unknown",
            "intent",
        )
        seo = relevance_gate_reason(
            "유튜브DB 판매",
            (
                "원하시는 키워드 맞춤 세팅과 구글 상위노출 마케팅을 진행합니다. "
                "직접 검색해서 찾아오는 진성 DB 유입을 만드는 광고 서비스입니다."
            ),
            "https://seo.example/youtube-db",
            "unknown",
            "intent",
        )
        self.assertEqual(formal, "excluded_inbound_lead_generation_service")
        self.assertEqual(seo, "excluded_inbound_lead_generation_service")

        self_reported_leads = relevance_gate_reason(
            "정책자금 영업 DB의 중요성",
            (
                "허위 광고 및 불법 내용들이 많지만 저희는 정부자금 관련 "
                "유입자만 선별합니다. 조건별 타겟 필터링이 가능하고 "
                "중복·허위 제거 후 검수한 뒤 자영업자 또는 소상공인 본인이 "
                "직접 작성한 리드만 제공합니다."
            ),
            "https://blog.example/policy-fund-leads",
            "unknown",
            "intent",
        )
        search_inbound = relevance_gate_reason(
            "정책자금디비DB 상담률 높이는 법",
            (
                "포털 검색 유입 중심의 진성 고객이 정보를 남긴 그 순간 "
                "실시간 전달합니다. 한 고객 정보를 여러 곳에 동시 판매하는 "
                "짓은 하지 않고 필요한 파트너에게만 단독 공급합니다."
            ),
            "https://blog.example/search-inbound-leads",
            "unknown",
            "intent",
        )
        self.assertEqual(
            self_reported_leads,
            "excluded_inbound_lead_generation_service",
        )
        self.assertEqual(
            search_inbound,
            "excluded_inbound_lead_generation_service",
        )

    def test_structured_opt_in_lead_services_are_not_illicit_distribution(self) -> None:
        cases = (
            (
                "유튜브 주식DB 전문 실행사",
                (
                    "유튜브 채널을 직접 운영하며 영상 시청자가 문자로 직접 "
                    "신청한 주식DB를 수집·납품합니다. 자체 유튜브센터에서 "
                    "상담 신청과 랜딩페이지를 운영하고 실시간 DB를 공급합니다."
                ),
            ),
            (
                "퍼미션 DB·노리워드 CPA 디비 전문",
                (
                    "퍼미션 DB 가격문의, 퍼미션 보험 DB 단순동의콜 가격문의, "
                    "퍼미션 보험 DB 상담확정 가격문의, 만남확정 퍼미션 DB, "
                    "CPA DB 상담 신청 상품을 안내합니다."
                ),
            ),
            (
                "실시간대출DB 계약 전 체크리스트",
                (
                    "개인정보 수집·이용 및 제3자 제공 동의를 완료한 상담 신청 "
                    "리드만 전달합니다. 동의 시점과 로그, 유입 출처 및 녹취 "
                    "확인 절차를 계약서에 명시합니다."
                ),
            ),
            (
                "보험전문 DB 플랫폼",
                (
                    "상담원의 실제 통화 내용을 확인하세요. TM퍼미션 DB의 "
                    "생산 과정과 TM퍼미션 DB 운영 기준을 공개합니다."
                ),
            ),
            (
                "보험DB판매업체",
                (
                    "샘플 녹취콜은 상품페이지에서 확인합니다. 방문확정 상품, "
                    "방문픽스 상품, 100% AS 상품을 빠르게 배정합니다."
                ),
            ),
            (
                "보험DB 실시간 제공 시스템",
                (
                    "실제 보험상품에 관심을 보인 고객의 정보만 제공합니다. "
                    "보험에 관심 있는 고객의 정보를 실시간으로 수집하여 "
                    "등록 회원에게 즉시 전달합니다."
                ),
            ),
            (
                "보험 리모델링DB | 실시간 보험 리모델링 디비",
                (
                    "보험 점검 및 리모델링 상담 신청 정보를 기반으로 구성된 "
                    "데이터입니다. 실제 상담 의사가 반영된 신청 데이터를 "
                    "중심으로 정리하며 보험 상담 연결과 마케팅 활용에 "
                    "적합한 데이터 형태로 운영합니다."
                ),
            ),
            (
                "온라인보험DB 원천 공급",
                (
                    "원천 보험DB 공장을 운영하며 최적의 광고 콘텐츠와 "
                    "카피 라이팅, 랜딩 페이지로 신청 고객을 모집합니다."
                ),
            ),
            (
                "DB 판매 사이트 | 퍼미션 DB·노리워드 CPA 디비 전문",
                (
                    "리드 전환율을 극대화하는 완벽한 DB 공급. 광고비는 "
                    "줄이고 상담 전환율을 높입니다. 고객이 먼저 찾아오는 "
                    "고품질 리드를 확보하세요."
                ),
            ),
        )
        for index, (title, text) in enumerate(cases):
            with self.subTest(index=index):
                self.assertEqual(
                    relevance_gate_reason(
                        title,
                        text,
                        f"https://opt-in.example/products/{index}",
                        "unknown",
                        "intent",
                    ),
                    "excluded_inbound_lead_generation_service",
                )

        hacked_offer = relevance_gate_reason(
            "해킹 주식DB 판매",
            (
                "광고 랜딩 페이지 관리자에서 탈취한 개인정보 DB를 판매합니다. "
                "최신 주식DB 샘플과 단가는 텔레그램 [ACCOUNT]으로 문의하세요."
            ),
            "https://illicit.example/hacked-db",
            "unknown",
            "intent",
        )
        self.assertEqual(hacked_offer, "")

        direct_lending_db = relevance_gate_reason(
            "실시간 대출 DB 공급 직장인 사업자 무직자 맞춤 디비",
            (
                "직장인, 사업자, 무직자 대출DB를 실시간 공급합니다. "
                "저가 DB는 지양하며 조건별 단가로 제공합니다. 소량 테스트와 "
                "주문은 텔레그램 [ACCOUNT]으로 문의하세요."
            ),
            "https://market.example/lending-db",
            "unknown",
            "intent",
        )
        self.assertEqual(direct_lending_db, "")

    def test_cpa_vendor_request_is_inbound_lead_generation(self) -> None:
        reason = relevance_gate_reason(
            "상조·가전결합 CPA DB 공급 대행사 찾습니다",
            (
                "자체 매체를 보유한 CPA 제공 업체를 찾습니다. 언론사 광고와 "
                "제휴 회원 DB 수집 마케팅, 커뮤니티 기반 리드 수집을 진행합니다. "
                "상담 DB 수집 항목은 이름과 연락처이며, 수집/전달 방식과 "
                "광고 소재 및 랜딩페이지 심의 절차를 회신해 주세요."
            ),
            "https://agency.example/cpa-request",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_inbound_lead_generation_service")

    def test_010_handmade_nonreal_account_supply_is_relevant(self) -> None:
        reason = relevance_gate_reason(
            "010 네이버비실계 성예사 전문판매",
            (
                "마케팅·커뮤니티 운영에 최적화된 국내 010 기반 수작업 "
                "생성 계정을 안정적으로 공급합니다. 네이버와 성예사 계정을 "
                "원청 라인에서 직접 생성하며 소량부터 대량까지 납품합니다. "
                "계정 구매 문의는 카카오톡 [ACCOUNT] 또는 텔레그램 "
                "[MESSENGER_ID]으로 연락주세요."
            ),
            "https://market.example/handmade-accounts",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "")

    def test_bulk_google_catalog_is_not_a_single_game_account(self) -> None:
        reason = relevance_gate_reason(
            "구글 깡통 계정 · 유튜브 채널 판매",
            (
                "해외, 국내 깡통 계정 소량, 대량 구입 문의 환영. 정상적인 "
                "계정만 검수하여 판매합니다. 광고 운영, 유튜브 운영, "
                "플레이 앱스토어 게임계정 분리 등 여러 용도로 대량 보유하고 "
                "있으며 대량 구매 시 할인 공급합니다. "
                "첫 로그인 불가 계정은 교환 가능하고 텔레그램 [ACCOUNT]으로 "
                "24시간 상담합니다."
            ),
            "https://accounts.example/",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "")

    def test_attributable_inventory_survives_service_template_spam_rule(self) -> None:
        reason = relevance_gate_reason(
            "010인증 문의텔ﾩ RUN55 네이버아이디판매",
            (
                "네이버 아이디 다량 보유, 즉시 거래 가능합니다. 안정 계정 "
                "제공과 맞춤형 아이디 추천, 모든 SNS 연동, 문자 인증을 "
                "지원합니다. 계정 생성 및 인증 세팅 후 서비스 문의 및 주문은 "
                "텔ﾩ RUN55로 연락하세요. 예약 확정, 날짜와 인원, 방문 및 "
                "이용, 심야 시간대, 첫 방문 고객 같은 템플릿 문구가 섞여 "
                "있습니다."
            ),
            "https://board.example/account-inventory",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "")

    def test_structured_cpa_volume_procurement_is_inbound_generation(self) -> None:
        reason = relevance_gate_reason(
            "암보험 CPA 광고 대행사 모집 (월 1,000건 이상)",
            (
                "CPA 공급 가능한 대행사를 찾습니다. 테스트 50~100건 후 "
                "월 1,000건 단위 계약, DB 당 6만원입니다. 선호 채널은 "
                "유튜브와 타불라이며 랜딩은 협의된 것만 사용합니다. 내부 "
                "API 연동이 필요하고 광고 소재, 계약률, 공급 수량과 AS율을 "
                "회사 소개서에 적어 보내주세요."
            ),
            "https://agency.example/cpa-volume-request",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_inbound_lead_generation_service")

    def test_paid_media_lead_vendor_procurement_is_inbound_generation(self) -> None:
        reason = relevance_gate_reason(
            "개인회생 db 실행사를 찾고 있습니다",
            (
                "저희는 법무법인, 법률사무소를 대신해 인스타/페이스북에서 "
                "메타광고를 통해 개인회생db를 생산하고 있는 대행사입니다. "
                "DB 수급량이 부족하여 추가 공급 실행사를 찾고 있습니다. "
                "현재는 일 10~20개 정도가 필요하며 단가는 3만원대입니다. "
                "광고 소재는 저희가 전달 가능합니다. 구글/유튜브/틱톡/블로그 "
                "등 다른 매체에서 생성된 DB이면 좋겠습니다. 이메일로 일 가능 "
                "DB 수량, 단가, 공급매체를 알려주세요."
            ),
            "https://agency.example/personal-rehabilitation-db-request",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_inbound_lead_generation_service")

    def test_sparse_formal_insurance_db_platform_is_consent_marketplace(self) -> None:
        reason = relevance_gate_reason(
            "정심에셋 | 보험전문 DB플랫폼",
            (
                "HOME DB소개 TM퍼미션DB 100%만남 대면DB 구매하기 "
                "마케팅DB 상품 카테고리 Account Login Cart 마이페이지"
            ),
            "https://formal-leads.example/shop",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_consent_based_db_marketplace")

    def test_db_member_distribution_notice_is_not_a_product_post(self) -> None:
        reason = relevance_gate_reason(
            "마이어시스트 | 보험DB 판매 | 영업DB 전문 업계 1위",
            (
                "대량구매 OPEN 모일수록 더 커진 할인혜택! 카카오톡 배분알림 "
                "채널 추가는 필수입니다. DB를 보내드리고 있으며 알림 누락은 "
                "교환 사유가 아니니 채널 추가를 완료해 주세요."
            ),
            "https://formal-leads.example/member-club",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_empty_container")

    def test_permission_campaign_and_inbound_review_are_not_illicit(self) -> None:
        campaign = relevance_gate_reason(
            "코인 퍼미션 DB 의뢰드립니다",
            (
                "월 광고료 2억원을 집행 중이며 DB 단가와 AS 조건을 협의합니다. "
                "당일 진행 건은 저녁에 전수 녹취를 전달해 주세요."
            ),
            "https://agency.example/permission-request",
            "unknown",
            "intent",
        )
        review = relevance_gate_reason(
            "보험 인바운드 DB 이용 후기",
            (
                "상담 신청 고객 정보를 구매 후 빠르게 DB 배분받는 서비스입니다. "
                "인바운드 상품을 꾸준히 재구매하며 10건 중 20%는 대면 상담과 "
                "계약으로 연결된다는 보험 영업 후기입니다."
            ),
            "https://insurance.example/inbound-review",
            "unknown",
            "intent",
        )
        self.assertEqual(campaign, "excluded_inbound_lead_generation_service")
        self.assertEqual(review, "excluded_inbound_lead_generation_service")

    def test_formal_insurance_opt_in_db_products_are_not_illicit(self) -> None:
        marketplace = relevance_gate_reason(
            "100% 안심 퍼미션 DB",
            (
                "보험상담 동의 고객 DB 상품입니다. 판매 일정은 상시이며 "
                "구매 단위는 10개, 최소 구매 수량은 10개입니다."
            ),
            "https://insurance.example/permission-db",
            "unknown",
            "intent",
        )
        appointment = relevance_gate_reason(
            "TM 100% 만남보장DB 구매하기",
            (
                "보장분석 신청고객DB로 접수 확정 후 이름, 연락처와 생년월일을 "
                "확인합니다. 정보 제공 거부와 동의 철회가 가능하며 녹취 자료를 "
                "보관합니다."
            ),
            "https://insurance.example/appointment-db",
            "unknown",
            "intent",
        )
        self.assertEqual(marketplace, "excluded_consent_based_db_marketplace")
        self.assertEqual(appointment, "excluded_consent_based_db_marketplace")

        voluntary = relevance_gate_reason(
            "퍼미션 보험 DB 상담확정 퍼미션 DB 판매 사이트",
            (
                "광고 배너와 랜딩페이지를 통해 자발적으로 본인이 직접 신청한 "
                "타겟으로 구성한 진성 고객 DB 상품입니다. 최소 구매수량은 "
                "50건이며 배분 후 5영업일 이내 A/S 신청이 가능합니다. "
                "구매하기와 장바구니에서 주문해 주세요. 카톡 문의도 가능합니다."
            ),
            "https://formal-leads.example/voluntary-applicants",
            "unknown",
            "intent",
        )
        self.assertEqual(voluntary, "excluded_consent_based_db_marketplace")

        stolen = relevance_gate_reason(
            "해킹 퍼미션 보험 DB 판매",
            (
                "광고 랜딩페이지에서 본인이 직접 신청한 타겟 정보를 관리자 "
                "서버에서 탈취한 해킹 DB입니다. 최신 고객 DB 샘플과 단가는 "
                "텔레그램 seller123 문의 바랍니다."
            ),
            "https://seller.example/stolen-applicants",
            "unknown",
            "intent",
        )
        self.assertEqual(stolen, "")

        permission_product = relevance_gate_reason(
            "공모주 주식 퍼미션DB 제공",
            (
                "광고 수신 동의 여부가 포함된 퍼미션 기반 DB만 제공합니다. "
                "중복 제거와 검증을 완료해 법적 리스크 없는 합법적 활용과 "
                "투명한 정산 시스템을 보장합니다."
            ),
            "https://permission.example/ipo-db",
            "unknown",
            "intent",
        )
        self.assertEqual(
            permission_product,
            "excluded_consent_based_db_marketplace",
        )

        campaign_supply = relevance_gate_reason(
            "정책자금 상담 DB 공급 제휴",
            (
                "상담 유입 캠페인을 운영하며 유입 조건과 리드 승인 기준을 "
                "합의한 뒤 테스트 공급합니다. 수집·이용, 광고성 수신, "
                "제3자 제공 범위를 협의한 문구에 따라 운영합니다. CRM에는 "
                "텔레그램 알림을 연동할 수 있습니다."
            ),
            "https://agency.example/funding-leads",
            "unknown",
            "intent",
        )
        permission_pipeline = relevance_gate_reason(
            "퍼미션DB 전문 법인기업",
            (
                "직접 DB 수집을 실행하고 1차 DB를 선별한 뒤 2차 해피콜로 "
                "퍼미션DB를 완성해 공급합니다. 한번 사용된 DB는 즉시 폐기 "
                "처리합니다."
            ),
            "https://permission.example/pipeline",
            "unknown",
            "intent",
        )
        self.assertEqual(campaign_supply, "excluded_consent_based_db_marketplace")
        self.assertEqual(
            permission_pipeline,
            "excluded_consent_based_db_marketplace",
        )

        structured_listing = relevance_gate_reason(
            "보험디비구매 괜찮은 곳",
            (
                "상품: 보험 2차 퍼미션 DB. 나이: 30세 이상, 지역: 서울·경기, "
                "특징: 월 보험료 10만원 이상, AS: 주문 수량의 20% 추가, "
                "금액: 88,000원. 100개 구입 시 20개를 추가 지급합니다. "
                "카톡 문의 vendor9812"
            ),
            "https://qna.example/insurance-permission-db",
            "unknown",
            "intent",
        )
        self.assertEqual(
            structured_listing,
            "excluded_consent_based_db_marketplace",
        )

    def test_landing_page_case_study_is_db_purchase_alternative(self) -> None:
        reason = relevance_gate_reason(
            "보험설계사 전용 랜딩페이지와 디비포스 구축기",
            (
                "단발성 영업에서 벗어나 나만의 랜딩페이지와 세일즈 파이프라인을 "
                "구축했습니다. 외부 DB 구매 비용은 60% 절감하고 고객이 직접 "
                "상담을 신청하게 만드는 CRM 시스템을 제공합니다."
            ),
            "https://developer.example/case-study",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_db_purchase_alternative")

    def test_db_force_crm_is_management_software_not_db_sale(self) -> None:
        reason = relevance_gate_reason(
            "주식·코인 영업관리 프로그램 디비포스",
            (
                "투자 고객 DB를 관리하는 특화 CRM 솔루션입니다. 전체 DB 수와 "
                "VIP 전환율을 대시보드로 제공하고 중복 DB를 자동 필터링합니다. "
                "본 서비스는 DB 관리 프로그램이며 광고 DB 자체를 판매하지는 "
                "않습니다."
            ),
            "https://software.example/db-force",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_db_management_software")

    def test_permission_db_course_is_education_product(self) -> None:
        reason = relevance_gate_reason(
            "TA 기본 스크립트 정립 퍼미션DB 2강",
            (
                "DB 구매 금액이 부담되는 설계사를 위한 강좌입니다. 미리보기와 "
                "강좌 회차별 커리큘럼을 확인하고 퍼미션DB 상담 스크립트를 "
                "연습하는 교육 상품입니다."
            ),
            "https://course.example/permission-db",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_education_product")

        rebuttal_course = relevance_gate_reason(
            "TA마스터 - TA반론스크립트(퍼미션DB)",
            (
                "거절과 함께 날아간 DB 구매 비용을 회복하고 싶다면 "
                "TA 반론 스크립트와 적나라한 경험담으로 노하우를 쌓아보세요. "
                "교육 영상 39분, 구매 필요."
            ),
            "https://course.example/ta-master",
            "unknown",
            "intent",
        )
        self.assertEqual(rebuttal_course, "excluded_education_product")

    def test_empty_detail_shell_is_classified_without_dropping_real_body(self) -> None:
        related_rows = "\n".join(
            f"다른 게시글 제목 작성자 2026-08-{day:02d}" for day in range(10, 14)
        )
        empty_text = (
            "등록된 댓글이 없습니다.\n회원에게만 댓글 작성 권한이 있습니다.\n"
            + related_rows
        )
        self.assertEqual(
            classify_page_type(
                "https://shop.example/article/q-a/6/150885/",
                "개인통장 매입 문의",
                empty_text,
            ),
            "empty_container",
        )
        self.assertEqual(
            relevance_gate_reason(
                "개인통장 매입 문의",
                empty_text,
                "https://shop.example/article/q-a/6/150885/",
                "unknown",
                "intent",
            ),
            "excluded_empty_container",
        )

        real_text = (
            "개인통장 5개 매입합니다. 텔레그램 [ACCOUNT] 문의하세요.\n"
            "등록된 댓글이 없습니다.\n" + related_rows
        )
        self.assertNotEqual(
            classify_page_type(
                "https://shop.example/article/q-a/6/150886/",
                "개인통장 매입 문의",
                real_text,
            ),
            "empty_container",
        )

    def test_intent_gate_rejects_empty_sales_keyword_container(self) -> None:
        text = "跳至主要内容 博文 此处没有可显示的博文！"
        page_type = classify_page_type(
            "https://sales-keywords.example/",
            "네이버아이디판매 텔레그램 문의",
            text,
        )
        self.assertEqual(page_type, "empty_container")
        self.assertEqual(
            relevance_gate_reason(
                "네이버아이디판매 텔레그램 문의",
                text,
                "https://sales-keywords.example/",
                page_type,
                "intent",
            ),
            "excluded_page_type",
        )

    def test_board_index_is_not_retained_as_a_post(self) -> None:
        text = (
            "이미지형 리스트형 게시물 검색 제목 글쓴이 아이디 "
            "고객DB 판매 2026-08-01 10:01 "
            "계정 매입 2026-08-01 09:40 "
            "통장 대여 2026-08-01 08:30"
        )
        page_type = classify_page_type(
            "https://shop.example/board/gallery/8/",
            "갤러리 - 정상 쇼핑몰",
            text,
        )
        self.assertEqual(page_type, "board_listing")
        self.assertEqual(
            relevance_gate_reason(
                "갤러리 - 정상 쇼핑몰",
                text,
                "https://shop.example/board/gallery/8/",
                page_type,
                "intent",
            ),
            "excluded_page_type",
        )

    def test_naver_influencer_content_index_is_not_retained_as_post(self) -> None:
        text = (
            "04:38 웍 런칭 안내 조회수 5,853\n"
            "02:36 소테팬 재출시 조회수 7,016\n"
            "04:57 틱톡 계정을 판매합니다 조회수 3,478\n"
            "04:32 스텐팬 관리법 조회수 5,711"
        )
        page_type = classify_page_type(
            "https://in.naver.com/creator/contents/external/1",
            "네이버 인플루언서",
            text,
        )
        self.assertEqual(page_type, "board_listing")

    def test_intent_gate_rejects_corporate_transfer_with_incidental_account(self) -> None:
        reason = relevance_gate_reason(
            "법인양도양수",
            (
                "법인 매입합니다. 자본금은 상관없고 한도제한 없는 은행통장이 "
                "있어야 합니다. 법무사에서 대표자 변경 서류를 작성합니다."
            ),
            "https://community.example/company-transfer/1",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_corporate_transfer")

    def test_intent_gate_rejects_public_business_contact_directory(self) -> None:
        reason = relevance_gate_reason(
            "전국 학원 주소록 연락처 DB 제공합니다",
            (
                "포털 등록 학원 DB 20만 건을 판매합니다. 자료 내역은 "
                "업장명, 구주소, 신주소, 우편번호, 팩스, 홈페이지, "
                "업종이며 결제 후 엑셀로 제공합니다."
            ),
            "https://market.example/business-directory",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_public_business_directory")

        formal_directory = relevance_gate_reason(
            "KT고객관리솔루션-KT비즈콜",
            (
                "업종명 / 사업장명 / 전화번호 / 주소. 최신화된 실존 "
                "사업자 DB를 주기적으로 업데이트합니다. DB서포터들이 "
                "직접 구축/가공한 검증된 데이터를 제공해 합법적 사용이 "
                "가능합니다. 통합 DB관리 프로그램도 제공합니다."
            ),
            "https://directory.example/db.php",
            "unknown",
            "intent",
        )
        self.assertEqual(
            formal_directory,
            "excluded_public_business_directory",
        )

        live_business_directory = relevance_gate_reason(
            "신규 사업자 실시간 확인 비즈니스디비",
            (
                "매일 업데이트되는 신규·법인 사업자 DB입니다. 포털DB와 "
                "인허가DB, 마켓DB, 소셜DB를 모아 제공하며 기업DB와 법인DB도 "
                "등록 업체 수와 오늘 신규 건수를 확인할 수 있습니다."
            ),
            "https://directory.example/business-db",
            "unknown",
            "intent",
        )
        self.assertEqual(
            live_business_directory,
            "excluded_public_business_directory",
        )

        company_page = relevance_gate_reason(
            "신규 사업자 실시간 확인 비즈니스디비.-회사소개",
            (
                "매일 업데이트되는 신규 사업자 DB를 제공합니다. 법인·개인 "
                "영역을 구분하고 엑셀 형태로 언제든지 다운로드할 수 있습니다. "
                "세무서 사업자 등록일 바로 다음날부터 정보를 뽑아내는 "
                "영업 데이터베이스 구축 솔루션입니다."
            ),
            "https://www.businessdb.net/company.php",
            "unknown",
            "intent",
        )
        self.assertEqual(
            company_page,
            "excluded_public_business_directory",
        )

        smartstore_seller_directory = relevance_gate_reason(
            "실시간 스마트스토어 DB 웹프로그램 제공",
            (
                "스마트스토어 인터넷 판매자 DB를 웹 프로그램으로 "
                "제공합니다. 업체명, 연락처, 수집시간, 스토어 URL과 전 "
                "카테고리 업종을 실시간 수집하고 업데이트합니다."
            ),
            "https://market.example/smartstore-seller-db",
            "unknown",
            "intent",
        )
        self.assertEqual(
            smartstore_seller_directory,
            "excluded_public_business_directory",
        )

        online_seller_directory = relevance_gate_reason(
            "온라인 판매자 최신 리스트 DB 제공",
            (
                "대표자 이름, 대표자 개인 번호와 대표자 개인 이메일은 "
                "제공되지 않습니다. 최신 온라인 판매자 DB 리스트는 "
                "쇼핑몰명, 업체명, 업체 이메일 주소, 주소, URL로 구성되어 "
                "있으며 카테고리별 엑셀 자료로 제공합니다."
            ),
            "https://market.example/online-seller-list",
            "unknown",
            "intent",
        )
        self.assertEqual(
            online_seller_directory,
            "excluded_public_business_directory",
        )

    def test_intent_gate_rejects_seo_explainer_without_direct_offer(self) -> None:
        reason = relevance_gate_reason(
            "보험퍼미션DB와 실시간DB 안내",
            (
                "보험퍼미션DB는 상담 현장에서 자주 언급됩니다. 운영 방향을 "
                "잡는 데 도움이 될 수 있습니다. 확인 항목을 살펴보는 것이 "
                "좋습니다. 실시간DB 페이지에 함께 담으면 좋은 내용과 구성 "
                "포인트를 설명합니다. 상담 문의를 남겨주시면 안내합니다."
            ),
            "https://guide.example/insurance-db",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_seo_explainer")

    def test_intent_gate_rejects_account_creation_software(self) -> None:
        reason = relevance_gate_reason(
            "텔레그램 계정 생성프로그램",
            (
                "카카오톡 다중 접속 프로그램과 아이디 생성 프로그램을 "
                "판매합니다. 보유 연락처 DB 자동 친구 추가도 지원합니다."
            ),
            "https://software.example/account-generator",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_account_creation_tool")

    def test_intent_gate_rejects_keyword_phrase_dump(self) -> None:
        phrases = [
            "카톡아이디판매",
            "운전면허증제작",
            "해외카톡구매",
            "인스타해킹가격",
            "각종해킹의뢰",
            "모바일신분증제작",
            "DB판매",
            "계정매입",
        ] * 8
        reason = relevance_gate_reason(
            "텔레그램 [ACCOUNT] 인스타해킹의뢰",
            "\n".join(phrases),
            "https://normal-office.example/qna/spam",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_keyword_stuffing")

        fifty_line_dump = [
            "대출DB판매",
            "주식디비구매",
            "최신DB문의",
            "계정매입",
            "통장판매",
        ] * 10
        reason = relevance_gate_reason(
            "대출DB 최신디비문의 [ACCOUNT]",
            "\n".join(fifty_line_dump),
            "https://hijacked-board.example/qna/1",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_keyword_stuffing")

    def test_intent_gate_rejects_reservation_template_with_account_keywords(self) -> None:
        reason = relevance_gate_reason(
            "카카오 계정 매입 텔레그램 [ACCOUNT]",
            (
                "COMPLETE GUIDE 심야 시간대 예약 가능. "
                "예약은 언제 하는 게 좋나요? 원하시는 날짜와 인원을 알려주세요. "
                "심야에도 이용 가능하며 편안하게 즐기실 수 있습니다. "
                "단골 고객 이용 후기와 재방문 의사가 높습니다. "
                "지피티결제와 카카오 계정 매입 차이를 안내합니다."
            ),
            "https://shop.example/qna/spam",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_keyword_stuffing")

    def test_intent_gate_keeps_structured_telegram_trade_list(self) -> None:
        block = """네이버 계정 매입가능
단가 15000
거래 양식
아이디
비밀번호
성함
전화번호
계좌번호
신분증 앞면
문의 [ACCOUNT]"""
        reason = relevance_gate_reason(
            "커뮤니티 계정 매입 채널 – Telegram",
            "\n".join([block] * 12),
            "https://t.me/s/account_buyer",
            "public_messenger_page",
            "intent",
        )
        self.assertEqual(reason, "")

    def test_intent_gate_does_not_treat_transaction_form_example_as_article(self) -> None:
        reason = relevance_gate_reason(
            "네이버 계정 매입 – Telegram",
            (
                "네이버 계정 대여 비용은 당일 6000원 정산합니다. "
                "매입업무 마감 후 입금드립니다. 거래 양식 예시 홍길동 "
                "전화번호 계좌번호를 보내주세요. 카톡 raw_handle 문의"
            ),
            "https://t.me/s/account_rental",
            "public_messenger_page",
            "intent",
        )
        self.assertEqual(reason, "")

    def test_intent_gate_rejects_market_update_article(self) -> None:
        reason = relevance_gate_reason(
            "텔레그램 업데이트 소식",
            (
                "업데이트 공지를 알려준 적이 있죠. 텔레그램 핸들을 거래 "
                "가능하게 하는 기능도 더 강화됐습니다. 시장을 관심 가져보는 "
                "걸 추천하며 개인 브랜딩 시대를 시사하는 것 같습니다."
            ),
            "https://t.me/s/marketing-news",
            "public_messenger_page",
            "intent",
        )
        self.assertEqual(reason, "excluded_informational_article")

        overview = relevance_gate_reason(
            "개인정보 DB 판매의 현황과 문제점",
            (
                "개인정보가 불법적으로 거래되는 사례가 사회적인 문제로 "
                "대두되고 있습니다. 금융 사기로 이어질 수 있어 각별한 "
                "주의가 필요하며, 불법 DB 거래 단속을 강화해야 합니다."
            ),
            "https://design.example/db-portfolio/overview",
            "unknown",
            "intent",
        )
        self.assertEqual(overview, "excluded_informational_article")

        operations_guide = relevance_gate_reason(
            "보험DB 납품 시각이 AS 신청률을 바꿉니다",
            (
                "이 글은 공급사 관점에서 납품 시간을 정리한 것입니다. "
                "관련 안내서에 따르면 야간 전송은 별도 동의가 필요합니다. "
                "이 내용은 사실관계 전달이며 가정 시나리오로 계산합니다. "
                "AS가 발생하면 대체 DB를 납품합니다."
            ),
            "https://blog.example/db-delivery-guide",
            "unknown",
            "intent",
        )
        self.assertEqual(operations_guide, "excluded_informational_article")

    def test_intent_gate_rejects_trade_motive_question(self) -> None:
        reason = relevance_gate_reason(
            "커넥트 계정 사는 애들은 뭐야?",
            "30만원에 계정 매입한다는 사람들은 왜 매입하는 거임? 사기인가?",
            "https://community.example/post/motive-question",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_trade_motive_question")

    def test_intent_gate_rejects_legal_prop_and_security_guide(self) -> None:
        prop = relevance_gate_reason(
            "여권 제작 촬영소품",
            (
                "영화 촬영용 여권 소품이며 VOID와 SAMPLE을 표기하고 "
                "실제 개인정보는 사용하지 않습니다. 실제 발급과 위조는 "
                "제공하지 않습니다."
            ),
            "https://props.example/passport",
            "unknown",
            "intent",
        )
        security = relevance_gate_reason(
            "운전면허증 위조 방지 기술",
            "운전면허증 위조 방지를 위한 카드 보안 구조와 안전 설계 기준입니다.",
            "https://props.example/license-security",
            "unknown",
            "intent",
        )
        self.assertEqual(prop, "excluded_legal_prop_or_security_guide")
        self.assertEqual(security, "excluded_legal_prop_or_security_guide")

    def test_intent_gate_rejects_telegram_directory_and_normal_telecom_page(self) -> None:
        directory = relevance_gate_reason(
            "계좌매입 텔레그램 채널",
            (
                "엄선된 Telegram 채널, 그룹, 봇을 한 곳에서 모두 봅니다. "
                "구독자와 카테고리 순위, 최근 게시물 업데이트를 제공합니다."
            ),
            "https://directory.example/channel/account-buying",
            "unknown",
            "intent",
        )
        telecom = relevance_gate_reason(
            "인터넷디비 | 인터넷3사DB",
            (
                "인터넷 3사 비교상담으로 SK KT LG 요금제와 IPTV 결합, "
                "설치 가능 지역 및 약정 조건을 안내합니다."
            ),
            "https://isp.example/compare",
            "unknown",
            "intent",
        )
        self.assertEqual(directory, "excluded_telegram_directory")
        self.assertEqual(telecom, "excluded_normal_telecom_service")

    def test_intent_gate_rejects_unrelated_body_despite_sales_keyword_title(self) -> None:
        mismatch = relevance_gate_reason(
            "#여권제작 #여권위조",
            "가수의 공연과 팬들의 응원에 관한 오래된 일기입니다.",
            "https://blog.example/unrelated-post",
            "unknown",
            "intent",
        )
        title_only_listing = relevance_gate_reason(
            "모든 인스타 계정 매입합니다",
            "카카오톡 오픈채팅 문의 https://open.kakao.com/o/example",
            "https://market.example/account-post",
            "unknown",
            "intent",
        )
        self.assertEqual(mismatch, "excluded_title_body_mismatch")
        self.assertEqual(title_only_listing, "")

        testimonial_event = relevance_gate_reason(
            "보험DB 판매 | 영업DB 전문",
            (
                "후기 영상을 보내주시면 파일 확인 후 M캐시 50,000원을 "
                "지급합니다. 이벤트로 접수한 영상은 마케팅 용도로 사용되며, "
                "촬영 규격에 맞지 않으면 재촬영을 요청할 수 있습니다."
            ),
            "https://db-market.example/testimonial-event",
            "unknown",
            "intent",
        )
        self.assertEqual(testimonial_event, "excluded_title_body_mismatch")

        login_gate = relevance_gate_reason(
            "마이어시스트 | 보험DB 판매 | 영업DB 전문 업계 1위",
            (
                "로그인 해 주세요! 체계적인 케어와 사후관리를 위해 "
                "구매 전 회원가입이 필요합니다. 적립금, 배분알림, "
                "고객지원까지 다양한 회원전용 서비스를 경험해보세요."
            ),
            "https://myassist.kr/login?back_url=product",
            "unknown",
            "intent",
        )
        self.assertEqual(login_gate, "excluded_title_body_mismatch")

    def test_offer_title_with_only_corporate_footer_is_rejected(self) -> None:
        reason = relevance_gate_reason(
            "본인인증 계정 삽니다",
            (
                "주식회사 메타플랫폼 | 대표 이승준 사업자번호 713-81-03460 "
                "통신판매업 2024-서울강남-05468 개인정보보호책임자 홍길동 "
                "contact@example.com 무단 도용 시 법적 불이익을 받습니다. "
                "Copyright © Example. All rights reserved."
            ),
            "https://forum.example/empty-post",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_title_body_mismatch")

    def test_past_coerced_account_sale_recovery_request_is_rejected(self) -> None:
        reason = relevance_gate_reason(
            "카카오톡 계정 판매",
            (
                "카카오톡 계정을 2022년에 협박 때문에 팔았는데 카카오페이와 "
                "연동된 계정이라 꼭 찾고 싶습니다. 신고든 다 할 수 있으니 "
                "계정 찾는 법만 알려주세요."
            ),
            "https://qna.example/account-recovery",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_account_recovery_context")

    def test_account_suspension_plan_document_is_not_account_trade(self) -> None:
        reason = relevance_gate_reason(
            "쿠팡계정정지, 쿠팡판매중지 해결방안 노하우",
            (
                "쿠팡 계정정지 해제를 위한 개선계획서 작성 자료입니다. "
                "직접 작성해 통과된 파일과 증빙자료를 제공하며 구매 후 "
                "자료문의 게시판에 글을 남겨주세요."
            ),
            "https://documents.example/coupang-plan",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_normal_product_context")

    def test_forged_document_offer_is_not_hidden_by_gift_card_shop_rule(self) -> None:
        reason = relevance_gate_reason(
            "민증제작 면허증위조 상품권제작",
            (
                "민증제작, 면허증제작, 백화점상품권제작, 상품권 영수증 "
                "제작과 가족관계증명서 작업을 받습니다. 완성된 위조 문서는 "
                "등기우편으로 발송하며 문의는 텔레그램 [ACCOUNT]입니다."
            ),
            "https://board.example/forged-documents/1",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "")

    def test_intent_gate_rejects_host_io_domain_metadata(self) -> None:
        reason = relevance_gate_reason(
            "ggidmc.com — Google 아이디 판매",
            (
                "Domain Name: ggidmc.com. Registrar and hosting metadata, "
                "DNS records, IP address, and website technology profile."
            ),
            "https://host.io/ggidmc.com",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_domain")

        normal_brand = relevance_gate_reason(
            "대량 주문 문의",
            "여행 가방과 백팩의 수량, 각인, 배송 일정을 문의합니다.",
            "https://www.tumi.co.kr/bulk-order-inquiry",
            "unknown",
            "intent",
        )
        self.assertEqual(normal_brand, "excluded_domain")

    def test_intent_gate_rejects_account_purchase_legality_question(self) -> None:
        reason = relevance_gate_reason(
            "비실명 계정 구매가 합법인가요?",
            (
                "네이버 비실명 계정을 구매하려고 하는데 법적으로 문제가 "
                "되는 건가요? 불법인지 궁금해서 질문드립니다."
            ),
            "https://community.example/legal-question",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_question_or_guide")

    def test_intent_gate_rejects_account_purchase_risk_discussion(self) -> None:
        reason = relevance_gate_reason(
            "블로그계정 구매하면 문제가 생길까요?",
            (
                "비실계 구매해서 대량으로 뿌리고 인증된 계정 몇 개 더 "
                "구매할까 생각 중입니다. 계정구매하면 법적 문제가 생기거나 "
                "계정이 쉽게 막힐지도 걱정이고 적정 단가도 모르겠습니다."
            ),
            "https://community.example/account-risk-question",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_question_or_guide")

    def test_intent_gate_rejects_account_seller_scam_warning(self) -> None:
        reason = relevance_gate_reason(
            "실명 아이디 판매 먹튀합니다. 조심하세요",
            (
                "아이디 판매글을 올리고 송금하면 사용 불가능한 아이디를 "
                "넘긴 뒤 AS를 미루다가 잠수 타는 사람입니다. 조심하세요."
            ),
            "https://community.example/account-scam-warning",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_trade_warning")

    def test_intent_gate_rejects_commentary_quoting_account_inventory(self) -> None:
        reason = relevance_gate_reason(
            "요새 작업용 계정들이 부쩍 많이 보입니다",
            (
                "검색해보면 커뮤니티 침투용 바이럴 판매 계정이더군요. "
                "안녕하세요, 3차 최적화 아이디 팝니다. 맘카페 침투용 "
                "성인 여성 아이디도 팝니다. 대량 소량 판매합니다. "
                "해당 아이디는 50년 징계 당했습니다."
            ),
            "https://community.example/account-commentary",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_aggregation_or_commentary")

    def test_intent_gate_rejects_enterprise_ai_account_deployment(self) -> None:
        reason = relevance_gate_reason(
            "공공기관 AI도입, 계정 대량 구매로 챗GPT 도입",
            (
                "공공기관과 공기업을 위한 AI 도입을 지원합니다. "
                "복잡한 해외 결제를 대신 처리하고 계정 구매 대행부터 "
                "SW 설치, 환경 세팅, 맞춤 교육, 사후 관리까지 원스톱으로 "
                "제공합니다. 기관별 보안 컨설팅과 교육 프로그램도 설계합니다."
            ),
            "https://consulting.example/public-ai-deployment",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_enterprise_account_deployment")

    def test_intent_gate_keeps_storefront_with_legal_safety_disclaimer(self) -> None:
        reason = relevance_gate_reason(
            "구글 아이디 판매 해외 국내 깡통 계정 전문",
            (
                "구글 계정 판매 전문 업체입니다. 정상 실사용 Google 계정을 대량 "
                "보유하고 결제 후 즉시 발송합니다. 구글 계정 판매와 구매 "
                "과정에서 법적 문제를 피하기 위해 관련 규정을 준수해야 "
                "합니다. 텔레그램 [MESSENGER_ID]으로 주문 문의하세요."
            ),
            "https://store.example/",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "")

    def test_intent_gate_rejects_target_found_only_in_related_listings(self) -> None:
        reason = relevance_gate_reason(
            "네이버클립 노출 의뢰합니다",
            (
                "네이버클립 영상 노출과 바이럴 마케팅 작업을 의뢰합니다. "
                "자유홍보/랜덤광고 네이버 비실명 계정 100개 판매합니다. "
                "텔레그램 [ACCOUNT] 문의"
            ),
            "https://market.example/clip-promotion",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_title_body_mismatch")

    def test_previous_next_navigation_does_not_define_current_post(self) -> None:
        reason = relevance_gate_reason(
            "토토사이트 신규가입 이벤트",
            (
                "신규 가입 시 보너스를 지급하는 스포츠 이벤트 광고입니다.\n"
                "삭제하려면 비밀번호를 입력하세요.\n이전\n"
                "개인통장 삽니다 텔레그램 [ACCOUNT] 문의\n다음\n"
                "통장 대여 전문 업체\n관련 글 보기\n게시물 목록"
            ),
            "https://shop.example/article/qna/6/20",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_title_body_mismatch")

    def test_intent_gate_keeps_real_offer_before_related_listings(self) -> None:
        reason = relevance_gate_reason(
            "네이버 비실명 계정 판매합니다",
            (
                "네이버 비실명 계정 100개 판매합니다. 개당 5천원이며 "
                "텔레그램 [ACCOUNT] 문의 바랍니다. 관련 글 보기 "
                "인스타 계정 판매 게시글"
            ),
            "https://market.example/account-offer",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "")

    def test_generic_account_offer_survives_navigation_and_scam_disclaimer(self) -> None:
        reason = relevance_gate_reason(
            "회원 가입 시 현금 지급합니다",
            (
                "대량 문자 발송 사이트 가입 후 아이디 매입합니다. "
                "계정당 7만원 매입하며 텔레그램 [ACCOUNT] 문의. "
                "유심 업자 단가 맞춰드립니다. 사칭주의\n이전\n"
                "네이버 비실명 계정 판매\n다음\n통장 판매 업체\n"
                "관련 글 보기\n게시물 목록"
            ),
            "https://community.example/post/account-buying",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "")

    def test_intent_gate_keeps_blog_account_storefront_offer(self) -> None:
        reason = relevance_gate_reason(
            "2월 구정 한정이벤트 진행중입니다",
            (
                "원청 블로그 수급사 최블 준최 엔비 보유 다수\n"
                "블로그 임대 매매\n단체전환 계정 (영구매매)\n"
                "비실명 영구계정 매매\n단가 조절 가능 (문의)\n"
                "카카오톡 : admk9988\n오픈카톡 : "
                "https://open.kakao.com/o/sZJEeX9h"
            ),
            "https://www.marketslab.net/boards/view.php?id=4686",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "")

    def test_intent_gate_keeps_optimized_id_storefront_offer(self) -> None:
        reason = relevance_gate_reason(
            "준최판매 NB판매 최적 알선 임대 최적화카페 전문",
            (
                "비실명 영구매매 단체계정 전환가능. 최블 임대 6개월 1년. "
                "N사 카페 최적 카페 판매, 최적 아이디 판매. "
                "010 국내 비실명. 문의환영 카톡 com57"
            ),
            "https://market.example/optimized-id-offer",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "")

    def test_intent_gate_keeps_optimized_blog_transfer_offer(self) -> None:
        reason = relevance_gate_reason(
            "최적화블로그 알선 임대",
            (
                "최적화블로그 매매(명의변경, 단체전환), "
                "최적화블로그 비실명. 카톡 dagn33. "
                "최적화블로그 건바이 80,000원이며 키워드마다 가격 상이."
            ),
            "https://market.example/optimized-blog-transfer",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "")

    def test_intent_gate_keeps_duplicated_legacy_blog_listing(self) -> None:
        block = [
            "최적화 블로그 임대",
            "6개월",
            "1년",
            "영구임대",
            "최블 알선 가능",
            "주제 선택 가능",
            "직거래 가능",
            "최블 판매합니다",
            "카카오톡 : qkffks56",
            "세금 계산서 가능",
        ]
        text = "\n".join(block * 5)
        reason = relevance_gate_reason(
            "최블 임대 알선 판매합니다",
            text,
            "https://legacy-board.example/blog/1",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "")

    def test_intent_gate_keeps_new_smartstore_shell_purchase_request(self) -> None:
        reason = relevance_gate_reason(
            "신규 개설한 스마트스토어를 양도 받고 싶습니다",
            (
                "구매 희망금액 3,000,000원. 프로젝트 등록일 2026. 3. 17. "
                "거래 희망 기간 6개월내. 이 프로젝트에 참여하고 싶으시다면 "
                "지금 바로 매물을 등록해 주세요."
            ),
            "https://sellernet.kr/project/example",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "")
        self.assertEqual(
            infer_collection_type(
                "포털ID",
                "신규 개설한 스마트스토어를 양도 받고 싶습니다",
                "구매 희망금액 3,000,000원",
            ),
            "계정·아이디·가입인증",
        )

    def test_intent_gate_keeps_direct_smartstore_transfer_offer(self) -> None:
        reason = relevance_gate_reason(
            "스마트스토어양도, 스마트스토어판매, 스마트스토어매매",
            (
                "상세 설명. 스마트스토어 양도 판매합니다. "
                "가격은 항시 협의 가능합니다 문의주세요."
            ),
            "https://owners.example/onbiz/smartstore-transfer",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "")

    def test_intent_gate_rejects_whole_shopping_site_transfer(self) -> None:
        reason = relevance_gate_reason(
            "2,30대 여성의류 쇼핑몰 > 사이트매매",
            (
                "2021년 6월 오픈한 신규쇼핑몰입니다. 네이버페이 연동, "
                "카페24 기반으로 직접 제작한 사이트이며 도메인 연동 "
                "되어있습니다. 판매중인 재고 함께 양도합니다. 매매사유 "
                "및 매매금액. 인수절차는 사업자등록증과 사업자명의통장사본 "
                "준비 후 매매계약서를 작성하고 쇼핑몰계정 명의이전, 사이트 "
                "운영 방법 교육, 잔금 결제로 진행합니다."
            ),
            "https://market.example/sites/fashion-shopping-mall",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_normal_site_business_transfer")

    def test_intent_gate_rejects_smartstore_fee_analysis(self) -> None:
        reason = relevance_gate_reason(
            "네이버 스마트스토어 수수료·광고 구조 완벽 분석 및 2025년 전망",
            (
                "스마트스토어 판매 수수료 구조는 카테고리와 결제 방식에 "
                "따라 달라집니다. 정산 및 광고 구조와 운영 전략을 상세히 "
                "분석하고 판매자가 확인할 최신 정책을 안내합니다."
            ),
            "https://guide.example/smartstore-fees",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_smartstore_operation_guide")

    def test_intent_gate_rejects_normal_blog_account_guide(self) -> None:
        reason = relevance_gate_reason(
            "블로그 계정 만드는 법",
            (
                "블로그 계정을 개설하고 프로필과 공개 범위를 설정하는 "
                "방법을 단계별로 안내합니다. 공식 고객센터를 확인하세요."
            ),
            "https://guide.example/blog-account",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_normal_product_context")

    def test_intent_gate_rejects_account_as_job_requirement(self) -> None:
        reason = relevance_gate_reason(
            "쓰레드 복붙 알바",
            (
                "제공해드린 원고 복사 붙여넣기만 하면 됩니다. "
                "제공내역은 건당 300원이며 지원조건은 쓰레드 "
                "신규계정도 상관없음. 연락처 카톡 lovelybbn"
            ),
            (
                "http://www.ww.selfmoa.com/bbs/board.php?"
                "bo_table=selfmoawongo&wr_id=14088"
            ),
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_account_job_context")

    def test_intent_gate_rejects_thin_profile_offer_stub(self) -> None:
        stub = relevance_gate_reason(
            "네이버 아이디 판매",
            (
                "네이버 아이디 판매 네이버 메일은 무료 이메일 서비스이며, "
                "사용자에게 무료 이메일 주소와 메일 송수신 관리 기능을 제공합니다."
            ),
            "https://profile.example/seller",
            "unknown",
            "intent",
        )
        real = relevance_gate_reason(
            "네이버 아이디 판매",
            "네이버 계정 10개 판매합니다. 개당 5천원, 텔레그램 [ACCOUNT] 문의.",
            "https://profile.example/real-seller",
            "unknown",
            "intent",
        )
        self.assertEqual(stub, "excluded_empty_container")
        self.assertEqual(real, "")

    def test_embedded_news_byline_is_reporting_context(self) -> None:
        reason = relevance_gate_reason(
            "블로그 아이디 매입? 임대?",
            (
                '"내가 판매자라고?" 네이버 계정 도용이 뭐지 '
                "[아시아경제 조유진 기자] 자신의 계정이 도용된 사건입니다. "
                "news.naver.com 기사 링크"
            ),
            "https://blog.example/repost/1",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_reporting_context")

    def test_past_offense_image_repost_is_reporting_context(self) -> None:
        reason = relevance_gate_reason(
            "27년간 의사 면허증 위조한 60대.jpg - 유저업로드",
            (
                "유저업로드 목록 아래로 위로 쓰기 검색 복사 신고 스크랩 "
                "좋아요 댓글 쓰기 주간 조회수 인기글"
            ),
            "https://community.example/issue/2617215",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_reporting_context")

    def test_unrelated_article_rejects_syndicated_illicit_forum_preview(self) -> None:
        reason = relevance_gate_reason(
            "Bundesliga market values: Bayern player drops again",
            (
                "The striker has been one of the biggest stars of the season. "
                "This article compares player values and performances in detail. "
                * 20
                + " Latest posts from the news forum "
                "우리는 위조 여권과 신분증을 판매합니다. "
                "가짜 한국 여권 구매는 이메일 [EMAIL]로 문의하십시오."
            ),
            "https://sports.example/view/news/423486",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_embedded_preview")

    def test_intent_gate_rejects_page_copywriting_seo_filler(self) -> None:
        reason = relevance_gate_reason(
            "자동차보험DB | 디비판매업체",
            (
                "이 안내형 페이지는 페이지 성격과 문단별 역할을 나누어 "
                "설명형 문장으로 구성합니다. 키워드를 반복 나열하지 않고 "
                "페이지 흐름이 자연스럽고 안정적으로 보이게 정리합니다."
            ),
            "https://seo.example/auto-insurance-db",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_seo_explainer")

    def test_intent_gate_rejects_ai_roleplay_plot_as_an_illicit_post(self) -> None:
        reason = relevance_gate_reason(
            "대포통장 판매 조직 - 제타",
            (
                "불법 도박사이트에 쓰이는 대포통장을 판매하는 조직. "
                "당신의 친누나이자 조직의 사장이라는 캐릭터 설정이다."
            ),
            "https://zeta-ai.io/ko/plots/example/profile",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_domain")

    def test_intent_gate_rejects_illicit_keyword_seo_agency_ad(self) -> None:
        reason = relevance_gate_reason(
            "수원꽃집 새아침식물원",
            (
                "네이버아이디판매 아이디 판매 유입 극대화의 핵심 비법. "
                "실행사의 상위 노출 유지력에 달려 있습니다. 최상단 고정이 "
                "안 되면 광고비를 받지 않겠습니다. 광고주 정보는 암호화하고 "
                "대행사의 키워드 노출 마케팅으로 구글 1페이지를 보장합니다."
            ),
            "https://m.sacgarden.net/board/gallery/read.html?no=2242&board_no=8",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_search_spam")

    def test_black_hat_keyword_placement_ad_is_not_account_trade(self) -> None:
        reason = relevance_gate_reason(
            "카카오톡계정구매 카톡계정판매",
            (
                "여러 플랫폼에 안정적인 광고 노출을 지원합니다. 검색 점유율 "
                "1위와 상단 고정, 무한 도배로 시장을 장악합니다. 광고주 정보는 "
                "암호화하며 블랙 마케팅 기술 개발사의 필터링 우회 기술로 영역 "
                "독점을 보장합니다. 네이버아이디판매 키워드를 홍보합니다."
            ),
            "https://cosmetics.example/board/qna/960",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_search_spam")

    def test_intent_gate_rejects_license_requirements_and_id_photo_tools(self) -> None:
        license_guide = relevance_gate_reason(
            "운전면허증에 대해 | 렌터카 안내",
            (
                "국제운전면허증 위조가 증가하여 발급국과 거주 증명을 "
                "확인합니다. 유효 기한과 제네바 조약 기준을 충족해야 "
                "차량 대여가 가능합니다."
            ),
            "https://rental.example/license.html",
            "unknown",
            "intent",
        )
        photo_tool = relevance_gate_reason(
            "온라인 신분증 사진 제작기",
            (
                "무료 신분증 사진 편집기로 여권과 면허증 사진을 "
                "자르고 배경, 조명, 크기를 조정해 만들어 보세요."
            ),
            "https://editor.example/id-photo-maker",
            "unknown",
            "intent",
        )
        passport_security = relevance_gate_reason(
            "여권 제작 규정",
            (
                "여권번호 천공은 모든 페이지를 관통하는 정상적인 보안 "
                "요소이며 위·변조 방지 기술입니다. 전자여권 제작 규정에 "
                "따라 발급 단계에서 고유번호로 천공 처리됩니다."
            ),
            "https://qna.example/passport-perforation",
            "unknown",
            "intent",
        )
        passport_issuance = relevance_gate_reason(
            "전역군인 여권 제작",
            (
                "구청에서 여권 재발급 시 전역예정증명서는 필요 없다고 "
                "했습니다. 전역 예정자도 10년짜리 여권을 발급받을 수 있나요?"
            ),
            "https://qna.example/passport-issuance",
            "unknown",
            "intent",
        )
        mock_passport = relevance_gate_reason(
            "모의 여권 제작",
            (
                "학교 수업 활동으로 모의 여권 도안을 A4에 출력해서 "
                "만들어 보려고 합니다."
            ),
            "https://qna.example/mock-passport",
            "unknown",
            "intent",
        )
        passport_copy = relevance_gate_reason(
            "여권 사본 제출",
            "프린트카페에서 여권을 스캔한 PDF를 출력해서 제출하면 되나요?",
            "https://qna.example/passport-copy",
            "unknown",
            "intent",
        )
        mobile_id_verification = relevance_gate_reason(
            "모바일 주민등록증 위조, 간단한 QR 검증으로 확인",
            (
                "모바일 주민등록증 검증앱에서 QR코드를 스캔하면 해당 "
                "신분증의 위·변조 여부를 빠르게 확인할 수 있습니다."
            ),
            "https://government.example/mobile-id-check",
            "unknown",
            "intent",
        )
        self.assertEqual(license_guide, "excluded_identity_document_guide")
        self.assertEqual(photo_tool, "excluded_identity_photo_guide")
        self.assertEqual(
            passport_security,
            "excluded_identity_document_guide",
        )
        self.assertEqual(passport_issuance, "excluded_identity_document_guide")
        self.assertEqual(mock_passport, "excluded_normal_product_context")
        self.assertEqual(passport_copy, "excluded_identity_document_guide")
        self.assertEqual(
            mobile_id_verification,
            "excluded_identity_document_guide",
        )

    def test_intent_gate_rejects_generic_article_with_sale_keyword_link(self) -> None:
        reason = relevance_gate_reason(
            "경제적인 비실명 ID로",
            (
                "온라인 입지를 확보하는 열쇠는 강력한 마케팅입니다. "
                "비실명 ID를 활용하는 것이 유리합니다. 이러한 ID는 "
                "소규모 회사에도 매력적인 옵션이며 다양한 접근을 가능하게 "
                "합니다. 네이버 아이디 판매 자세한 내용은 웹사이트를 참고하세요."
            ),
            "https://notes.example/article",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_seo_explainer")

    def test_intent_gate_rejects_news_aggregator_even_when_title_has_trade_terms(self) -> None:
        reason = relevance_gate_reason(
            "19년간 실종 한국인 행세…여권 위조 입국 집행유예",
            (
                "총 9개의 출처 보기. 여권을 위조해 입국한 중국인들에게 "
                "법원이 집행유예를 선고했다."
            ),
            "https://aagag.com/issue/?idx=1649889",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_press_domain")

    def test_intent_gate_rejects_foreign_news_repost_about_account_purchase(self) -> None:
        reason = relevance_gate_reason(
            "페이스북 인증 계정 구매 후 사기 악용",
            (
                "공안부에 따르면 범죄자들은 기존 인증 페이스북 계정을 "
                "구매한 후 여행사를 사칭해 피해자의 돈을 갈취했다. "
                "규제 기관은 사기 피해 예방 수칙을 당부했다."
            ),
            "https://www.vietnam.vn/ko/example-news",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_press_domain")

        domestic = relevance_gate_reason(
            "북한 외화벌이 돈세탁, 대포통장 매입 공고",
            (
                "한국일보 취재 결과 텔레그램에서 미국인 명의 계좌를 "
                "구한다는 공고를 분석했다. 재무부 제재와 국제 공조 필요성을 "
                "전문가가 설명했다."
            ),
            "https://www.hankookilbo.com/news/article/example",
            "unknown",
            "intent",
        )
        self.assertEqual(domestic, "excluded_press_domain")

        asiae = relevance_gate_reason(
            "미성년자 대상 신분증 위조 판매 실태",
            "위조 의뢰 가격과 판매 수법을 취재한 사회면 기사입니다.",
            "https://view.asiae.co.kr/article/2025010813495733917",
            "unknown",
            "intent",
        )
        self.assertEqual(asiae, "excluded_press_domain")

    def test_intent_gate_rejects_documentary_id_card_market_and_db_pc_job(self) -> None:
        documentary = relevance_gate_reason(
            "한국여권 위조해 불법체류하는 중국인들",
            (
                "보더 시큐리티 프로그램입니다. 시드니 공항 직원이 "
                "승객을 인터뷰하고 통역사를 통해 위조 여권을 구했다는 "
                "사실을 확인합니다."
            ),
            "https://forum.example/documentary",
            "unknown",
            "intent",
        )
        id_card_market = relevance_gate_reason(
            "2025 신분증 제작 업체 트렌드",
            (
                "신분증 제작 업체 시장 개요와 제조사를 비교합니다. "
                "스마트 카드, NFC, RFID 기술은 보안과 접근 제어를 "
                "강화하는 정상적인 기업용 제품입니다."
            ),
            "https://market.example/id-card-trend",
            "unknown",
            "intent",
        )
        db_pc_job = relevance_gate_reason(
            "송내DB PC 평일오전 알바님을 찾습니다",
            "DB PC카페 시급 10,320원, 월~금 매장관리 알바 구인 공고입니다.",
            "https://jobs.example/job/1",
            "unknown",
            "intent",
        )
        self.assertEqual(documentary, "excluded_reporting_context")
        self.assertEqual(id_card_market, "excluded_normal_product_context")
        self.assertEqual(db_pc_job, "excluded_db_job_context")

    def test_intent_gate_rejects_financial_research_db_feed(self) -> None:
        reason = relevance_gate_reason(
            "안회수 DB 이차전지/철강금속",
            (
                "기업명: LG에너지솔루션(시가총액: 93조) 보고서명: "
                "단일판매ㆍ공급계약체결 계약상대: Ford Motor Company "
                "공급지역: 유럽 매출대비: 10% 공시링크: https://example.com "
                "<안회수 DB 이차전지 Daily News> 철강 생산량과 출하량 소식"
            ),
            "https://t.me/s/financial_research_feed/1",
            "public_messenger_page",
            "intent",
        )
        self.assertEqual(reason, "excluded_financial_research_feed")

    def test_intent_gate_rejects_new_live_normal_domain_false_positives(self) -> None:
        cases = (
            (
                "상품권대리 구매사기, 대포계좌 피해 대금 환수를 위한 민·형사 법리 전략 - 법무법인 초원",
                (
                    "상품권 구매 사기 피해자의 대포통장 계좌를 추적하여 법원 기습 "
                    "가압류를 단행합니다. 형사 고소와 강제추심으로 피해 금액을 "
                    "환수하는 법률 서비스입니다."
                ),
                "https://www.chowonlaw.com/legal/recovery",
                "excluded_legal_service_context",
            ),
            (
                "BRANDNARA - 택배 매입 신청",
                (
                    "개인정보 수집 및 이용 동의 제1조 개인정보의 처리목적 회사는 "
                    "홈페이지 회원 가입 및 관리를 위해 처리합니다. 제2조 개인정보의 "
                    "처리 및 보유기간을 안내합니다. 제3조 개인정보의 제3자 제공과 "
                    "제4조 개인정보 처리업무의 위탁을 공개합니다."
                ),
                "https://brandnara.co.kr/buy/privacy",
                "excluded_privacy_policy_document",
            ),
            (
                "병원DB 효과적으로 사용하는 팁3가지",
                (
                    "전국병원리스트는 진료과목, 병원명, 주소, 전화번호, 개원일을 "
                    "제공합니다. 모든 데이터는 각 지자체, 공공기관, 세무서에서 "
                    "제공하는 DB이며 공개되지 않거나 사용자가 원치 않는 정보는 "
                    "절대 수집 및 제공하지 않습니다."
                ),
                "https://mediking.net/public-hospital-list",
                "excluded_public_business_directory",
            ),
            (
                "주민등록증 없는 세상을 꿈꾸는 사람들",
                (
                    "좌담 사회: 주민등록증과 지문날인 제도가 정보인권에 미치는 "
                    "영향을 이야기하겠습니다. 참석자: 개인정보 유출과 주민등록증 "
                    "위조가 늘어나는 사회적 원인을 먼저 살펴봐야 합니다."
                ),
                "https://digitaljustice.kr/archive/id-card-roundtable",
                "excluded_reporting_context",
            ),
        )
        for title, text, url, expected in cases:
            with self.subTest(url=url):
                self.assertEqual(
                    relevance_gate_reason(title, text, url, "unknown", "intent"),
                    expected,
                )

    def test_intent_gate_rejects_new_normal_service_false_positives(self) -> None:
        cases = (
            (
                "여권스탬프투어제작 관광책자제작",
                "여권 크기의 수첩을 실제 여권처럼 구성한 공공기관 전문 디자인입니다.",
                "https://design.example/tour-booklet",
                "excluded_normal_product_context",
            ),
            (
                "주식대여서비스란?",
                "보유 주식을 기관투자자에게 빌려주고 대여수수료를 받는 대차거래입니다.",
                "https://securities.example/stock-lending",
                "excluded_investment_trading_service",
            ),
            (
                "중고오토바이매입 전국 출장 상담",
                "모델, 연식, 주행거리, 사고차 상태를 확인하고 명의이전을 진행합니다.",
                "https://vehicle.example/buy",
                "excluded_normal_product_context",
            ),
            (
                "장안평중고차시장",
                (
                    "HOME > 차량상세보기. 차종 중형차, 연식 18년식, "
                    "배기량 2,000cc, 주행거리 50,000km. 판매자정보와 "
                    "성능점검 기록입니다. 신용할부 구비서류는 운전면허증 "
                    "사본, 자동이체통장 사본, 신용카드와 신분증입니다."
                ),
                "https://vehicle.example/inventory/1",
                "excluded_normal_product_context",
            ),
            (
                "맘카페 침투 전문 업체 구합니다",
                "홍보/마케팅 대행 전문업체는 오픈카톡으로 문의해주세요.",
                "https://marketing.example/request/1",
                "excluded_marketing_service_request",
            ),
            (
                "한국 사는 외국인 필수 앱 — 본인인증 장벽부터 가입 팁까지",
                (
                    "외국인등록증 ARC가 나오기 전에는 여러 앱 가입이 막힙니다. "
                    "본인인증 장벽을 먼저 읽고 카테고리별 필수 앱과 가입 꿀팁, "
                    "외국인이 자주 막히는 지점과 해법을 확인하세요."
                ),
                "https://guide.example/foreign-resident-apps",
                "excluded_question_or_guide",
            ),
            (
                "네이버 비실명 아이디 거래와 정보통신망법",
                "계정 거래가 침입죄에 해당하는지 법적 쟁점과 법령, 판매자·구매자 처벌을 설명합니다.",
                "https://law.example/account-trade",
                "excluded_question_or_guide",
            ),
        )
        for title, text, url, expected in cases:
            with self.subTest(url=url):
                self.assertEqual(
                    relevance_gate_reason(title, text, url, "unknown", "intent"),
                    expected,
                )

    def test_intent_gate_rejects_normal_employee_id_card_designer(self) -> None:
        reason = relevance_gate_reason(
            "무료 신분증 만들기, 직원 ID 카드 온라인 제작",
            (
                "AI 신분증 생성기로 직원, 학생, 방문자 ID 카드를 만드세요. "
                "전문 템플릿에서 로고를 업로드하고 PDF, PNG로 인쇄할 수 있습니다."
            ),
            "https://design.example/id-card-maker",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_normal_product_context")

    def test_intent_gate_rejects_known_normal_institutional_domains(self) -> None:
        for url in (
            "https://www.kbsec.com/service/stock-lending",
            "https://www.imfnsec.com/mobile/bank/bank03.jsp",
            "https://www.insungsavingsbank.co.kr/product/loan",
            "https://www.jobkorea.co.kr/Recruit/GI_Read_Comt_Ifrm?Gno=1",
            "https://securities.miraeasset.com/imf/200/imf703.do",
            "https://www.cafe24.com/commerce/channel/kakao.html",
            "https://www.seogu.gwangju.kr/passport",
        ):
            with self.subTest(url=url):
                self.assertEqual(
                    relevance_gate_reason(
                        "계정 판매 여권 제작 통장 대여",
                        "정상 서비스 안내입니다.",
                        url,
                        "unknown",
                        "intent",
                    ),
                    "excluded_domain",
                )

    def test_daum_view_host_is_classified_and_excluded_as_press(self) -> None:
        url = "https://v.daum.net/v/20260830000000000"
        title = '"여권 만들어줘" 위조는 안 된다는 AI'
        text = "생성형 AI로 위조 여권을 만들 수 있는지 실험한 기사입니다."
        self.assertEqual(
            classify_page_type(url, title, text),
            "news_or_education",
        )
        self.assertEqual(
            relevance_gate_reason(title, text, url, "unknown", "intent"),
            "excluded_press_domain",
        )

    def test_intent_gate_rejects_public_dataset_and_account_purchase_guide(self) -> None:
        public_dataset = relevance_gate_reason(
            "병의원 DB 신규 개원 의료기관 데이터",
            (
                "전국 병의원 리스트를 엑셀로 다운로드합니다. 병원명, "
                "주소, 진료과목, 개원일, 의사수가 포함된 공공데이터 기반 "
                "상품으로 개인정보 미포함 DB입니다."
            ),
            "https://directory.example/hospital-db",
            "unknown",
            "intent",
        )
        guide = relevance_gate_reason(
            "저렴한 가격으로 TikTok 계정을 구매하는 방법",
            (
                "TikTok 계정 구매처를 알아보세요. Fameswap과 Famebolt, "
                "틱톡 마켓플레이스의 판매용 계정 가격과 팔로워를 비교합니다."
            ),
            "https://guide.example/tiktok-account-for-sale",
            "unknown",
            "intent",
        )
        self.assertEqual(public_dataset, "excluded_public_business_directory")
        self.assertEqual(guide, "excluded_market_guide")

    def test_intent_gate_rejects_account_advertorial_and_subscription_sharing(self) -> None:
        advertorial = relevance_gate_reason(
            "구글 계정 판매 사이트 비교 및 추천 TOP 5",
            (
                "인기 있는 구글 계정 판매 사이트 5곳의 특징과 장단점, "
                "추천 이유를 정리합니다. 마지막에는 추천 사이트를 알려드립니다."
            ),
            "https://board.example/post/1",
            "unknown",
            "intent",
        )
        subscription = relevance_gate_reason(
            "유튜브 프리미엄 가족계정 1달 5000원",
            (
                "유튜브 프리미엄 가족그룹은 12개월에 한 번만 변경 가능합니다. "
                "아이디를 주시면 초대 이메일을 드리고 남은 기간 환불합니다."
            ),
            "https://community.example/post/2",
            "unknown",
            "intent",
        )
        telegram_premium = relevance_gate_reason(
            "텔레그램 프리미엄계정 대행 서비스",
            (
                "텔레그램 프리미엄 계정 결제 대행은 PG 수수료 2%부터 "
                "제공합니다. 프리미엄 활성화 후 채널 제한을 높이고 그룹 "
                "인원 유입과 프로그램 솔루션 임대를 지원합니다."
            ),
            "https://agency.example/telegram-premium",
            "unknown",
            "intent",
        )
        telegram_account_sale = relevance_gate_reason(
            "텔레그램 프리미엄 계정 판매",
            (
                "프리미엄 활성화된 텔레그램 계정을 대량 판매합니다. "
                "구매 문의는 텔레그램 [ACCOUNT]으로 연락하세요."
            ),
            "https://market.example/telegram-premium-account",
            "unknown",
            "intent",
        )
        self.assertEqual(advertorial, "excluded_market_guide")
        self.assertEqual(subscription, "excluded_normal_subscription_service")
        self.assertEqual(
            telegram_premium,
            "excluded_normal_subscription_service",
        )
        self.assertEqual(telegram_account_sale, "")

    def test_intent_gate_rejects_search_placement_ads_with_account_keywords(self) -> None:
        reason = relevance_gate_reason(
            "카카오톡계정판매 대포유심",
            (
                "계정 판매 검색 시 광고 노출이 가능합니다. 상단 고정과 상단 장악, "
                "광고주 정보 암호화, 실행사 다이렉트 소스로 구글 첫 페이지를 "
                "보장합니다. 상위 노출 키워드와 노출 마케팅 서비스를 제공합니다. "
                "카카오톡계정판매 네이버계정판매 구글계정판매 계정 구매 판매."
            ),
            "https://shop.example/article/6/1",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_search_spam")

    def test_dense_forgery_catalog_with_operator_contact_is_kept(self) -> None:
        catalog = " ".join(
            [
                "위조민증 위조신분증 위조주민등록증 위조운전면허증 "
                "위조여권 모든서류위조가능"
            ]
            * 6
        )
        reason = relevance_gate_reason(
            "여권위조제작 여권위조업체",
            (
                f"{catalog} 텔레그램 @document_seller 24시상담 "
                "모든작업 원본100% 고퀄로 당일진행합니다"
            ),
            "https://board.example/post/forgery-catalog",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "")

    def test_threat_research_forgery_demo_is_not_a_trade_post(self) -> None:
        reason = relevance_gate_reason(
            "OpenAI 이미지 생성기로 가능해진 위조 여권 제작 | Cato Networks",
            (
                "Cato CTRL 위협 연구. 1. 개요 2. 기술적 개요 3. 결론 "
                "2025 위협 보고서에 따르면 생성형 AI로 사이버 범죄의 "
                "진입 장벽이 낮아졌습니다. 아래에서는 위조 여권 제작 과정을 "
                "시연합니다. 조직은 피싱과 맬웨어, 문서 기반 공격에 대비하고 "
                "사기 탐지 메커니즘을 업데이트해야 합니다."
            ),
            "https://security.example/blog/fake-passport-research",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_reporting_context")

    def test_intent_gate_keeps_account_storefront_with_marketing_service_copy(self) -> None:
        reason = relevance_gate_reason(
            "구글아이디판매 계정판매전문CAJA",
            (
                "구글 아이디 판매, 구글 애드워즈계정 판매, 유튜브아이디판매. "
                "저희는 수년째 계정을 판매중이며 정상 계정이 아닐 경우 "
                "A/S와 사후관리를 제공합니다. 구글 검색 광고는 광고주가 "
                "키워드 입찰을 통해 광고 노출을 얻습니다. 구글 검색에서 "
                "광고주가 키워드를 고르고 구글 검색 광고를 운영합니다. "
                + "구글 광고 서비스 안내 " * 14
            ),
            "https://www.googleidcaja.com/",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "")

    def test_platform_account_storefront_is_not_a_game_account_listing(self) -> None:
        reason = relevance_gate_reason(
            "구글 계정·유튜브 채널 판매",
            (
                "저희는 정상 구글 계정 판매합니다. 대량 구매 문의는 "
                "카카오톡 [MESSENGER_ID]으로 주세요. 플레이 앱스토어에서 "
                "게임계정분리 용도로도 사용할 수 있습니다. 365일 상담과 "
                "A/S 지원을 제공합니다."
            ),
            "https://accounts.example/",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "")

    def test_isolated_browser_renderer_accepts_public_same_domain_text(self) -> None:
        payload = json.dumps(
            {
                "source": "browser",
                "finalUrl": "https://shop.example/products",
                "title": "구글 계정 판매",
                "content": "구글 계정을 대량 판매합니다. 카카오톡 문의 가능합니다.",
                "safetyFlags": [],
            },
            ensure_ascii=False,
        )
        completed = Mock(returncode=0, stdout=payload)
        with (
            patch(
                "collector.collect_candidates.is_public_http_url",
                return_value=(True, ""),
            ),
            patch("collector.collect_candidates.shutil.which", return_value="/bin/agbrowse"),
            patch("collector.collect_candidates.subprocess.run", return_value=completed) as run,
        ):
            title, text, final_url, reason = render_public_text(
                "https://shop.example/"
            )
        self.assertEqual(title, "구글 계정 판매")
        self.assertIn("대량 판매", text)
        self.assertEqual(final_url, "https://shop.example/products")
        self.assertEqual(reason, "")
        command = run.call_args.args[0]
        self.assertIn("isolated", command)
        self.assertNotIn("user", command)

    def test_intent_gate_keeps_account_storefront_faq_question(self) -> None:
        reason = relevance_gate_reason(
            "네이버 아이디 구매 비실명 계정 판매 전문 아이디피플",
            (
                "네이버 아이디 구매 하면 모든 서비스 이용이 가능한가요? "
                "실명인증이 필요한 서비스 외에는 이용 가능합니다. "
                "저희 네이버 아이디 판매 전문 업체에서는 비실명 아이디를 "
                "판매 합니다. 상담/판매시간은 10:00부터 24:00까지이며 "
                "최소 5개부터 판매합니다. 카카오톡 [MESSENGER_ID]로 문의하세요."
            ),
            "https://www.opensellnids.com/",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "")

    def test_intent_gate_rejects_generic_review_template_keyword_spam(self) -> None:
        reason = relevance_gate_reason(
            "대포통장구매 텔레 seller 토토통장즉시매입",
            (
                "직접 경험해보니 만족스러웠습니다. 이용 후기 모음과 이용 팁입니다. "
                "처음이신 분들도 괜찮고 인원이 많으실 경우 미리 말씀해주세요. "
                "특별한 날이라면 분위기에 맞게 안내합니다. 대포통장 구매와 "
                "통장 즉시매입을 문의하세요."
            ),
            "https://shop.example/forum/post/1",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_keyword_stuffing")

    def test_masked_hospitality_template_spam_is_still_rejected(self) -> None:
        title = mask_text(
            "010인증 |텔레 raw_handle| 네이버맘카페용아이디구매 "
            "토스계정임대 상품 Q&A"
        )
        text = mask_text(
            "010인증 문의 텔레그램 raw_handle · 광고방법.COM. 토스계정임대와 "
            "네이버맘카페용아이디구매를 "
            "선택하실 때는 가격뿐 아니라 만족도를 보세요. 직접 경험해보시면 "
            "왜 찾는지 알 수 있습니다. 이용 후기와 이용 팁을 안내합니다. "
            "궁금한 점을 미리 정리해 문의하세요. 처음이신 분들도 많이 "
            "찾아주십니다. 코스와 인원에 따라 예약 가격이 달라집니다."
        )
        self.assertEqual(
            relevance_gate_reason(
                title,
                text,
                "https://shop.example/article/q-a/6/1",
                "unknown",
                "intent",
            ),
            "excluded_keyword_stuffing",
        )

    def test_intent_gate_rejects_bank_freeze_advice_question(self) -> None:
        reason = relevance_gate_reason(
            "카뱅 지급정지 후 서면접수 당했습니다. 자문 구합니다.",
            (
                "제 계좌가 지급정지되었고 다른 은행은 비대면거래제한만 걸렸습니다. "
                "환수당해야 하는지에 대한 부분이 궁금한데 답답합니다."
            ),
            "https://community.example/post/2",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_question_or_guide")

    def test_intent_gate_rejects_pawnshop_derivatives_and_free_giveaway(self) -> None:
        pawnshop = relevance_gate_reason(
            "하이앤바이전당포대부",
            (
                "전문 감정사가 고객 물품의 시세를 감정한 후 매입 "
                "계약서를 작성하고 전당 또는 위탁판매를 진행합니다."
            ),
            "https://pawn.example/buy",
            "unknown",
            "intent",
        )
        derivatives = relevance_gate_reason(
            "국내선물 대여계좌",
            (
                "국내선물 실계좌 기반 대여계좌로 증거금 30만원에 "
                "코스피·환율 선물 거래를 제공합니다."
            ),
            "https://derivatives.example/rental-account",
            "unknown",
            "intent",
        )
        giveaway = relevance_gate_reason(
            "활성화아이디 무료나눔 이벤트",
            (
                "댓글 선착순 10명에게 페이스북 연동 로그인 방식 "
                "활성화 아이디 1계정 제공, 무료 나눔합니다."
            ),
            "https://community.example/free-account",
            "unknown",
            "intent",
        )
        self.assertEqual(pawnshop, "excluded_normal_product_context")
        self.assertEqual(derivatives, "excluded_investment_trading_service")
        self.assertEqual(giveaway, "excluded_noncommercial_giveaway")

    def test_xe_board_home_with_many_dated_rows_is_a_listing(self) -> None:
        text = "\n".join(
            f"- 08.{day:02d} 기타 고객 DB 관련 게시글 {day}"
            for day in range(20, 28)
        )
        self.assertEqual(
            classify_page_type(
                "https://forum.example/index.php?mid=board_MsdG13",
                "셀프모아",
                text,
            ),
            "board_listing",
        )

    def test_cafe24_board_category_with_dated_rows_is_a_listing(self) -> None:
        text = "\n".join(
            f"- [] 위조문서 제작 문의 작성자 2026-04-{day:02d} 조회 {day} 추천 0"
            for day in range(20, 25)
        )
        self.assertEqual(
            classify_page_type(
                "https://m.portal.example/board/%EC%9E%90%EB%A3%8C%EC%8B%A4/7",
                "자료실 - 경북포털",
                text,
            ),
            "board_listing",
        )

    def test_intent_gate_rejects_victim_and_regulatory_explainers(self) -> None:
        cases = (
            (
                "AI 연산력의 달콤한 함정",
                (
                    "법무법인 금융센터에서 사기 사건 피해자분들을 위해 "
                    "작성했습니다. 사칭 사기 조직이므로 주의 바랍니다."
                ),
                "excluded_reporting_context",
            ),
            (
                "편법DB영업 근절, 보험업계 무엇이 달라질까요?",
                (
                    "금융감독원은 개인정보 보호를 핵심 감독 과제로 선정하고 "
                    "소비자경보를 발령했으며 제도 개선 방침을 발표했습니다."
                ),
                "excluded_reporting_context",
            ),
            (
                "내 번호가 불법사채 리스트가 된 이유",
                (
                    "저희 센터를 찾는 피해자분들이 공통적으로 호소합니다. "
                    "DB 유통업자가 정보를 사고팝니다. 그 생태계를 상세히 분석합니다."
                ),
                "excluded_reporting_context",
            ),
        )
        for title, text, expected in cases:
            with self.subTest(title=title):
                self.assertEqual(
                    relevance_gate_reason(
                        title,
                        text,
                        "https://blog.example/post/1",
                        "unknown",
                        "intent",
                    ),
                    expected,
                )

    def test_intent_gate_rejects_fiction_procurement_and_crime_reposts(self) -> None:
        cases = (
            (
                "칼럼 조합이면 너무 위험하지 않음?",
                (
                    "윌키가 가짜신분증 제작하고 에반이 배포하는 팀이 "
                    "만들어지면 어캄. 생각만으로도 어질어질해."
                ),
                "excluded_fictional_context",
            ),
            (
                "아기주민등록증 발급 알아보기",
                (
                    "영주시가 출생을 축하하는 출생 축하증입니다. 실제 신분증처럼 "
                    "생겼지만 법적인 효력은 없는 기념용 증서입니다."
                ),
                "excluded_normal_product_context",
            ),
            (
                "모바일 운전면허증 제작용 보안카드 단가계약",
                (
                    "한국도로교통공단 발주 RF-PVC 보안카드 입찰 개찰 결과와 "
                    "업체별 투찰금액을 안내합니다."
                ),
                "excluded_public_procurement",
            ),
            (
                "19년간 실종 한국인 행세, 여권 위조 입국",
                "여권 위조 입국 후 불법 체류한 사건을 전합니다.",
                "excluded_reporting_context",
            ),
            (
                "캡틴아메리카 남성, 가짜 CIA 신분증 제작",
                "한 남성이 웹사이트에서 가짜 CIA 신분증을 제작한 사건입니다.",
                "excluded_reporting_context",
            ),
            (
                "AI 신분증 위조 92% 우회, 신원 검증 재점검 시급",
                (
                    "생성 AI 신분증 테스트에서 높은 우회율이 나타나 기업 신원 "
                    "검증 체계의 핵심 위험으로 제기됐습니다."
                ),
                "excluded_reporting_context",
            ),
            (
                "셀피글로벌 모바일 운전면허증 보안카드 수주",
                (
                    "계약금액은 매출 대비 8.72%입니다. 주가 영향 분석과 실적 "
                    "반영 가능성, 투자자 유의사항을 정리합니다."
                ),
                "excluded_financial_research_feed",
            ),
        )
        for title, text, expected in cases:
            with self.subTest(title=title):
                self.assertEqual(
                    relevance_gate_reason(
                        title,
                        text,
                        "https://community.example/post/1",
                        "unknown",
                        "intent",
                    ),
                    expected,
                )

    def test_intent_gate_rejects_stock_disclosure_channel_feed(self) -> None:
        reason = relevance_gate_reason(
            "실시간 주식 공시 정리채널",
            (
                "기업명: 탑코미디어(시가총액: 1,701억) 보고서명: "
                "교환청구권행사 청구주식 121,153주 공시링크: https://example.com "
                "회사정보: https://example.com/company"
            ),
            "https://t.me/s/disclosure_feed",
            "public_messenger_page",
            "intent",
        )
        self.assertEqual(reason, "excluded_financial_research_feed")

    def test_deleted_post_message_is_structurally_inaccessible(self) -> None:
        self.assertEqual(
            classify_page_type(
                "https://forum.example/board_read?idx=1",
                "자유게시판",
                "주민등록증 제작 문의 게시글 삭제 되었습니다. 목록보기",
            ),
            "deleted_or_inaccessible",
        )
    def test_revalidation_recomputes_page_type_from_current_review_text(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            masked_path = Path(temp_dir) / "candidates_masked.csv"
            review_path = Path(temp_dir) / "data.csv"
            audit_path = Path(temp_dir) / ".private" / "revalidation_removals.csv"
            listing_text = "\n".join(
                f"- 08.{day:02d} 기타 고객 DB 게시글 {day}"
                for day in range(20, 28)
            )
            masked_rows = []
            review_rows = []
            for sample_id, url, title, text in (
                (
                    "EG-000001",
                    "https://forum.example/index.php?mid=board_MsdG13",
                    "셀프모아",
                    listing_text,
                ),
                (
                    "EG-000002",
                    "https://other.example/post/2",
                    "네이버 계정 판매합니다",
                    "네이버 계정 판매합니다. 텔레그램 seller123 문의",
                ),
            ):
                masked = {name: "" for name in SCHEMA}
                masked.update(
                    {
                        "sample_id": sample_id,
                        "page_type": "unknown",
                        "masked_title": mask_text(title),
                        "masked_text": mask_text(text),
                    }
                )
                masked_rows.append(masked)
                review_rows.append(
                    {
                        "sample_id": sample_id,
                        "collected_at": "2026-08-29T00:00:00+09:00",
                        "source_url": url,
                        "registrable_domain": "example",
                        "title": title,
                        "text": text,
                    }
                )
            with masked_path.open("w", encoding="utf-8-sig", newline="") as f:
                writer = csv.DictWriter(f, fieldnames=SCHEMA)
                writer.writeheader()
                writer.writerows(masked_rows)
            with review_path.open("w", encoding="utf-8-sig", newline="") as f:
                writer = csv.DictWriter(f, fieldnames=RESTRICTED_REVIEW_SCHEMA)
                writer.writeheader()
                writer.writerows(review_rows)

            removed = revalidate_existing_records(
                masked_path,
                review_path,
                b"test-key",
                "intent",
                1,
                1,
                audit_path,
            )

            self.assertEqual(removed["excluded_page_type"], 1)
            with masked_path.open(encoding="utf-8-sig", newline="") as f:
                kept = list(csv.DictReader(f))
            self.assertEqual([row["sample_id"] for row in kept], ["EG-000002"])
            self.assertEqual(kept[0]["page_type"], "unknown")
            self.assertEqual(kept[0]["registrable_domain"], "other.example")
            with review_path.open(encoding="utf-8-sig", newline="") as f:
                kept_review = list(csv.DictReader(f))
            self.assertEqual(kept_review[0]["registrable_domain"], "other.example")
            with audit_path.open(encoding="utf-8-sig", newline="") as f:
                audit_rows = list(csv.DictReader(f))
            self.assertEqual(len(audit_rows), 1)
            self.assertEqual(audit_rows[0]["sample_id"], "EG-000001")
            self.assertEqual(audit_rows[0]["reason"], "excluded_page_type")
            self.assertEqual(
                audit_rows[0]["source_url"],
                "https://forum.example/index.php?mid=board_MsdG13",
            )

    def test_intent_gate_keeps_weak_price_signal_with_concrete_contact(self) -> None:
        reason = relevance_gate_reason(
            "고객 DB 안내",
            "고객 DB 건당 단가는 텔레그램 [MESSENGER_ID]로 문의주세요.",
            "https://board.example/db-price",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "")

    def test_intent_gate_keeps_db_needed_offer_with_raw_messenger_handle(self) -> None:
        reason = relevance_gate_reason(
            "실시간 대출 디비 전문 실행사입니다",
            (
                "실시간 대출DB 필요하신 대표님들이라면 책임지고 "
                "진행해드립니다. 텔레그램 thomnews 로 문의주세요."
            ),
            "https://marketing.example/db-service/1",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "")

    def test_intent_gate_keeps_messenger_account_sales_offer(self) -> None:
        reason = relevance_gate_reason(
            "국내외 텔레그램·카카오톡 계정 유통점",
            (
                "카톡계정과 텔레그램 계정을 한 달에 10000개 이상 거래 "
                "가능합니다. 계정 판매 취급점은 구매상담 텔레그램 "
                "@talkseller1 로 연락주세요."
            ),
            "https://t.me/s/account_shop",
            "public_messenger_page",
            "intent",
        )
        self.assertEqual(reason, "")

    def test_intent_gate_keeps_compound_offer_with_raw_messenger_url(self) -> None:
        reason = relevance_gate_reason(
            "카톡 실계정 매입·판매",
            (
                "카톡 실계정 실시간 매입, 당일 입금합니다. "
                "모바일톡과 피시톡 계정을 대량·소량 공급합니다. "
                "문의 https://t.me/raw_seller"
            ),
            "https://t.me/s/source_channel/1",
            "public_messenger_page",
            "intent",
        )
        self.assertEqual(reason, "")

    def test_intent_gate_keeps_phone_and_email_data_sales_offer(self) -> None:
        reason = relevance_gate_reason(
            "텔레그램 전화번호 데이터",
            (
                "휴대전화 번호 데이터와 이메일 데이터를 제공합니다. "
                "최신 연락처 데이터베이스 패키지 구매는 Telegram "
                "@data_seller 로 문의주세요."
            ),
            "https://data.example/telegram-numbers",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "")

    def test_intent_gate_keeps_db_purchase_inquiry(self) -> None:
        reason = relevance_gate_reason(
            "대출DB 구매 문의드립니다",
            (
                "대부업 대출DB 단가는 어느 정도인가요? 거절난 대부업DB는 "
                "얼마에 구매해야 맞는지 문의드립니다."
            ),
            "https://community.example/question/buy-db",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "")

        no_contact_inquiry = relevance_gate_reason(
            "대출DB 문의드립니다",
            (
                "일수 관련 대부업DB의 단가는 어느 정도인가요? 일반 "
                "신용대출보다 다 거절난 대부업DB는 얼마에 구매해야 "
                "맞는 건지 모르겠습니다."
            ),
            "https://community.example/question/db-price",
            "unknown",
            "intent",
        )
        self.assertEqual(no_contact_inquiry, "")

    def test_intent_gate_rejects_normal_foreign_exchange_account(self) -> None:
        reason = relevance_gate_reason(
            "씨티은행 인터넷 뱅킹",
            (
                "계좌간 환전 서비스는 외화통장에 외환을 매입해 두거나 "
                "보유 외환을 매각하는 정상 은행 서비스입니다."
            ),
            "https://bank.example/foreign-exchange",
            "unknown",
            "intent",
        )
        self.assertNotEqual(reason, "")

    def test_normal_account_opening_guide_is_not_bank_account_trade(self) -> None:
        reason = relevance_gate_reason(
            "비대면 예금계좌개설 준비서류와 본인인증 절차",
            (
                "은행 앱에서 계좌 개설 메뉴를 선택하고 본인 명의 휴대폰 인증, "
                "신분증 촬영, 타행 계좌 인증 순서로 진행하는 방법을 정리합니다."
            ),
            "https://finance.example/account-opening-guide",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_normal_product_context")

    def test_used_phone_reset_guide_is_not_account_trade(self) -> None:
        reason = relevance_gate_reason(
            (
                "쓰던 폰 중고로 넘기기 전, 초기화 버튼부터 누르면 안 "
                "됩니다 (계정·유심·SD카드 순서)"
            ),
            (
                "판매자 입장에서 안전하게 넘기는 순서를 정리했습니다. "
                "삼성 계정 로그아웃과 구글 계정 삭제를 먼저 합니다. "
                "유심과 eSIM을 분리하고 SD카드를 빼서 보관한 뒤 설정 "
                "메뉴에서 디바이스 전체 초기화를 실행합니다. 초기화는 "
                "맨 마지막입니다."
            ),
            "https://m.blog.naver.com/vidsavepro/224364516644",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_normal_product_context")

    def test_direct_account_offer_survives_used_phone_reset_rule(self) -> None:
        reason = relevance_gate_reason(
            "네이버 계정 대량 판매합니다",
            (
                "네이버 아이디와 계정 판매합니다. 중고폰을 팔기 전 계정 "
                "로그아웃과 공장 초기화를 하라는 안내가 아니라 계정 "
                "자체를 대량 판매하는 글입니다. 텔레그램 [ACCOUNT]으로 "
                "문의하세요."
            ),
            "https://seller.example/naver-accounts",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "")

    def test_account_freeze_remedy_guides_are_not_bank_account_trade(self) -> None:
        examples = (
            (
                "계좌지급정지 해제 사기의심계좌 통장 거래 제한 법적 해법은",
                (
                    "금융기관에 제출할 거래 내역과 증빙 자료를 정리하고 "
                    "수사기관 진술과 자금 이동 경위를 일치시켜야 합니다. "
                    "피해 신고와 연결된 입금인지 확인한 뒤 해제 절차를 "
                    "준비합니다."
                ),
            ),
            (
                "계좌지급정지 해제와 통장대여 전기통신금융사기 구제법",
                (
                    "은행 안내를 받은 뒤 입출금 내역과 객관적 자료를 "
                    "준비했습니다. 선의의 계좌 명의인은 사실관계를 정리해 "
                    "이의신청할 수 있다는 금융위원회 안내입니다."
                ),
            ),
        )
        for title, text in examples:
            with self.subTest(title=title):
                reason = relevance_gate_reason(
                    title,
                    text,
                    "https://m.blog.naver.com/legal-guide/123",
                    "unknown",
                    "intent",
                )
                self.assertEqual(reason, "excluded_legal_service_context")

    def test_direct_bank_offer_survives_account_freeze_remedy_rule(self) -> None:
        reason = relevance_gate_reason(
            "지급정지 걱정 없는 법인계좌 판매합니다",
            (
                "법인 통장과 계좌를 대량 판매합니다. 거래 내역 확인 후 "
                "즉시 정산하며 텔레그램 [ACCOUNT]으로 문의하세요."
            ),
            "https://seller.example/bank-accounts",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "")

    def test_kyc_account_buying_channel_is_not_excluded_as_game_trade(self) -> None:
        reason = relevance_gate_reason(
            "게임 매입방",
            (
                "KYC 인증 계정 삽니다. 딥코인 인증 개당 2만원 즉시 정산하며 "
                "신분증, 운전면허증, 여권 중 하나를 보유해야 진행 가능합니다. "
                "문의는 텔레그램 [ACCOUNT]으로 주세요."
            ),
            "https://t.me/s/example_kyc_room",
            "public_messenger_page",
            "intent",
        )
        self.assertEqual(reason, "")

    def test_intent_gate_rejects_additional_single_game_accounts(self) -> None:
        reason = relevance_gate_reason(
            "배틀그라운드 카카오 계정 판매",
            "제가 사용하던 계정을 10만원에 팝니다.",
            "https://market.example/post/7",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_single_account_trade")

        linked_game_account = relevance_gate_reason(
            "구글 계정 팝니다",
            "아이온 캐릭터 레벨 45인 계정 하나를 판매합니다.",
            "https://market.example/post/8",
            "unknown",
            "intent",
        )
        self.assertEqual(linked_game_account, "excluded_single_account_trade")

        league_account = relevance_gate_reason(
            "롤계정 대량 판매합니다",
            "리그오브레전드 계정을 판매합니다. 카톡 문의 가능합니다.",
            "https://game.example/lol-account",
            "unknown",
            "intent",
        )
        self.assertEqual(league_account, "excluded_single_account_trade")

        cookie_run_account = relevance_gate_reason(
            "카카오 쿠키런 계정 구매합니다",
            (
                "크리스탈 수급량이 높은 쿠키런 게임 계정을 구합니다. "
                "희망 가격과 계정 상세 정보를 보내주세요."
            ),
            "https://game.example/cookie-run-account",
            "unknown",
            "intent",
        )
        self.assertEqual(cookie_run_account, "excluded_single_account_trade")

        marketplace_listing = relevance_gate_reason(
            "틱라 장기미접 또는 본인인증 삽니다",
            (
                "구매수량 1개, 구매금액 2만원. 계정종류 게임사 로그인, "
                "캐릭터직업 기타, 레벨 111인 게임 계정을 구합니다."
            ),
            "https://game-market.example/buy/1",
            "unknown",
            "intent",
        )
        self.assertEqual(marketplace_listing, "excluded_single_account_trade")

        nexon_account = relevance_gate_reason(
            "넥슨 A/S 가능한 본인인증계정 판매합니다",
            "핵 이용자에게 계정을 공급하며 보호조치 해제를 해드립니다.",
            "https://game.example/maple-story/account",
            "unknown",
            "intent",
        )
        self.assertEqual(nexon_account, "excluded_single_account_trade")

    def test_ncsoft_accounts_for_personal_game_use_are_excluded(self) -> None:
        reason = relevance_gate_reason(
            "엔씨 NCsoft 계정 삽니다",
            (
                "순수 게임 이용 목적으로 NCsoft 계정 최대 10개까지 삽니다. "
                "한국 휴대폰 본인 인증이 가능해야 하며 직접 만나 거래합니다."
            ),
            "https://community.example/game/1",
            "unknown",
            "intent",
        )
        self.assertEqual(reason, "excluded_single_account_trade")

    def test_press_page_classification_uses_path_and_byline(self) -> None:
        page_type = classify_page_type(
            "https://media.example/news/articleView.html?idxno=10",
            '"ID 삽니다" 국내외 계정 매매 성행',
            "홍길동 기자 = 관련 업계에 따르면 계정 거래가 늘고 있다.",
        )
        self.assertEqual(page_type, "news_or_education")

    def test_newsroom_page_classification_does_not_require_a_byline(self) -> None:
        page_type = classify_page_type(
            "https://security.example/about/news_list/view/360",
            "모바일 주민등록증 시스템 구축 - Security Newsroom",
            "자사 기술을 적용해 관련 사업을 수주했으며 시스템을 구축할 계획이다.",
        )
        self.assertEqual(page_type, "news_or_education")

    def test_strict_gate_rejects_reporting_and_generic_links(self) -> None:
        reporting = relevance_gate_reason(
            "고객 DB 판매 게시물 적발",
            "경찰이 텔레그램을 통해 거래한 사건을 검거했다는 기사입니다.",
            "https://news.example/article/1",
            "news_or_education",
            "strict",
        )
        generic_link = relevance_gate_reason(
            "고객 DB 판매합니다",
            "대량 보유 중이며 주문 문의는 홈페이지 [CONTACT_URL]에서 받습니다.",
            "https://board.example/post/1",
            "unknown",
            "strict",
        )
        self.assertEqual(reporting, "excluded_page_type")
        self.assertEqual(generic_link, "missing_concrete_contact")

    def test_strict_gate_rejects_scam_warning_disguised_by_sales_title(self) -> None:
        reason = relevance_gate_reason(
            "SNS 계정 싸게 팝니다",
            "사기입니다. 돈을 보낸 뒤 피해를 입어 사이버수사대에 신고했습니다.",
            "https://community.example/post/2",
            "unknown",
            "strict",
            "SNS 계정 판매 연락 [PHONE]",
        )
        self.assertEqual(reason, "excluded_reporting_context")

    def test_strict_gate_rejects_bulk_messaging_service_context(self) -> None:
        reason = relevance_gate_reason(
            "대량문자 발송 서비스",
            (
                "연락처를 주소록에 저장해 대량문자를 발송합니다. "
                "회원 할인 단가를 제공하며 상담은 텔그 [MESSENGER_ID]"
            ),
            "https://blog.example/post/3",
            "unknown",
            "strict",
        )
        self.assertEqual(reason, "missing_relevant_target")

    def test_strict_gate_keeps_account_verification_buying_post(self) -> None:
        reason = relevance_gate_reason(
            "가입인증 삽니다",
            "성인 본인인증 자료를 매입합니다. 문의 텔레그램 [MESSENGER_ID]",
            "https://board.example/post/4",
            "unknown",
            "strict",
        )
        self.assertEqual(reason, "")

    def test_strict_gate_rejects_single_game_account_trade(self) -> None:
        reason = relevance_gate_reason(
            "게임 계정 하나 팝니다",
            "실사용하던 계정을 10만원에 판매합니다. 문의 [PHONE]",
            "https://market.example/post/5",
            "unknown",
            "strict",
        )
        self.assertEqual(reason, "excluded_single_account_trade")

        linked_platform_account = relevance_gate_reason(
            "구글 계정 팝니다",
            "로드 모바일에서 실사용하던 구글 연동 계정을 판매합니다. 문의 [PHONE]",
            "https://market.example/post/6",
            "unknown",
            "strict",
        )
        self.assertEqual(linked_platform_account, "excluded_single_account_trade")

    def test_strict_gate_keeps_forgery_service_offer(self) -> None:
        reason = relevance_gate_reason(
            "각종 신분증 위조 전문",
            "주민등록증과 운전면허증 제작 가능. 의뢰 문의 텔레그램 [MESSENGER_ID]",
            "https://board.example/post/7",
            "unknown",
            "strict",
        )
        self.assertEqual(reason, "")

    def test_strict_gate_does_not_treat_body_link_words_as_document_type(self) -> None:
        reason = relevance_gate_reason(
            "스토리 계좌매입 채널",
            (
                "계좌매입 후 즉시 정산합니다. 편하게 문의주십쇼 "
                "텔레그램 [MESSENGER_ID]. 채널 바로가기와 AV위키 제휴 안내."
            ),
            "https://t.me/public_channel",
            "unknown",
            "strict",
        )
        self.assertEqual(reason, "")

    def test_strict_gate_rejects_explicitly_negated_offer(self) -> None:
        reason = relevance_gate_reason(
            "종토방 제휴",
            (
                "유심, 통장대여, 코인이체, 계정매입 등 피싱과 관련된 "
                "제휴는 받지 않습니다. 텔레그램 [MESSENGER_ID]"
            ),
            "https://t.me/public_channel",
            "public_messenger_page",
            "strict",
        )
        self.assertNotEqual(reason, "")

    def test_strict_gate_can_use_strong_discovery_evidence(self) -> None:
        reason = relevance_gate_reason(
            "서비스 홍보",
            "주식 고객 DB를 대량 보유하고 판매합니다. 건당 단가 안내 가능합니다.",
            "https://community.example/service/1",
            "unknown",
            "strict",
            "주식 디비 판매합니다. 텔레그램 https://t.me/private_handle 문의",
        )
        self.assertEqual(reason, "")

    def test_strict_gate_rejects_unrelated_destination_despite_search_snippet(self) -> None:
        reason = relevance_gate_reason(
            "기업 홈페이지",
            "ICT 인프라 구축과 기술 컨설팅 서비스를 제공합니다.",
            "https://company.example/",
            "unknown",
            "strict",
            "고객 DB 판매합니다. 텔레그램 https://t.me/private_handle 문의",
        )
        self.assertEqual(reason, "missing_relevant_target")

    def test_input_parameter_reflection_is_classified_as_search_reflection(self) -> None:
        page_type = classify_page_type(
            "https://calculator.example/input?i=customer+db+sale",
            "customer db sale - calculator",
            "customer db sale natural language input",
        )
        self.assertEqual(page_type, "search_reflection")

    def test_detail_route_with_record_id_ignores_search_form_state(self) -> None:
        self.assertEqual(
            classify_page_type(
                "https://board.example/bbs/view.php?no=36420&search=subject",
                "토익 성적표 위조 판매",
                "제작 전 확인 후 배송하며 텔레그램 [ACCOUNT] 문의",
            ),
            "unknown",
        )
        self.assertEqual(
            classify_page_type(
                "https://board.example/board/read?article_id=7&query=subject",
                "운전면허증 위조 제작",
                "직거래 가능하며 카톡 [ACCOUNT] 문의",
            ),
            "unknown",
        )

    def test_search_route_and_detail_route_without_id_remain_search_pages(self) -> None:
        self.assertEqual(
            classify_page_type(
                "https://board.example/search?no=36420&search=subject",
                "게시물 검색",
                "제목 검색 결과 목록",
            ),
            "search_result_list",
        )
        self.assertEqual(
            classify_page_type(
                "https://board.example/bbs/view.php?search=subject",
                "게시물 검색",
                "제목 검색 결과 목록",
            ),
            "search_result_list",
        )

    def test_address_search_p1_parameter_reflection_is_structural(self) -> None:
        query = "카카오톡계정판매 텔레 seller123"
        page_type = classify_page_type(
            "https://address.example/addreng.asp?p1=" + quote_plus(query),
            query,
            (
                "지역명에 번지수와 건물명을 붙여 검색하세요. 전체 주소전체 "
                "지번주소 도로명주소 한글주소 영문주소 우편번호"
            ),
        )
        self.assertEqual(page_type, "search_reflection")
        self.assertEqual(
            classify_page_type(
                "https://address.example/addreng.asp?p1=" + quote_plus(query),
                mask_text(query),
                (
                    "전체 주소전체 지번주소 도로명주소 한글주소 "
                    "영문주소 우편번호"
                ),
            ),
            "search_reflection",
        )

    def test_dsearch_and_invalid_slug_error_are_search_reflections(self) -> None:
        query = "네이버생성아이디거래 텔레 seller123"
        self.assertEqual(
            classify_page_type(
                "https://search.example/dsearch.php?k1=" + quote_plus(query),
                query + " : 통합검색",
                f'"{query}" 에 대한 상품 검색결과가 없습니다.',
            ),
            "search_reflection",
        )
        self.assertEqual(
            classify_page_type(
                "https://city.example/in/%EC%9D%B8%EC%8A%A4%ED%83%80%EA%B3%84%EC%A0%95%ED%8C%90%EB%A7%A4",
                "Cannot find city id",
                "Cannot find city id for 인스타계정판매",
            ),
            "search_reflection",
        )

    def test_integrated_zero_result_search_is_a_search_reflection(self) -> None:
        query = "카카오톡계정 판매 seller123"
        self.assertEqual(
            classify_page_type(
                "https://institute.example/intgSearch.do?intgSw="
                + quote_plus(query),
                "연구원",
                f"검색어 {query}에 대한 검색결과 총 0건이 있습니다.",
            ),
            "search_reflection",
        )

    def test_story_tag_and_ranked_vendor_lists_are_structural_pages(self) -> None:
        self.assertEqual(
            classify_page_type(
                "https://fiction.example/stories/google-account-sale",
                "Google account sale Stories",
                "Refine by tag: google-account-sale 1 Story Sort by: Hot",
            ),
            "search_result_list",
        )
        self.assertEqual(
            classify_page_type(
                "https://www.wattpad.com/stories/youtube-account-sale",
                "YouTube account sale Stories - Wattpad",
                "One preview contains a seller advertisement and contact link.",
            ),
            "search_result_list",
        )
        self.assertEqual(
            classify_page_type(
                "https://ranking.example/",
                "네이버 아이디 판매 업체",
                (
                    "순위 서비스 타입 종합 점수 바로가기 "
                    "1 HOT 98.2 / 100 2 RISING 94.2 / 100 "
                    "3 NEW 88.7 / 100"
                ),
            ),
            "search_result_list",
        )

    def test_forum_index_is_classified_as_search_result_list(self) -> None:
        page_type = classify_page_type(
            "https://board.example/pds",
            "자료실",
            "번호 제목 작성자 작성일 추천 조회 3010 개인통장 매입 문의",
        )
        self.assertEqual(page_type, "search_result_list")

    def test_board_php_without_record_id_is_a_board_listing(self) -> None:
        self.assertEqual(
            classify_page_type(
                "https://community.example/bbs/board.php?bo_table=market",
                "중고거래장터 1 페이지",
                "제목 유흥 DB 판매합니다 2025.03.12 다음 글 목록",
            ),
            "board_listing",
        )
        self.assertEqual(
            classify_page_type(
                "https://community.example/bbs/board.php?bo_table=market&wr_id=7",
                "유흥 DB 판매합니다",
                "실시간 고객 DB를 판매합니다. 문의 [ACCOUNT]",
            ),
            "unknown",
        )

    def test_legacy_bbs_list_endpoint_is_a_board_listing(self) -> None:
        page_type = classify_page_type(
            "https://sellerocean.example/bbs_list.php?tb=board_bestseller",
            "매매게시판 : 셀러오션",
            (
                "제목에 특수기호 이모티콘등은 삭제 대상입니다\n"
                "상표권 판매합니다.\n먹스타그램 계정 양도 희망합니다\n"
                "최적화 카페 양도\n부업 최적화 스마트스토어 양도합니다"
            ),
        )
        self.assertEqual(page_type, "board_listing")

    def test_cafe24_board_list_route_is_a_board_listing(self) -> None:
        page_type = classify_page_type(
            "https://shop.example/board/gallery/list.html?board_no=8",
            "건축자재 전문 도매몰",
            (
                "운전면허증위조 작성자 제작소신 작성일 2025-05-14 "
                "운전면허증위조 작성자 제작소신 작성일 2025-05-14"
            ),
        )
        self.assertEqual(page_type, "board_listing")
        category_page_type = classify_page_type(
            "https://shop.example/board/%EC%83%81%ED%92%88-qa/6/",
            "상품 Q&A - Example Shop",
            "\n".join(
                f"개인통장판매 작성자 2026-07-0{day} 11:02:30 조회 {day}"
                for day in range(1, 5)
            ),
        )
        self.assertEqual(category_page_type, "board_listing")
        self.assertEqual(
            classify_page_type(
                "https://shop.example/article/%EC%83%81%ED%92%88-qa/6/2900/",
                "선불심 판매 상품 Q&A - Example Shop",
                "타인 명의 선불심을 판매합니다. 텔레그램 [ACCOUNT] 문의",
            ),
            "unknown",
        )
        self.assertEqual(
            relevance_gate_reason(
                "매매게시판 : 셀러오션",
                "상표권 판매합니다. 먹스타그램 계정 양도 희망합니다.",
                "https://sellerocean.example/bbs_list.php?tb=board_bestseller",
                page_type,
                "intent",
            ),
            "excluded_page_type",
        )

        mode_list = classify_page_type(
            "http://alumni.example/xboard/board.php?mode=list&tbnum=13",
            "HOME > 공지사항",
            (
                "공지사항 게시물 목록 연번 제목 작성자 작성일 조회 파일 "
                "1251 네이버실명아이디판매 텔레그램계정판매 홍채영 "
                "1250 동창골프 결산보고 이홍철"
            ),
        )
        self.assertEqual(mode_list, "board_listing")
        self.assertEqual(
            relevance_gate_reason(
                "HOME > 공지사항",
                "게시물 목록 1251 네이버실명아이디판매 텔레그램계정판매",
                "http://alumni.example/xboard/board.php?mode=list&tbnum=13",
                mode_list,
                "intent",
            ),
            "excluded_page_type",
        )

    def test_public_telegram_page_has_distinct_page_type(self) -> None:
        page_type = classify_page_type(
            "https://t.me/public_channel",
            "계좌매입 채널",
            "계좌 매입 문의 텔레그램 [MESSENGER_ID]",
        )
        self.assertEqual(page_type, "public_messenger_page")

    def test_labeling_gate_keeps_topical_hard_negative(self) -> None:
        reason = relevance_gate_reason(
            "개인정보 유출 사고 안내",
            "피해 확인 방법을 설명하며 판매나 구매 의사는 없는 기사입니다.",
            "https://news.example/article/7",
            "news_or_education",
            "labeling",
        )
        self.assertEqual(reason, "")
        sidebar_noise = relevance_gate_reason(
            "여름맞이 경품 이벤트",
            "이벤트 안내입니다. 인기글: 고객 DB 판매 관련 문의",
            "https://community.example/event/9",
            "unknown",
            "labeling",
        )
        self.assertEqual(sidebar_noise, "missing_relevant_target")
        trade_title = relevance_gate_reason(
            "대량 판매합니다",
            "자세한 거래 대상은 게시물 본문을 확인하세요.",
            "https://community.example/post/10",
            "unknown",
            "labeling",
        )
        contact_with_target_lead = relevance_gate_reason(
            "텔레그램 문의",
            "고객 DB와 계정 관련 내용을 안내합니다.",
            "https://community.example/post/11",
            "unknown",
            "labeling",
        )
        generic_contact = relevance_gate_reason(
            "문의하기",
            "행사 참여 방법을 안내합니다.",
            "https://community.example/event/12",
            "unknown",
            "labeling",
        )
        self.assertEqual(trade_title, "")
        self.assertEqual(contact_with_target_lead, "")
        self.assertEqual(generic_contact, "missing_relevant_target")

    def test_relevance_gate_rejects_policy_and_search_pages(self) -> None:
        policy_reason = relevance_gate_reason(
            "개인정보 처리방침",
            "고객정보를 보유하며 상품 구매 문의는 고객센터로 연락하세요.",
            "https://shop.example/privacy",
            "unknown",
            "review",
        )
        search_reason = relevance_gate_reason(
            "고객 DB 판매 검색",
            "검색 결과입니다. 텔레그램 문의",
            "https://board.example/search?q=test",
            "search_reflection",
            "review",
        )
        self.assertEqual(policy_reason, "excluded_document_type")
        self.assertEqual(search_reason, "excluded_page_type")
        self.assertEqual(
            relevance_gate_reason(
                "검색 결과",
                "공개 게시물 목록",
                "https://board.example/search?q=test",
                "search_result_list",
                "off",
            ),
            "excluded_page_type",
        )

    def test_review_gate_requires_title_signal_but_keeps_topical_news(self) -> None:
        generic = relevance_gate_reason(
            "전자계약 서비스 주요 기능",
            "고객 개인정보를 보유하며 계약 거래 문의는 [EMAIL]로 받습니다.",
            "https://service.example/features",
            "unknown",
            "review",
        )
        generic_account_trade = relevance_gate_reason(
            "거래중지 계좌 해지 방법",
            "오래 사용하지 않은 통장의 거래를 다시 시작하고 싶습니다.",
            "https://qna.example/question/2",
            "unknown",
            "review",
        )
        topical_news = relevance_gate_reason(
            "고객 DB 판매 게시물 적발",
            "개인정보 명단을 대량으로 거래한 사례가 확인됐다는 보도입니다.",
            "https://news.example/article/1",
            "news_or_education",
            "review",
        )
        self.assertEqual(generic, "missing_title_signal")
        self.assertEqual(generic_account_trade, "missing_title_signal")
        self.assertEqual(topical_news, "")

    def test_seed_csv_deduplicates_urls(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "seeds.local.csv"
            path.write_text(
                "url,detection_type,query_group\n"
                "https://example.com/post,기타,seed\n"
                "https://example.com/post#fragment,기타,seed\n",
                encoding="utf-8",
            )
            candidates = load_seed_candidates(path)
            self.assertEqual(len(candidates), 1)
            self.assertEqual(candidates[0].url, "https://example.com/post")

    def test_seed_csv_accepts_restricted_raw_url_column(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "urls.csv"
            path.write_text(
                "sample_id,raw_url\n"
                "LP-000001,https://example.com/public/post\n",
                encoding="utf-8",
            )
            candidates = load_seed_candidates(path)
            self.assertEqual(len(candidates), 1)
            self.assertEqual(candidates[0].url, "https://example.com/public/post")

    def test_seed_csv_accepts_shareable_source_url_column(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "data.csv"
            path.write_text(
                "sample_id,source_url\n"
                "LP-000001,https://example.com/public/post\n",
                encoding="utf-8",
            )
            candidates = load_seed_candidates(path)
            self.assertEqual(len(candidates), 1)
            self.assertEqual(candidates[0].url, "https://example.com/public/post")

    def test_prior_search_queue_is_prefiltered_but_manual_seed_is_retained(self) -> None:
        direct = Candidate(
            "https://example.com/direct",
            "group",
            "개인정보DB",
            discovery_text=(
                "고객 DB 판매합니다. 텔레그램 https://t.me/direct 문의"
            ),
        )
        reporting = Candidate(
            "https://example.com/report",
            "group",
            "개인정보DB",
            discovery_text="고객 DB 판매 사건을 경찰이 적발했다는 기사",
        )
        manual = Candidate(
            "https://example.com/manual",
            "seed",
            "기타",
            source_type="seed",
        )
        empty_search_artifact = Candidate(
            "https://search.example/navigation",
            "group",
            "기타",
            source_type="search",
        )
        filtered = prefilter_seed_candidates(
            [direct, reporting, manual, empty_search_artifact], "strict"
        )
        self.assertEqual(filtered, [direct, manual])

    def test_resume_prefilter_can_tighten_a_broad_search_queue(self) -> None:
        broad_false_positive = Candidate(
            "https://news.example/report",
            "group",
            "기타",
            discovery_text="불법 신분증 위조 판매 사건을 경찰이 적발했다",
            search_provider="naver",
        )
        direct_offer = Candidate(
            "https://seller.example/post",
            "group",
            "기타",
            discovery_text="신분증 위조 제작합니다 텔레그램 문의",
            search_provider="naver",
        )
        self.assertEqual(
            prefilter_seed_candidates(
                [broad_false_positive, direct_offer],
                "intent",
            ),
            [direct_offer],
        )

    def test_private_candidate_queue_round_trip(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / ".private" / "candidate_queue.jsonl"
            expected = [Candidate("https://example.com/post", "group", "기타")]
            save_candidate_queue(path, expected)
            loaded = load_candidate_queue(path)
            self.assertEqual(loaded, expected)
            self.assertEqual(load_seed_candidates(path), expected)
            self.assertEqual(path.stat().st_mode & 0o777, 0o600)

    def test_refresh_discovery_preserves_previous_nonempty_queue(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / ".private" / "candidate_queue.jsonl"
            expected = [Candidate("https://example.com/post", "group", "기타")]
            save_candidate_queue(path, expected)
            backup_path = backup_candidate_queue(path)
            self.assertIsNotNone(backup_path)
            self.assertEqual(load_candidate_queue(backup_path), expected)
            self.assertEqual(backup_path.stat().st_mode & 0o777, 0o600)

            path.write_text("", encoding="utf-8")
            self.assertIsNone(backup_candidate_queue(path))

    def test_prior_sample_urls_can_be_loaded_for_holdout_exclusion(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "prior.csv"
            path.write_text(
                "source_url,title\n"
                "https://example.com/post?a=1&utm_source=test,first\n"
                "https://example.com/other,second\n",
                encoding="utf-8",
            )
            self.assertEqual(
                load_excluded_urls([path]),
                {
                    "https://example.com/post?a=1",
                    "https://example.com/other",
                },
            )

    def test_naver_blog_desktop_mobile_urls_share_document_identity(self) -> None:
        desktop = canonicalize_url(
            "https://blog.naver.com/qnaj2330/223196767893?from=postList"
        )
        mobile = canonicalize_url(
            "https://m.blog.naver.com/qnaj2330/223196767893"
        )
        self.assertEqual(desktop, mobile)

    def test_marketplace_post_search_parameters_do_not_create_new_documents(self) -> None:
        plain = canonicalize_url("https://creativebox.kr/igtrade/5742")
        filtered = canonicalize_url(
            "https://www.creativebox.kr/igtrade/5742?"
            "sfl=mb_id%2C1&stx=example&page=3"
        )
        self.assertEqual(plain, filtered)

    def test_cafe24_article_page_context_does_not_create_new_documents(self) -> None:
        plain = canonicalize_url(
            "https://mongata.ai/article/1문의/1002/1836/"
        )
        contextual = canonicalize_url(
            "https://mongata.ai/article/1%EB%AC%B8%EC%9D%98/1002/1836/"
            "page/40/?board_no=1002&page=40&utm_source=board"
        )
        self.assertEqual(plain, contextual)

    def test_gnuboard_search_context_does_not_create_new_documents(self) -> None:
        plain = canonicalize_url(
            "https://forum.example/bbs/board.php?bo_table=market&wr_id=712"
        )
        contextual = canonicalize_url(
            "https://forum.example/bbs/board.php?"
            "bo_table=market&sca=trade&sfl=wr_subject&stx=sample&wr_id=712&page=4"
        )
        self.assertEqual(plain, contextual)

    def test_post_identity_uses_board_and_post_id(self) -> None:
        first = post_identity_descriptor(
            "https://www.example.co.kr/bbs/board.php?"
            "bo_table=market&wr_id=712&sfl=wr_subject&stx=sample"
        )
        second = post_identity_descriptor(
            "http://example.co.kr/bbs/board.php?bo_table=market&wr_id=712"
        )
        different_post = post_identity_descriptor(
            "https://example.co.kr/bbs/board.php?bo_table=market&wr_id=713"
        )
        self.assertEqual(first, second)
        self.assertNotEqual(first, different_post)

    def test_url_canonicalization_preserves_ipv6_brackets(self) -> None:
        self.assertEqual(
            canonicalize_url("https://[2001:4860:4860::8888]/post#fragment"),
            "https://[2001:4860:4860::8888]/post",
        )

    def test_prior_sample_fingerprints_are_loaded_for_holdout_exclusion(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "prior.csv"
            path.write_text(
                "source_url,near_duplicate_fingerprint\n"
                "https://example.com/one,simhash64:1111\n"
                "https://example.com/two,simhash64:2222\n",
                encoding="utf-8",
            )
            self.assertEqual(
                load_excluded_fingerprints([path]),
                {"simhash64:1111", "simhash64:2222"},
            )

    def test_terminal_failures_are_skipped_on_resume(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "collection_log.csv"
            path.write_text(
                "url_hmac,query_group,outcome,http_status,reason,text_chars,extraction_method\n"
                "done,g,skipped,,robots_disallowed,0,\n"
                "filtered,g,skipped,200,missing_relevant_target,100,main_container\n"
                "oversized,g,skipped,200,content_too_large,0,\n"
                "nonhtml,g,skipped,200,non_html_content,0,\n"
                "domain,g,skipped,200,reserved_for_domain_diversity,100,main_container\n"
                "type,g,skipped,200,reserved_for_type_diversity,100,main_container\n"
                "campaign,g,skipped,200,campaign_record_limit,100,main_container\n"
                "short,g,failed,200,insufficient_text,20,visible_body_fallback\n"
                "retry,g,failed,,ReadTimeout,0,\n",
                encoding="utf-8",
            )
            self.assertEqual(
                terminal_attempt_hashes(path),
                {"done", "filtered", "oversized", "nonhtml", "short"},
            )
            self.assertEqual(
                terminal_attempt_hashes(path, retry_filtered=True),
                {"done", "short"},
            )
            self.assertEqual(
                terminal_attempt_hashes(
                    path,
                    retry_filtered=True,
                    retry_renderable=True,
                ),
                {"done"},
            )

    def test_collection_log_upgrade_adds_attempt_time_column(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "collection_log.csv"
            path.write_text(
                "url_hmac,query_group,outcome,http_status,reason,text_chars,extraction_method\n"
                "abc,g,failed,200,insufficient_text,20,visible_body_fallback\n",
                encoding="utf-8",
            )
            upgrade_collection_log_schema(path)
            with path.open(encoding="utf-8-sig", newline="") as handle:
                reader = csv.DictReader(handle)
                rows = list(reader)
            self.assertIn("attempted_at", reader.fieldnames)
            self.assertEqual(rows[0]["text_chars"], "20")

    def test_existing_dataset_upgrade_preserves_legacy_fields(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "candidates_masked.csv"
            legacy_fields = [
                field
                for field in SCHEMA
                if field not in {"extraction_status", "near_duplicate_fingerprint"}
            ]
            row = {field: "" for field in legacy_fields}
            row.update(
                {
                    "sample_id": "EG-0001",
                    "source_type": "public_web_search",
                    "live_status": "true",
                    "page_type": "reflected_search_page",
                    "near_duplicate_cluster": "simhash64:1234567890abcdef",
                }
            )
            with path.open("w", encoding="utf-8-sig", newline="") as handle:
                writer = csv.DictWriter(handle, fieldnames=legacy_fields)
                writer.writeheader()
                writer.writerow(row)
            upgrade_existing_csv_schema(path)
            with path.open(encoding="utf-8-sig", newline="") as handle:
                upgraded = next(csv.DictReader(handle))
            self.assertEqual(upgraded["sample_id"], "EG-0001")
            self.assertEqual(upgraded["source_type"], "search")
            self.assertEqual(upgraded["live_status"], "accessible")
            self.assertEqual(upgraded["page_type"], "search_reflection")
            self.assertEqual(
                upgraded["near_duplicate_fingerprint"],
                upgraded["near_duplicate_cluster"],
            )

    def test_extraction_failure_has_standard_status_and_no_raw_url(self) -> None:
        log = CollectionLog(
            "digest-only",
            "group",
            "failed",
            "200",
            "insufficient_text",
            25,
            "main_container_short",
        )
        failure = extraction_failure_record(log)
        self.assertIsNotNone(failure)
        self.assertEqual(failure["extraction_status"], "partial")
        self.assertRegex(failure["attempted_at"], r"^\d{4}-\d{2}-\d{2}T")

    def test_record_uses_standard_handoff_values(self) -> None:
        candidate = Candidate(
            "https://example.com/post/1", "seed-group", "기타", source_type="seed"
        )
        record = make_record(
            7,
            candidate,
            candidate.url,
            200,
            "테스트 제목",
            "충분한 공개 예시 본문입니다. " * 10,
            b"test-key",
        )
        self.assertEqual(record["sample_id"], "EG-000007")
        self.assertEqual(record["source_type"], "seed")
        self.assertEqual(record["live_status"], "accessible")
        self.assertEqual(record["extraction_status"], "success")
        self.assertEqual(
            record["near_duplicate_cluster"],
            record["near_duplicate_fingerprint"],
        )

    def test_masking_report_and_manifest_include_integrity_data(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            out = Path(temp_dir)
            csv_path = out / "candidates_masked.csv"
            csv_path.write_text(
                "masked_title,masked_text\n제목,연락처 [PHONE]\n",
                encoding="utf-8",
            )
            report = masking_validation(csv_path, "pilot-v1")
            self.assertTrue(report["passed"])
            manifest = data_manifest(out, "pilot-v1", [csv_path], {"target": 1})
            self.assertEqual(manifest["files"][0]["rows"], 1)
            self.assertRegex(manifest["files"][0]["sha256"], r"^[0-9a-f]{64}$")
            self.assertEqual(manifest["settings"]["target"], 1)

    def test_masking_report_detects_obfuscated_mobile_phone(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            csv_path = Path(temp_dir) / "candidates_masked.csv"
            csv_path.write_text(
                "masked_title,masked_text\n문의,ⓞ①ⓞ-③④③⑤-⑤①⑦⑦\n",
                encoding="utf-8",
            )
            report = masking_validation(csv_path, "pilot-v1")
            self.assertFalse(report["passed"])
            self.assertEqual(report["residual_obfuscated_phone_hits"], 1)

    def test_extracts_short_post_from_main_container(self) -> None:
        html = """
        <html><head><title>테스트 게시물</title></head><body>
        <nav>메뉴 메뉴 메뉴</nav>
        <main><p>이것은 본문 추출 테스트를 위한 공개 예시 문장입니다.</p>
        <p>게시물의 핵심 내용이 메뉴보다 우선해서 저장되어야 합니다.</p></main>
        </body></html>
        """
        title, text, method = extract_title_text(html, "https://example.com/post")
        self.assertEqual(title, "테스트 게시물")
        self.assertIn("본문 추출 테스트", text)
        self.assertNotIn("메뉴 메뉴", text)
        self.assertIn(method, {"trafilatura_precision", "main_container"})

    def test_legacy_board_post_beats_longer_footer(self) -> None:
        html = """
        <html><head><title>여권발급이나 새신분이 필요하신분</title></head>
        <body>
          <table><tr><td class="con_f">
            여권과 신분증 위조 제작 가능합니다. 판매 문의는
            텔레그램 sample_handle 또는 카톡 sample_chat으로 주세요.
            신청 대상과 제작 종류를 확인한 뒤 신속하게 안내한다는 게시물입니다.
          </td></tr></table>
          <footer>고객센터 문의와 저작권 안내입니다. """ + "일반 안내 " * 80 + """</footer>
        </body></html>
        """
        title, text, method = extract_title_text(
            html, "https://board.example/public/post/1"
        )
        self.assertIn("여권과 신분증 위조 제작", text)
        self.assertNotIn("저작권 안내", text)
        self.assertEqual(method, "strong_post_container")

    def test_dcinside_write_container_beats_truncated_title(self) -> None:
        html = """
        <html><head><title>개인통장 팝니다 010-3435-517</title></head><body>
          <div class="writing_view_box"><div class="write_div">
            <p>개인통장과 법인통장을 판매합니다.</p>
            <p>빠른 배송과 한 달 A/S를 보장합니다.</p>
            <p>인터넷뱅킹과 보안카드 옵션은 상담 후 바로 안내합니다.</p>
            <p>문의 ⓞ①ⓞ-③④③⑤-⑤①⑦⑦</p>
          </div></div>
        </body></html>
        """
        title, text, method = extract_title_text(
            html,
            "https://gall.dcinside.com/board/view/?id=iphone&no=529622",
        )
        self.assertIn("법인통장을 판매", text)
        self.assertIn("⑤①⑦⑦", text)
        self.assertEqual(method, "strong_post_container")

    def test_challenge_page_fails_text_quality_gate(self) -> None:
        text = "Checking your browser before accessing the requested public page."
        self.assertEqual(
            text_quality_reason(text, 40), "challenge_or_access_page"
        )

    def test_korean_text_gate_rejects_english_only_page(self) -> None:
        text = "This is a sufficiently long English page with meaningful words."
        self.assertEqual(
            text_quality_reason(text, 40, minimum_korean_chars=5),
            "insufficient_korean_text",
        )

    def test_related_internal_links_are_bounded_and_same_site(self) -> None:
        html = """
        <a href="/board/view?id=2">고객 DB 관련 게시물</a>
        <a href="/login">로그인</a>
        <a href="https://outside.example/post/3">개인정보</a>
        <a href="/search?q=개인정보">검색</a>
        """
        links = discover_related_internal_links(
            html, "https://example.com/board/view?id=1", 5
        )
        self.assertEqual(links, ["https://example.com/board/view?id=2"])

    def test_template_is_preserved_and_rows_expand(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            output = Path(temp_dir) / "restricted" / "result.xlsx"
            workbook, sheet = prepare_detection_workbook(TEMPLATE, output, resume=False)
            entries = [
                DetectionEntry(
                    detected_on=dt.date(2026, 8, 17),
                    url=f"https://example.com/post/{index}",
                    detection_type="개인정보DB",
                    registrant="테스트",
                )
                for index in range(1, 35)
            ]
            append_detection_entries(sheet, entries)
            save_restricted_workbook(workbook, output)

            saved = load_workbook(output)["8월"]
            self.assertEqual(saved["A4"].value, 1)
            self.assertEqual(saved["A37"].value, 34)
            self.assertEqual(saved["D4"].value, "개인정보DB")
            self.assertEqual(saved["E4"].value, "테스트")
            self.assertEqual(saved["C4"].hyperlink.target, "https://example.com/post/1")
            self.assertIn("예시 내용", saved["H4"].value)
            self.assertEqual(output.stat().st_mode & 0o777, 0o600)


if __name__ == "__main__":
    unittest.main()
